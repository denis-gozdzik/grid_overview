"""Serialized workbook regression checks for evidence readiness presentation."""
from openpyxl import load_workbook
import pytest

from infoblox_inventory.models import CollectionResult
from infoblox_inventory.profile_readiness import PROFILE_READINESS_HEADERS, PROFILE_SUMMARY_HEADERS
from infoblox_inventory.report import write_reports
from infoblox_inventory.xlsx_validation import validate_xlsx


def _result():
    reference = 'network/example:192.0.2.0/24/default'
    raw = {'_ref': reference, 'network': '192.0.2.0/24', 'network_view': 'default',
           'options': [{'num': 51, 'name': 'dhcp-lease-time', 'value': '3600', 'use_option': True}]}
    return CollectionResult(
        grid='LAB', grid_url='https://grid.example', wapi_version='2.13.7',
        records={'network': [raw], 'range': []},
        effective_records={'network': [{**raw, 'options': [
            {'inherited': False, 'source': '', 'values': raw['options']}]}], 'range': []},
        schemas={'network': {'fields': [{'name': 'options', 'supports': 'r',
                                        'overridden_by': 'use_options'}]},
                 'range': {'fields': [{'name': 'options', 'supports': 'r',
                                      'overridden_by': 'use_options'}]}},
        coverage=[{'Grid': 'LAB', 'Object': object_type, 'Query': query,
                   'Collection Status': 'COMPLETE' if object_type == 'network' else 'EMPTY',
                   'Objects Found': 1 if object_type == 'network' else 0}
                  for object_type in ('network', 'range') for query in ('raw', 'effective')],
    )


@pytest.fixture
def report_path(tmp_path):
    write_reports([_result()], tmp_path)
    return tmp_path / 'current_state_inventory.xlsx'


def _rows(sheet):
    return [dict(zip([cell.value for cell in sheet[1]], values))
            for values in sheet.iter_rows(min_row=2, values_only=True)]


def test_matrix_sheet_order_candidate_order_and_serialized_table_validity(report_path):
    validated = validate_xlsx(report_path)
    assert validated['zip_crc_ok'] is True
    assert validated['sheet_rows']['Profile_Readiness'] == 39
    workbook = load_workbook(report_path)
    try:
        assert workbook.sheetnames[:13] == [
            'Overview', 'Profile_Readiness', 'Profile_Populations', 'Profile_Usefulness',
            'Profile_Discovery', 'Profiles', 'Profile_Objects', 'Standardization',
            'Decisions', 'Exceptions', 'Grid_Comparison', 'Coverage', 'Manual_Review',
        ]
        sheet = workbook['Profile_Readiness']
        assert [cell.value for cell in sheet[1]] == list(PROFILE_READINESS_HEADERS)
        rows = _rows(sheet)
        assert [row['Profile'] for row in rows] == ['Network Profile v1'] * 21 + ['Range Profile v1'] * 18
        assert all(row['Scope'] == 'Network' for row in rows[:21])
        assert all(row['Scope'] == 'Range' for row in rows[21:])
        lease = next(row for row in rows if row['Parameter ID'] == 'dhcp.lease_time.network')
        assert lease['Population Objects'] == lease['Confirmed Objects'] == 1
        assert lease['Resolved Evidence %'] == 100
        assert lease['Readiness'] == 'READY'
        assert all(row['Readiness'] == 'NOT_APPLICABLE' for row in rows[21:])
        assert sheet.freeze_panes == 'A2'
        assert sheet.auto_filter.ref is None
        assert len(sheet.tables) == 1
        table = next(iter(sheet.tables.values()))
        assert table.autoFilter.ref == table.ref == 'A1:X40'
        assert all(cell.data_type != 'f' for row in sheet for cell in row)
        populations = workbook['Profile_Populations']
        population_rows = _rows(populations)
        assert any(row['Population Basis'] == 'ALL_NETWORKS' and row['Object Count'] == 1
                   for row in population_rows)
        assert any(row['Population Basis'] == 'DHCP_RELEVANT_NETWORK_CANDIDATES'
                   and row['Object Count'] == 1 for row in population_rows)
        assert len(_rows(workbook['Profile_Usefulness'])) == 39
        discovery = _rows(workbook['Profile_Discovery'])
        assert {row['Profile Type'] for row in discovery} == {'Network', 'Range'}
        network = next(row for row in discovery if row['Profile Type'] == 'Network')
        assert network['Applicable Objects'] == network['Profiled Objects'] == 1
        assert network['Distinct Profiles'] == 1
        assert len(_rows(workbook['Profiles'])) == 1
        assignments = _rows(workbook['Profile_Objects'])
        assert len(assignments) == 1
        assert assignments[0]['Profile Status'] == 'PROFILED'
    finally:
        workbook.close()


def test_overview_summary_has_input_denominator_and_both_profile_drilldowns(report_path):
    workbook = load_workbook(report_path)
    try:
        overview = workbook['Overview']
        title = next(cell for row in overview for cell in row if cell.value == 'Profile discovery readiness')
        header_row = title.row + 1
        assert [overview.cell(header_row, column).value for column in range(1, 10)] == list(PROFILE_SUMMARY_HEADERS)
        assert overview.cell(header_row, 10).value == 'Input details'
        summaries = [dict(zip(PROFILE_SUMMARY_HEADERS,
                              [overview.cell(row, column).value for column in range(1, 10)]))
                     for row in range(header_row + 1, header_row + 3)]
        assert summaries[0]['Candidate Inputs'] == 21
        assert summaries[0]['Applicable Inputs'] == 20
        assert summaries[0]['READY'] == 8
        assert summaries[0]['NOT_READY'] == 12
        assert summaries[0]['DEFERRED'] == 1
        assert summaries[0]['Ready Input %'] == 40
        assert summaries[1]['Candidate Inputs'] == summaries[1]['NOT_APPLICABLE'] == 18
        assert summaries[1]['Applicable Inputs'] == 0
        assert summaries[1]['DEFERRED'] == 0
        assert summaries[1]['Ready Input %'] is None
        for row_number, expected_target in ((header_row + 1, 2), (header_row + 2, 23)):
            link = overview.cell(row_number, 10).hyperlink
            assert link.location == f"'Profile_Readiness'!A{expected_target}"
            assert link.target is None
            assert workbook['Profile_Readiness'].cell(expected_target, 1).value == overview.cell(row_number, 1).value
        note = overview.cell(header_row + 3, 1).value
        assert 'READY inputs / applicable candidate inputs' in note
        assert 'not object coverage' in note
        assert 'not an approved standard' in note
        assert str(header_row + 3) in str(overview.print_area)
    finally:
        workbook.close()


def test_existing_overview_hotspots_still_resolve_by_scoped_parameter(report_path):
    workbook = load_workbook(report_path)
    try:
        overview = workbook['Overview']
        assert overview['A16'].value == 'Top standardization hotspots'
        standardization = workbook['Standardization']
        headers = {cell.value: cell.column for cell in standardization[1]}
        links = [row[0] for row in overview.iter_rows(min_row=18) if row[0].hyperlink]
        assert links
        for cell in links:
            assert cell.hyperlink.location.startswith("'Standardization'!A")
            row = int(cell.hyperlink.location.rsplit('A', 1)[1])
            parameter = standardization.cell(row, headers['Parameter']).value
            scope = standardization.cell(row, headers['Scope']).value
            assert cell.value == f'{parameter} [{scope}]'
            assert standardization.cell(row, headers['Parameter ID']).value
    finally:
        workbook.close()


def test_profile_matrix_preserves_literal_strings_and_escapes_control_characters(tmp_path, monkeypatch):
    from infoblox_inventory import report

    original = report.build_profile_readiness

    def with_literal_data(standardization, *args, **kwargs):
        rows = original(standardization, *args, **kwargs)
        rows[0]['Readiness Reason'] = '=1+1\x01literal'
        return rows

    monkeypatch.setattr(report, 'build_profile_readiness', with_literal_data)
    write_reports([_result()], tmp_path)
    path = tmp_path / 'current_state_inventory.xlsx'
    assert validate_xlsx(path)['zip_crc_ok']
    workbook = load_workbook(path)
    try:
        sheet = workbook['Profile_Readiness']
        column = next(cell.column for cell in sheet[1] if cell.value == 'Readiness Reason')
        cell = sheet.cell(2, column)
        assert cell.value == '=1+1\\u0001literal'
        assert cell.data_type == 's'
    finally:
        workbook.close()


def test_no_collection_results_keeps_unknown_evidence_visible(tmp_path):
    write_reports([], tmp_path)
    path = tmp_path / 'current_state_inventory.xlsx'
    assert validate_xlsx(path)['sheet_rows']['Profile_Readiness'] == 39
    workbook = load_workbook(path)
    try:
        rows = _rows(workbook['Profile_Readiness'])
        assert all(row['Readiness'] == 'NOT_READY' for row in rows)
        assert all(row['Population Objects'] is None for row in rows)
        assert all(row['Collection Status'] == 'UNKNOWN' for row in rows)
    finally:
        workbook.close()
