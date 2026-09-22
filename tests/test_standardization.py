from pathlib import Path

from openpyxl import load_workbook
import yaml

from infoblox_inventory.models import CollectionResult
from infoblox_inventory.report import write_reports
from infoblox_inventory.standardization import (
    build_standardization, exception_rows, grid_comparison_rows, load_decisions,
)
from infoblox_inventory.xlsx_validation import validate_xlsx


def _option(grid, obj_type, obj_ref, obj_name, value, *, configured_here, inherited,
            source_level, status="COMPLETE", option_number=51, parameter="dhcp-lease-time"):
    return {
        "grid": grid, "object_type": obj_type, "object_ref": obj_ref, "object_name": obj_name,
        "network_view": "default", "parent_network": "", "parameter": parameter,
        "option_number": option_number, "vendor_class": "DHCP", "effective_value": value,
        "configured_here": configured_here, "inherited": inherited, "multisource": False,
        "source_level": source_level, "source_object": obj_name if configured_here else "LAB",
        "source_ref": obj_ref if configured_here else "grid:dhcpproperties/example:LAB", "status": status,
    }


def _coverage(grid, object_type, query, status, count):
    return {"Grid": grid, "Area": "DHCP options and inheritance", "Object": object_type,
            "Query": query, "Field": "", "Collection Status": status,
            "Objects Found": count, "Notes": ""}


def _field_coverage(grid, object_type, field, status="NOT_EXPOSED_BY_WAPI"):
    return {"Grid": grid, "Area": "DHCP options and inheritance", "Object": object_type,
            "Query": "schema", "Field": field, "Collection Status": status,
            "Objects Found": "", "Notes": "Field unavailable"}


def _scalar(grid, obj_type, obj_ref, obj_name, parameter, value=None, *,
            configured_here=None, inherited=None, source_level="UNKNOWN", status="PARTIAL"):
    return {
        "grid": grid, "object_type": obj_type, "object_ref": obj_ref, "object_name": obj_name,
        "network_view": "default", "parent_network": "", "parameter": parameter,
        "effective_value": value, "configured_here": configured_here, "inherited": inherited,
        "multisource": False, "source_level": source_level, "source_object": "",
        "source_ref": "", "status": status,
    }


def _result(grid="LAB", networks=1, ranges=1, *, network_effective=None, range_effective=None,
            network_effective_status="COMPLETE", range_effective_status="COMPLETE"):
    network_effective = networks if network_effective is None else network_effective
    range_effective = ranges if range_effective is None else range_effective
    grid_ref = f"grid:dhcpproperties/example:{grid}"
    network_rows = [{"_ref": f"network/{grid}/{index}", "network": f"10.{index}.0.0/24"}
                    for index in range(networks)]
    range_rows = [{"_ref": f"range/{grid}/{index}", "start_addr": f"10.0.0.{index + 10}",
                   "end_addr": f"10.0.0.{index + 20}"} for index in range(ranges)]
    result = CollectionResult(
        grid=grid,
        records={
            "grid:dhcpproperties": [{"_ref": grid_ref}],
            "network": network_rows,
            "range": range_rows,
        },
        effective_records={
            "network": network_rows[:network_effective],
            "range": range_rows[:range_effective],
        },
        coverage=[
            _coverage(grid, "grid:dhcpproperties", "raw", "COMPLETE", 1),
            _coverage(grid, "network", "raw", "COMPLETE", networks),
            _coverage(grid, "network", "effective", network_effective_status, network_effective),
            _coverage(grid, "range", "raw", "COMPLETE", ranges),
            _coverage(grid, "range", "effective", range_effective_status, range_effective),
        ],
    )
    return result


def _standardization(decisions=None):
    result = _result()
    rows = [
        _option("LAB", "grid:dhcpproperties", "grid:dhcpproperties/example:LAB", "LAB", "28800",
                configured_here=True, inherited=False, source_level="Grid"),
        _option("LAB", "network", "network/LAB/0", "10.0.0.0/24", "28800",
                configured_here=False, inherited=True, source_level="Grid"),
        _option("LAB", "range", "range/LAB/0", "10.0.0.10-10.0.0.20", "3600",
                configured_here=True, inherited=False, source_level="Range"),
    ]
    return build_standardization([result], result.coverage, [], rows, decisions or {})


def _lease(rows, scope):
    found = [row for row in rows if row["Parameter ID"] == f"dhcp.lease_time.{scope.lower()}"]
    assert len(found) == 1
    return found[0]


def test_scope_specific_distributions_do_not_mix_grid_network_and_range():
    rows = _standardization()
    grid = _lease(rows, "Grid")
    network = _lease(rows, "Network")
    range_row = _lease(rows, "Range")

    assert grid["Observed Values"] == "28800 (1/1, 100.0%)"
    assert network["Observed Values"] == "28800 (1/1, 100.0%)"
    assert range_row["Observed Values"] == "3600 (1/1, 100.0%)"
    assert network["Common Value %"] == 100.0
    assert range_row["Common Value %"] == 100.0
    assert network["Consistency Classification"] == "CONSISTENT"
    assert range_row["Consistency Classification"] == "LOCAL_OVERRIDES"
    assert range_row["Local Override Count"] == 1
    assert range_row["Local Override %"] == 100.0
    assert grid["Override Eligible Objects"] == 0


def test_fixedaddress_option_cannot_contaminate_network_or_range_statistics():
    result = _result()
    rows = [
        _option("LAB", "network", "network/LAB/0", "10.0.0.0/24", "28800",
                configured_here=False, inherited=True, source_level="Grid"),
        _option("LAB", "range", "range/LAB/0", "10.0.0.10-10.0.0.20", "3600",
                configured_here=True, inherited=False, source_level="Range"),
        _option("LAB", "fixedaddress", "fixedaddress/LAB/1", "10.0.0.99", "86400",
                configured_here=True, inherited=False, source_level="Fixed Address"),
    ]
    analysis = build_standardization([result], result.coverage, [], rows, {})
    assert _lease(analysis, "Network")["Observed Values"] == "28800 (1/1, 100.0%)"
    assert _lease(analysis, "Range")["Observed Values"] == "3600 (1/1, 100.0%)"
    assert all("86400" not in _lease(analysis, scope)["Observed Values"] for scope in ("Grid", "Network", "Range"))


def test_partial_effective_collection_changes_coverage_and_uses_population_denominator():
    result = _result(networks=4, ranges=0, network_effective=3, network_effective_status="PARTIAL")
    result.errors.append({"grid": "LAB", "area": "DHCP options and inheritance",
                          "object_type": "network", "query": "effective", "error": "ReadTimeout"})
    rows = [
        _option("LAB", "network", "network/LAB/0", "n0", "28800",
                configured_here=False, inherited=True, source_level="Grid"),
        _option("LAB", "network", "network/LAB/1", "n1", "28800",
                configured_here=True, inherited=False, source_level="Network"),
    ]
    lease = _lease(build_standardization([result], result.coverage, [], rows, {}), "Network")
    assert lease["Population Objects"] == 4
    assert lease["Query Evidence Objects"] == 3
    assert lease["Query Coverage %"] == 75.0
    assert lease["Collection Status"] == "PARTIAL"
    assert lease["Confirmed Objects"] == 2
    assert lease["Confirmed Value %"] == 50.0
    assert lease["Explicit Not Configured"] == 0
    assert lease["Resolved Evidence %"] == 50.0
    assert lease["Unresolved Objects"] == 2
    assert lease["Unresolved %"] == 50.0
    assert lease["Common Value Count"] == 2
    assert lease["Common Value %"] == 50.0
    assert lease["Observed Values"] == "28800 (2/4, 50.0%)"
    assert lease["Local Override Count"] == 1
    assert lease["Override Eligible Objects"] == 4
    assert lease["Local Override %"] == 25.0
    assert lease["Inherited Count"] == 1
    assert lease["Inherited %"] == 25.0
    assert lease["Source Distribution"] == "Grid (1); Network (1)"
    assert lease["Coverage"] == "MEDIUM"
    assert lease["Evidence Status"] == "PARTIAL"
    assert lease["Standardization Candidate"] == "NEEDS_ANALYSIS"


def test_unrelated_missing_schema_field_does_not_degrade_parameter_collection():
    result = _result(networks=4, ranges=0, network_effective=4, network_effective_status="PARTIAL")
    result.coverage.append(_field_coverage("LAB", "network", "ddns_ttl"))
    rows = [
        _option("LAB", "network", f"network/LAB/{index}", f"n{index}", "28800",
                configured_here=False, inherited=True, source_level="Grid")
        for index in range(4)
    ]
    lease = _lease(build_standardization([result], result.coverage, [], rows, {}), "Network")
    assert lease["Required Field"] == "options"
    assert lease["Required Field Status"] == "AVAILABLE_OR_NOT_FLAGGED"
    assert lease["Collection Status"] == "COMPLETE"
    assert lease["Query Coverage %"] == 100.0
    assert lease["Confirmed Objects"] == 4
    assert lease["Confirmed Value %"] == 100.0
    assert lease["Unresolved Objects"] == 0
    assert lease["Coverage"] == "HIGH"
    assert lease["Evidence Status"] == "COMPLETE"


def test_required_field_unavailable_blocks_only_that_parameter():
    result = _result(networks=2, ranges=0, network_effective=2, network_effective_status="PARTIAL")
    result.coverage.append(_field_coverage("LAB", "network", "options"))
    rows = [
        _option("LAB", "network", "network/LAB/0", "n0", "28800",
                configured_here=False, inherited=True, source_level="Grid"),
    ]
    lease = _lease(build_standardization([result], result.coverage, [], rows, {}), "Network")
    assert lease["Required Field Status"] == "NOT_EXPOSED_BY_WAPI"
    assert lease["Coverage"] == "LOW"
    assert lease["Evidence Status"] == "NOT_EXPOSED_BY_WAPI"
    assert lease["Standardization Candidate"] == "BLOCKED_BY_DATA"


def test_full_collection_with_127_of_130_confirmed_separates_collection_from_evidence():
    result = _result(networks=0, ranges=130, range_effective=130, range_effective_status="PARTIAL")
    result.coverage.append(_field_coverage("LAB", "range", "ddns_ttl"))
    rows = [
        _option("LAB", "range", f"range/LAB/{index}", f"r{index}", "14400",
                configured_here=False, inherited=True, source_level="Grid")
        for index in range(127)
    ]
    lease = _lease(build_standardization([result], result.coverage, [], rows, {}), "Range")
    assert lease["Population Objects"] == 130
    assert lease["Query Evidence Objects"] == 130
    assert lease["Query Coverage %"] == 100.0
    assert lease["Collection Status"] == "COMPLETE"
    assert lease["Confirmed Objects"] == 127
    assert lease["Confirmed Value %"] == 97.7
    assert lease["Explicit Not Configured"] == 0
    assert lease["Resolved Evidence %"] == 97.7
    assert lease["Unresolved Objects"] == 3
    assert lease["Unresolved %"] == 2.3
    assert lease["Coverage"] == "MEDIUM"
    assert lease["Evidence Status"] == "PARTIAL"


def test_explicit_not_configured_is_resolved_separately_from_unresolved():
    result = _result(networks=2, ranges=0)
    rows = [
        _scalar("LAB", "network", "network/LAB/0", "n0", "bootserver", "10.0.0.10",
                configured_here=False, inherited=True, source_level="Grid", status="COMPLETE"),
        _scalar("LAB", "network", "network/LAB/1", "n1", "bootserver",
                configured_here=False, inherited=True, source_level="NOT_DEFINED", status="NOT_CONFIGURED"),
    ]
    analysis = build_standardization([result], result.coverage, rows, [], {})
    item = next(row for row in analysis if row["Parameter ID"] == "pxe.bootserver.network")
    assert item["Population Objects"] == 2
    assert item["Confirmed Objects"] == 1
    assert item["Explicit Not Configured"] == 1
    assert item["Resolved Evidence %"] == 100.0
    assert item["Unresolved Objects"] == 0
    assert item["Coverage"] == "HIGH"
    assert item["Evidence Status"] == "COMPLETE"


def test_member_effective_limitation_remains_unresolved():
    result = CollectionResult(
        "LAB",
        records={"member:dhcpproperties": [{"_ref": f"member:dhcpproperties/LAB/{i}"} for i in range(3)]},
        effective_records={"member:dhcpproperties": [{"_ref": f"member:dhcpproperties/LAB/{i}"} for i in range(3)]},
        coverage=[
            _coverage("LAB", "member:dhcpproperties", "raw", "COMPLETE", 3),
            _coverage("LAB", "member:dhcpproperties", "effective", "COMPLETE", 3),
        ],
    )
    rows = [
        _scalar("LAB", "member:dhcpproperties", f"member:dhcpproperties/LAB/{i}", f"m{i}", "nextserver")
        for i in range(3)
    ]
    analysis = build_standardization([result], result.coverage, rows, [], {})
    item = next(row for row in analysis if row["Parameter ID"] == "pxe.nextserver.member")
    assert item["Collection Status"] == "COMPLETE"
    assert item["Query Coverage %"] == 100.0
    assert item["Confirmed Objects"] == 0
    assert item["Unresolved Objects"] == 3
    assert item["Coverage"] == "MEDIUM"
    assert item["Evidence Status"] == "PARTIAL"
    assert item["Standardization Candidate"] == "BLOCKED_BY_DATA"


def test_conflicting_duplicate_normalized_rows_do_not_inflate_confirmation():
    result = _result(networks=1, ranges=0)
    rows = [
        _option("LAB", "network", "network/LAB/0", "n0", "28800",
                configured_here=False, inherited=True, source_level="Grid"),
        _option("LAB", "network", "network/LAB/0", "n0", "3600",
                configured_here=True, inherited=False, source_level="Network"),
    ]
    lease = _lease(build_standardization([result], result.coverage, [], rows, {}), "Network")
    assert lease["Population Objects"] == 1
    assert lease["Confirmed Objects"] == 0
    assert lease["Unresolved Objects"] == 1
    assert lease["Confirmed Value %"] == 0.0
    assert lease["Coverage"] == "MEDIUM"


def test_missing_collection_coverage_blocks_review_even_when_normalized_rows_exist():
    result = CollectionResult("LAB", records={"network": [{"_ref": "network/LAB/0"}]},
                              effective_records={"network": [{"_ref": "network/LAB/0"}]})
    rows = [_option("LAB", "network", "network/LAB/0", "n0", "28800",
                    configured_here=False, inherited=True, source_level="Grid")]
    lease = _lease(build_standardization([result], [], [], rows, {}), "Network")
    assert lease["Confirmed Objects"] == 1
    assert lease["Coverage"] == "MEDIUM"
    assert lease["Evidence Status"] == "PARTIAL"
    assert lease["Standardization Candidate"] == "NEEDS_ANALYSIS"


def test_approved_target_creates_deviation_only_for_the_selected_scope():
    decisions = {"dhcp.lease_time.range": {
        "status": "APPROVED", "approved_target": "28800", "exceptions_allowed": False,
        "proposed_target": None, "exception_rule": "", "owner": "PO", "decision_date": "2026-09-21",
        "notes": "", "comparison_mode": "exact", "approved_exceptions": [],
    }}
    rows = exception_rows(_standardization(decisions))
    assert [(row["Object Type"], row["Parameter ID"], row["Assessment State"]) for row in rows] == [
        ("range", "dhcp.lease_time.range", "DEVIATION")
    ]
    assert rows[0]["Approved Target"] == "28800"
    assert rows[0]["Owner"] == "PO"


def test_explicit_exception_selector_can_approve_a_scoped_deviation():
    ref = "range/LAB/0"
    decisions = {"dhcp.lease_time.range": {
        "status": "APPROVED", "approved_target": 28800, "exceptions_allowed": True,
        "proposed_target": None, "exception_rule": "Short-lived test range", "owner": "PO", "decision_date": "",
        "notes": "", "comparison_mode": "exact", "approved_exceptions": [{"object_ref": ref}],
    }}
    row, = exception_rows(_standardization(decisions))
    assert row["Assessment State"] == "APPROVED_EXCEPTION"
    assert row["Exception Rule"] == "Short-lived test range"


def test_unscoped_decision_id_is_rejected_with_migration_hint(tmp_path):
    path = tmp_path / "decisions.yaml"
    path.write_text("decisions:\n  dhcp.lease_time:\n    status: UNDER_REVIEW\n", encoding="utf-8")
    try:
        load_decisions(path)
    except ValueError as exc:
        text = str(exc)
        assert "Unscoped decision IDs" in text
        assert "dhcp.lease_time.network" in text
        assert "dhcp.lease_time.range" in text
    else:
        raise AssertionError("legacy unscoped decision id should fail")


def test_scoped_decision_loader_never_infers_approval(tmp_path):
    path = tmp_path / "decisions.yaml"
    path.write_text("decisions:\n  dhcp.lease_time.network:\n    status: UNDER_REVIEW\n    proposed_target: 28800\n",
                    encoding="utf-8")
    loaded = load_decisions(path)
    assert loaded["dhcp.lease_time.network"]["status"] == "UNDER_REVIEW"
    assert loaded["dhcp.lease_time.network"]["approved_target"] is None


def test_single_grid_comparison_is_explicitly_not_applicable():
    rows, headers = grid_comparison_rows(_standardization(), ["LAB"])
    assert {"LAB", "Scope", "Object Type"} <= set(headers)
    assert rows == [{
        "Category": "Comparison", "Parameter ID": "", "Parameter": "Cross-Grid comparison",
        "Scope": "", "Object Type": "", "LAB": "", "Distinct Values": "", "Common Observed Value": "",
        "Consistency": "NOT_APPLICABLE", "Local Overrides": "", "Coverage": "",
        "Notes": "Cross-Grid comparison requires at least two Grids.",
    }]


def test_multi_grid_comparison_keeps_scope_and_grid_denominators_separate():
    a = _result("GRID-A", networks=2, ranges=0)
    b = _result("GRID-B", networks=1, ranges=0)
    rows = [
        _option("GRID-A", "network", "network/GRID-A/0", "a0", "28800",
                configured_here=False, inherited=True, source_level="Grid"),
        _option("GRID-A", "network", "network/GRID-A/1", "a1", "28800",
                configured_here=False, inherited=True, source_level="Grid"),
        _option("GRID-B", "network", "network/GRID-B/0", "b0", "43200",
                configured_here=False, inherited=True, source_level="Grid"),
    ]
    coverage = a.coverage + b.coverage
    analysis = build_standardization([a, b], coverage, [], rows, {})
    lease = _lease(analysis, "Network")
    assert lease["Consistency Classification"] == "MULTIPLE_VALUES"
    assert lease["Common Observed Value"] == "28800"
    assert lease["Common Value Count"] == 2
    assert lease["Common Value %"] == 66.7
    comparison, headers = grid_comparison_rows(analysis, ["GRID-A", "GRID-B"])
    assert {"GRID-A", "GRID-B", "Scope"} <= set(headers)
    row = next(item for item in comparison if item["Parameter ID"] == "dhcp.lease_time.network")
    assert row["Scope"] == "Network"
    assert row["GRID-A"] == "28800 (2/2, 100.0%)"
    assert row["GRID-B"] == "43200 (1/1, 100.0%)"


def test_report_puts_scoped_decision_support_first_and_links_overview_by_parameter_id(tmp_path):
    grid_ref = "grid:dhcpproperties/example:LAB"
    network_ref = "network/example:10.0.0.0/24/default"
    range_ref = "range/example:10.0.0.10/10.0.0.20/default"
    result = CollectionResult(
        grid="LAB", grid_url="https://grid", wapi_version="2.13.7", collected_at="2026-09-21T06:00:00Z",
        records={
            "grid:dhcpproperties": [{"_ref": grid_ref, "options": [{"num": 51, "name": "dhcp-lease-time", "value": "28800", "use_option": True}]}],
            "network": [{"_ref": network_ref, "network": "10.0.0.0/24", "network_view": "default",
                         "options": [{"num": 51, "name": "dhcp-lease-time", "value": "43200", "use_option": False}]}],
            "range": [{"_ref": range_ref, "start_addr": "10.0.0.10", "end_addr": "10.0.0.20", "network_view": "default",
                       "options": [{"num": 51, "name": "dhcp-lease-time", "value": "3600", "use_option": True}]}],
        },
        effective_records={
            "network": [{"_ref": network_ref, "network": "10.0.0.0/24", "network_view": "default",
                         "options": [{"inherited": True, "source": grid_ref,
                                      "values": [{"num": 51, "name": "dhcp-lease-time", "value": "28800", "use_option": False}]}]}],
            "range": [{"_ref": range_ref, "start_addr": "10.0.0.10", "end_addr": "10.0.0.20", "network_view": "default",
                       "options": [{"inherited": False, "source": "",
                                    "values": [{"num": 51, "name": "dhcp-lease-time", "value": "3600", "use_option": True}]}]}],
        },
        schemas={"network": {"fields": [{"name": "options", "overridden_by": "use_options"}]},
                 "range": {"fields": [{"name": "options", "overridden_by": "use_options"}]}},
        coverage=[
            _coverage("LAB", "grid:dhcpproperties", "raw", "COMPLETE", 1),
            _coverage("LAB", "network", "raw", "COMPLETE", 1),
            _coverage("LAB", "network", "effective", "COMPLETE", 1),
            _coverage("LAB", "range", "raw", "COMPLETE", 1),
            _coverage("LAB", "range", "effective", "COMPLETE", 1),
        ],
    )
    decisions = tmp_path / "decisions.yaml"
    decisions.write_text(yaml.safe_dump({"decisions": {"dhcp.lease_time.range": {
        "status": "APPROVED", "approved_target": 28800, "exceptions_allowed": False,
        "owner": "Infoblox Product Owner",
    }}}), encoding="utf-8")
    output = tmp_path / "report"
    write_reports([result], output, decisions_path=decisions)
    validated = validate_xlsx(output / "current_state_inventory.xlsx")
    assert validated["zip_crc_ok"] is True
    workbook = load_workbook(output / "current_state_inventory.xlsx")
    try:
        assert workbook.sheetnames[:9] == [
            "Overview", "Profile_Readiness", "Profile_Populations", "Standardization",
            "Decisions", "Exceptions", "Grid_Comparison", "Coverage", "Manual_Review"
        ]
        standardization = workbook["Standardization"]
        headers = {cell.value: cell.column for cell in standardization[1]}
        parameter_id_col = headers["Parameter ID"]
        parameter_col = headers["Parameter"]
        scope_col = headers["Scope"]
        links = []
        for row_number in range(18, workbook["Overview"].max_row + 1):
            cell = workbook["Overview"].cell(row_number, 1)
            if cell.hyperlink is None:
                continue
            links.append(cell)
            target = cell.hyperlink.location or cell.hyperlink.target
            assert target is not None
            target = target.removeprefix("#")
            assert target.startswith("'Standardization'!A")
            target_row = int(target.rsplit("A", 1)[1])
            target_parameter = standardization.cell(target_row, parameter_col).value
            target_scope = standardization.cell(target_row, scope_col).value
            assert cell.value == f"{target_parameter} [{target_scope}]"
            target_id = standardization.cell(target_row, parameter_id_col).value
            assert target_id
        assert links

        decision_headers = [cell.value for cell in workbook["Decisions"][1]]
        decision_rows = [dict(zip(decision_headers, values)) for values in workbook["Decisions"].iter_rows(min_row=2, values_only=True)]
        lease = [row for row in decision_rows if row["Decision ID"] == "dhcp.lease_time.range"][0]
        assert lease["Scope"] == "Range"
        assert lease["Status"] == "APPROVED"
        assert lease["Approved Target"] == "28800"

        exception_headers = [cell.value for cell in workbook["Exceptions"][1]]
        exceptions = [dict(zip(exception_headers, values)) for values in workbook["Exceptions"].iter_rows(min_row=2, values_only=True)]
        assert [(row["Object Type"], row["Parameter ID"], row["Assessment State"]) for row in exceptions] == [
            ("range", "dhcp.lease_time.range", "DEVIATION")
        ]
    finally:
        workbook.close()
    template = yaml.safe_load((output / "standardization_decisions.template.yaml").read_text(encoding="utf-8"))
    assert template["version"] == 2
    assert template["decisions"]["dhcp.lease_time.range"]["status"] == "APPROVED"


def test_cli_parser_accepts_decision_overlay_for_offline_reporting():
    from infoblox_inventory.cli import build_parser

    args = build_parser().parse_args(["--offline", "raw", "--decisions", "decisions.yaml"])
    assert args.offline == ["raw"]
    assert args.decisions == "decisions.yaml"
