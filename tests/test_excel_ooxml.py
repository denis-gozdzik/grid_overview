"""Inspect serialized XLSX parts, including mutations Excel may otherwise repair."""
from copy import deepcopy
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.table import Table
import pytest
import requests

from infoblox_inventory.models import CollectionResult
from infoblox_inventory.report import _write_inventory_sheet, write_reports
from infoblox_inventory.xlsx_validation import MAIN, PACKAGE_REL, REL, XlsxValidationError, validate_xlsx


NS = {"s": MAIN, "r": REL}


def test_generated_report_is_valid_ooxml_across_collector_families(tmp_path, monkeypatch):
    def no_network(*args, **kwargs):
        raise AssertionError("Excel repair and generation must stay offline")

    monkeypatch.setattr(requests.Session, "request", no_network)
    result = CollectionResult(
        grid="SYNTHETIC", wapi_version="2.13.7", collected_at="2026-09-20T12:00:00Z",
        records={
            "networkview": [{"_ref": "networkview/synthetic:default", "name": "default"}],
            "network": [{"_ref": "network/synthetic:192.0.2.0/24/default", "network": "192.0.2.0/24",
                         "network_view": "default", "options": [{"num": 51, "value": "43200", "use_option": False}]}],
            "networktemplate": [{"_ref": "networktemplate/synthetic:access", "name": "access", "cidr": 24,
                                 "options": [{"num": 51, "value": "43200", "use_option": True}]}],
            "dhcpoptiondefinition": [{"_ref": "dhcpoptiondefinition/synthetic:51", "name": "dhcp-lease-time",
                                      "code": 51, "space": "DHCP", "type": "32-bit unsigned integer"}],
            "fixedaddress": [{"_ref": "fixedaddress/synthetic:192.0.2.10/default", "name": "Printer",
                              "ipv4addr": "192.0.2.10", "mac": "00:00:5e:00:53:01",
                              "network_view": "default", "logic_filter_rules": [{"filter": "Corporate", "type": "MAC"}]}],
            "filtermac": [{"_ref": "filtermac/synthetic:Corporate", "name": "Corporate", "disable": False}],
            "filteroption": [],
        },
        effective_records={"network": [{"_ref": "network/synthetic:192.0.2.0/24/default", "network": "192.0.2.0/24",
                                        "options": [{"inherited": True, "source": "grid:dhcpproperties/synthetic:LAB",
                                                     "values": [{"num": 51, "value": "28800", "use_option": False}]}]}]},
        coverage=[{"Grid": "SYNTHETIC", "Area": "DHCP", "Collection Status": "COMPLETE", "Objects Found": 1},
                  {"Grid": "SYNTHETIC", "Area": "Approval process", "Collection Status": "MANUAL_REVIEW_REQUIRED",
                   "Objects Found": 0}],
    )
    before = deepcopy(result)
    write_reports([result], tmp_path)
    assert result == before
    path = tmp_path / "current_state_inventory.xlsx"
    validated = validate_xlsx(path)
    assert validated["zip_crc_ok"] is True
    assert validated["table_count"] >= 10
    assert {"Grid_Summary", "Coverage", "Manual_Review", "Networks", "Network_Templates",
            "Option_Definitions", "Fixed_Reservations", "MAC_Filters"} <= validated["sheet_rows"].keys()
    for name in ("Network_Templates", "Option_Definitions", "Fixed_Reservations", "MAC_Filters"):
        assert validated["sheet_rows"][name] == 1
    assert validated["sheet_rows"]["Option_Filters"] == 0
    workbook = load_workbook(path)
    try:
        for sheet in workbook:
            assert sheet.freeze_panes == "A2"
            if sheet.tables:
                assert sheet.auto_filter.ref is None
                assert all(table.autoFilter is not None for table in sheet.tables.values())
            if sheet.max_row == 1:
                assert not sheet.tables
        data = list(workbook["DHCP_Options"].values)
        rows = [dict(zip(data[0], values)) for values in data[1:]]
        assert any(row["raw_value"] == "43200" and row["effective_value"] == "28800" for row in rows)
    finally:
        workbook.close()


def test_inventory_sheet_titles_are_excel_safe_and_collision_resistant(tmp_path):
    workbook = Workbook()
    workbook.remove(workbook.active)
    logical_first = "A" * 40
    logical_second = "A" * 39 + "B"

    first = _write_inventory_sheet(
        workbook, 1, logical_first, [{"Value": "first"}], ["Value"]
    )
    second = _write_inventory_sheet(
        workbook, 2, logical_second, [{"Value": "second"}], ["Value"]
    )
    sanitized = _write_inventory_sheet(
        workbook, 3, "Unsafe[]:*?/\\Name", [{"Value": "third"}], ["Value"]
    )

    assert first.title == "A" * 31
    assert second.title == "A" * 29 + "_2"
    assert first.title != second.title
    assert sanitized.title == "Unsafe_______Name"
    assert all(len(title) <= 31 for title in workbook.sheetnames)

    first["A2"].hyperlink = Hyperlink(
        ref="A2", location=f"'{second.title}'!A2"
    )
    path = tmp_path / "safe-sheet-titles.xlsx"
    workbook.save(path)
    workbook.close()

    validate_xlsx(path)
    reopened = load_workbook(path)
    try:
        assert reopened.sheetnames == ["A" * 31, "A" * 29 + "_2", "Unsafe_______Name"]
        link = reopened[reopened.sheetnames[0]]["A2"].hyperlink
        assert link.location == f"'{reopened.sheetnames[1]}'!A2"
        assert link.target is None
    finally:
        reopened.close()


def test_empty_report_contains_no_header_only_tables(tmp_path):
    write_reports([], tmp_path)
    path = tmp_path / "current_state_inventory.xlsx"
    validated = validate_xlsx(path)
    # Candidate readiness rows remain visible even when collection evidence is absent.
    assert validated["table_count"] == 1
    assert validated["sheet_rows"]["Profile_Readiness"] == 39
    assert not any(count for name, count in validated["sheet_rows"].items()
                   if name != "Profile_Readiness")
    workbook = load_workbook(path)
    try:
        assert len(workbook["Profile_Readiness"].tables) == 1
        assert workbook["Profile_Readiness"].max_row == 40
        assert all(not sheet.tables for sheet in workbook if sheet.title != "Profile_Readiness")
    finally:
        workbook.close()


def test_case_colliding_and_blank_source_keys_have_valid_headers_without_losing_values(tmp_path):
    result = CollectionResult(grid="SYNTHETIC", records={"networkview": [
        {"_ref": "networkview/synthetic:1", "name": "lowercase", "Name": "uppercase",
         "": "blank-key-data", " ": "space-key-data", "Column": "real-column-data"},
    ]})
    write_reports([result], tmp_path)
    path = tmp_path / "current_state_inventory.xlsx"
    validate_xlsx(path)
    workbook = load_workbook(path)
    try:
        sheet = workbook["Network_Views"]
        headers = [cell.value for cell in sheet[1]]
        assert all(isinstance(value, str) and value.strip() for value in headers)
        assert len({value.casefold() for value in headers}) == len(headers)
        assert {"lowercase", "uppercase", "blank-key-data", "space-key-data", "real-column-data"} <= {
            cell.value for cell in sheet[2]}
    finally:
        workbook.close()


@pytest.fixture
def simple_workbook(tmp_path):
    """Two valid tables, built independently of the report writer being tested."""
    path = tmp_path / "valid.xlsx"
    workbook = Workbook()
    workbook.remove(workbook.active)
    for index in (1, 2):
        sheet = workbook.create_sheet(f"Sheet{index}")
        sheet.append(["Name", "Value"])
        sheet.append([f"Object{index}", "Evidence"])
        sheet.add_table(Table(displayName=f"ValidTable{index}", ref="A1:B2"))
    workbook.save(path)
    validate_xlsx(path)
    return path


def _mutate(source: Path, target: Path, change):
    with ZipFile(source) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    change(parts)
    with ZipFile(target, "w", ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def _change_xml(parts, name, change):
    root = ET.fromstring(parts[name])
    change(root)
    parts[name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _defect(parts, defect):
    table1, table2 = "xl/tables/table1.xml", "xl/tables/table2.xml"
    sheet1, rels1 = "xl/worksheets/sheet1.xml", "xl/worksheets/_rels/sheet1.xml.rels"
    if defect == "overlapping_filter":
        # Partial overlap also violates ownership; the ranges need not be identical.
        _change_xml(parts, sheet1, lambda root: ET.SubElement(root, f"{{{MAIN}}}autoFilter", {"ref": "B1:C2"}))
    elif defect in {"duplicate_table_id", "duplicate_name", "duplicate_display_name"}:
        attribute, value = {"duplicate_table_id": ("id", "1"), "duplicate_name": ("name", "validtable1"),
                            "duplicate_display_name": ("displayName", "VALIDTABLE1")}[defect]
        _change_xml(parts, table2, lambda root: root.set(attribute, value))
    elif defect in {"invalid_ref", "header_only", "wrong_width"}:
        value = {"invalid_ref": "A0:B2", "header_only": "A1:B1", "wrong_width": "A1:C2"}[defect]
        _change_xml(parts, table1, lambda root: root.set("ref", value))
    elif defect == "wrong_count":
        _change_xml(parts, table1, lambda root: root.find("s:tableColumns", NS).set("count", "3"))
    elif defect in {"blank_header", "duplicate_header", "header_mismatch", "duplicate_column_id"}:
        def change(root):
            columns = root.findall("s:tableColumns/s:tableColumn", NS)
            attribute, value = {"blank_header": ("name", " "), "duplicate_header": ("name", "NAME"),
                                "header_mismatch": ("name", "Different"), "duplicate_column_id": ("id", "1")}[defect]
            columns[1].set(attribute, value)
        _change_xml(parts, table1, change)
    elif defect == "numeric_header":
        def change(root):
            cell = root.find("s:sheetData/s:row/s:c", NS)
            cell.set("t", "n")
            for child in list(cell):
                cell.remove(child)
            ET.SubElement(cell, f"{{{MAIN}}}v").text = "42"
        _change_xml(parts, sheet1, change)
    elif defect == "missing_data_rows":
        def change(root):
            data = root.find("s:sheetData", NS)
            for row in list(data)[1:]:
                data.remove(row)
        _change_xml(parts, sheet1, change)
    elif defect == "missing_target":
        _change_xml(parts, rels1, lambda root: root[0].set("Target", "/xl/tables/missing.xml"))
    elif defect == "external_table_relationship":
        _change_xml(parts, rels1, lambda root: root[0].set("TargetMode", "External"))
    elif defect == "duplicate_relationship_id":
        _change_xml(parts, rels1, lambda root: root.append(deepcopy(root[0])))
    elif defect == "orphan_relationship":
        _change_xml(parts, sheet1, lambda root: root.remove(root.find("s:tableParts", NS)))
    elif defect == "orphan_table":
        parts["xl/tables/orphan.xml"] = parts[table1]
    elif defect == "orphan_relationship_part":
        parts["xl/worksheets/_rels/absent.xml.rels"] = parts[rels1]
    elif defect == "table_parts_count":
        _change_xml(parts, sheet1, lambda root: root.find("s:tableParts", NS).set("count", "2"))
    elif defect == "wrong_relationship_type":
        _change_xml(parts, rels1, lambda root: root[0].set("Type", f"{REL}/worksheet"))
    elif defect == "missing_table_filter":
        _change_xml(parts, table1, lambda root: root.remove(root.find("s:autoFilter", NS)))
    elif defect == "malformed_xml":
        parts[table1] = b"<table"
    else:
        raise AssertionError(f"Unknown defect: {defect}")


@pytest.mark.parametrize(("defect", "message"), [
    ("overlapping_filter", "Worksheet AutoFilter overlaps"),
    ("duplicate_table_id", "Duplicate table ID"),
    ("duplicate_name", "Duplicate table name/displayName"),
    ("duplicate_display_name", "Duplicate table name/displayName"),
    ("invalid_ref", "invalid ref"),
    ("header_only", "Header-only or empty table"),
    ("wrong_width", "does not match table width"),
    ("wrong_count", "tableColumn count mismatch"),
    ("blank_header", "Blank table header"),
    ("duplicate_header", "Duplicate table header"),
    ("header_mismatch", "Table header/cell mismatch"),
    ("duplicate_column_id", "Duplicate tableColumn ID"),
    ("numeric_header", "header is not a string cell"),
    ("missing_data_rows", "no serialized data cells"),
    ("missing_target", "Missing or invalid relationship target"),
    ("external_table_relationship", "Invalid table relationship"),
    ("duplicate_relationship_id", "Duplicate relationship Id"),
    ("orphan_relationship", "Orphan table relationship"),
    ("orphan_table", "Orphan table part"),
    ("orphan_relationship_part", "Orphan relationship part"),
    ("table_parts_count", "tableParts count mismatch"),
    ("wrong_relationship_type", "Invalid table relationship"),
    ("missing_table_filter", "Table must own exactly one AutoFilter"),
    ("malformed_xml", "Malformed XML"),
])
def test_validator_rejects_serialized_excel_repair_defects(simple_workbook, tmp_path, defect, message):
    broken = tmp_path / "broken.xlsx"
    _mutate(simple_workbook, broken, lambda parts: _defect(parts, defect))
    with pytest.raises(XlsxValidationError, match=message):
        validate_xlsx(broken)
