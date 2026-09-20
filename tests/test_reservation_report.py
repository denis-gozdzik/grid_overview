"""Actual empty LAB captures and clearly synthetic populated Excel cases."""
from copy import deepcopy
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook
import requests

from infoblox_inventory.models import CollectionResult
from infoblox_inventory.report import write_reports
from infoblox_inventory.reservations import RESERVATION_SHEETS, normalize_reservations
from infoblox_inventory.storage import path_component


CAPTURE = Path(__file__).parent / "fixtures" / "reservations_lab_20260920"


def schemas():
    return {obj: json.loads((CAPTURE / "schemas" / f"{path_component(obj)}.json").read_text(encoding="utf-8"))
            for obj in RESERVATION_SHEETS}


def empty_live_result():
    records, effective = {}, {}
    for obj in RESERVATION_SHEETS:
        for mode, target in (("raw", records), ("effective", effective)):
            path = CAPTURE / "responses" / path_component(obj) / mode / "page-000001.json"
            if path.exists():
                target[obj] = json.loads(path.read_text(encoding="utf-8"))["result"]
    assert records["superhost"] == []
    # Dependency conclusion; no successful live child response is fabricated.
    records["superhostchild"] = []
    return CollectionResult(grid="LAB-GRID", records=records, effective_records=effective,
                            schemas=schemas(), coverage=[
        {"Grid": "LAB-GRID", "Area": "Reservations and filters", "Object": obj,
         "Query": "raw", "Collection Status": "EMPTY", "Objects Found": 0,
         "Notes": "No child GET: complete parent inventory empty" if obj == "superhostchild" else ""}
        for obj in records])


def rows(sheet):
    values = list(sheet.values)
    return [dict(zip(values[0], row)) for row in values[1:]]


def matrix(path):
    workbook = load_workbook(path)
    try:
        return {sheet.title: list(sheet.values) for sheet in workbook}
    finally:
        workbook.close()


def test_real_capture_hashes_empty_responses_and_parent_requirement():
    provenance = json.loads((CAPTURE / "provenance.json").read_text(encoding="utf-8"))
    for relative, expected in provenance["sha256"].items():
        assert hashlib.sha256((CAPTURE / relative).read_bytes()).hexdigest() == expected
    result = empty_live_result()
    assert len(result.records) == 12 and not any(result.records.values())
    assert result.effective_records == {"fixedaddress": []}
    assert normalize_reservations(result) == []
    error = json.loads((CAPTURE / "diagnostics/superhostchild-parent-required.json").read_text(encoding="utf-8"))
    assert error["status"] == 400
    assert error["body"]["text"] == "Search fields must contain 'parent'"
    assert not (CAPTURE / "responses/superhostchild/raw/page-000001.json").exists()


def test_empty_live_reports_have_meaningful_columns_and_repeat_offline(tmp_path, monkeypatch):
    def no_network(*args, **kwargs):
        raise AssertionError("Offline report attempted network access")
    monkeypatch.setattr(requests.Session, "request", no_network)
    result = empty_live_result()
    before = deepcopy(result)
    write_reports([result], tmp_path / "one")
    write_reports([result], tmp_path / "two")
    assert result == before
    assert matrix(tmp_path / "one/current_state_inventory.xlsx") == matrix(tmp_path / "two/current_state_inventory.xlsx")
    workbook = load_workbook(tmp_path / "one/current_state_inventory.xlsx")
    assert set(RESERVATION_SHEETS.values()) | {"Reservation_Options", "Reservation_Links"} <= set(workbook.sheetnames)
    for name in RESERVATION_SHEETS.values():
        sheet = workbook[name]
        assert sheet.max_row == 1
        assert {"grid", "object_ref", "normalization_status", "data_representation"} <= {cell.value for cell in sheet[1]}
    assert {cell.value for cell in workbook["Fixed_Reservations"][1]} >= {
        "ipv4addr", "mac", "dhcp_client_identifier", "match_client", "network_view",
        "enable_ddns", "ddns_hostname", "ddns_domainname", "extattrs", "disable", "ms_server"}
    assert {cell.value for cell in workbook["Option_Filters"][1]} >= {"expression", "option_list", "apply_as_class"}
    assert {cell.value for cell in workbook["MAC_Filter_Addresses"][1]} >= {"filter", "mac", "expiration_time"}
    assert all(sheet.freeze_panes == "A2" and sheet.auto_filter.ref for sheet in workbook)
    assert not any(row["Query"] == "normalization" for row in rows(workbook["Coverage"]))
    workbook.close()


def test_synthetic_reservations_keep_fixed_effective_pipeline_and_literal_filter_configuration(tmp_path):
    assert empty_live_result().records["fixedaddress"] == []
    ref = "fixedaddress/synthetic:192.0.2.10/default"
    raw = {"_ref": ref, "ipv4addr": "192.0.2.10", "name": "=1+1", "mac": "00:00:5e:00:53:01",
           "network_view": "default", "match_client": "RESERVED", "disable": False,
           "enable_ddns": False, "extattrs": {"Owner": {"value": ["opaque", "metadata"]}},
           "options": [{"num": 51, "name": "dhcp-lease-time", "value": "43200", "use_option": False}],
           "use_options": False, "logic_filter_rules": [{"filter": "Corporate", "type": "MAC"}],
           "use_logic_filter_rules": False, "new_field": {"unknown": [1, 2]}}
    result = CollectionResult(grid="SYNTHETIC", records={
        "fixedaddress": [raw],
        "filteroption": [{"name": "Allowed", "expression": "=1+1", "apply_as_class": False,
                          "option_list": [{"num": 60, "value": "opaque", "vendor_class": "custom"}]}],
        "macfilteraddress": [{"mac": "00:00:5e:00:53:01", "filter": "Corporate", "never_expires": True}],
    }, effective_records={"fixedaddress": [{"_ref": ref, "ipv4addr": "192.0.2.10", "network_view": "default",
        "options": [{"inherited": True, "source": "grid:dhcpproperties/synthetic:LAB", "values": [
            {"num": 51, "name": "dhcp-lease-time", "value": "28800", "use_option": False}]}]}]}, schemas=schemas())
    before = deepcopy(result)
    write_reports([result], tmp_path)
    assert result == before
    workbook = load_workbook(tmp_path / "current_state_inventory.xlsx")
    fixed, = rows(workbook["Fixed_Reservations"])
    assert fixed["data_representation"] == "RAW_CONFIGURATION" and fixed["disable"] is False
    assert fixed["match_client"] == "RESERVED"
    assert json.loads(fixed["extra_fields"])["new_field"] == {"unknown": [1, 2]}
    options = rows(workbook["Reservation_Options"])
    stored = next(row for row in options if row["object_type"] == "fixedaddress")
    assert stored["stored_value"] == "43200" and stored["use_options"] is False
    lease, = [row for row in rows(workbook["DHCP_Options"]) if row["option_number"] == 51]
    assert lease["effective_value"] == "28800" and lease["raw_value"] == "43200"
    assert lease["status"] == "COMPLETE"
    assert all(row["object_type"] == "fixedaddress" for row in rows(workbook["DHCP_Options"]))
    option_filter, = rows(workbook["Option_Filters"])
    assert option_filter["name"] == "Allowed" and option_filter["expression"] == "=1+1"
    assert "policy" not in option_filter and "allow" not in option_filter
    links = rows(workbook["Reservation_Links"])
    assert {row["target"] for row in links} == {"Corporate"}
    assert any(row["use_logic_filter_rules"] is False for row in links)
    assert any(row["Query"] == "policy_interpretation" and row["Collection Status"] == "MANUAL_REVIEW_REQUIRED"
               for row in rows(workbook["Coverage"]))
    assert not any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row)
    workbook.close()


def test_synthetic_multigrid_host_flags_and_source_indexes_are_visible(tmp_path):
    assert empty_live_result().records["record:host"] == []
    records = {"record:host": [{"name": "same.example", "ipv4addrs": [17,
        {"_ref": "record:host_ipv4addr/synthetic:1", "configure_for_dhcp": False,
         "options": [False, {"num": 51, "value": "3600"}]}]}],
               "record:host_ipv4addr": [{"ipv4addr": "192.0.2.10", "configure_for_dhcp": False}]}
    write_reports([CollectionResult(grid=grid, records=deepcopy(records)) for grid in ("A", "B")], tmp_path)
    workbook = load_workbook(tmp_path / "current_state_inventory.xlsx")
    hosts = rows(workbook["Host_DHCP_Context"])
    assert {row["grid"] for row in hosts} == {"A", "B"}
    assert all(row["normalization_status"] == "PARTIAL" for row in hosts)
    assert all(row["dhcp_configuration_status"] == "NOT_CONFIGURED" for row in rows(workbook["Host_IPv4_DHCP"]))
    options = rows(workbook["Reservation_Options"])
    assert all(row["option_field"] == "ipv4addrs[1].options" and row["option_index"] == 1 for row in options)
    assert workbook["DHCP_Options"].max_row == 1
    workbook.close()
