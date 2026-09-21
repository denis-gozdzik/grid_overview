"""Synthetic end-to-end cumulative collection and immutable offline composition."""
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
from urllib.parse import parse_qs, urlsplit

from openpyxl import load_workbook
import pytest
import responses

from infoblox_inventory import cli
from infoblox_inventory.dataset import combine_collections
from infoblox_inventory.models import CollectionResult
from infoblox_inventory.report import write_reports
from infoblox_inventory.storage import RawStore, load_raw
from infoblox_inventory.xlsx_validation import validate_xlsx


ORIGIN = "https://synthetic-grid.example"
BASE = ORIGIN + "/wapi/v2.13.7/"
CORE_TIME = "2026-09-20T12:00:00Z"
CONFIG_TIME = "2026-09-20T13:00:00Z"


def _family_records():
    return {
        "network": [{"_ref": "network/synthetic:192.0.2.0/24/default", "network": "192.0.2.0/24",
                     "network_view": "default", "options": [{"num": 51, "value": "43200", "use_option": False}],
                     "use_options": False}],
        "networktemplate": [{"_ref": "networktemplate/synthetic:access", "name": "access", "cidr": 24}],
        "dhcpoptiondefinition": [{"_ref": "dhcpoptiondefinition/synthetic:51", "name": "dhcp-lease-time",
                                  "code": 51, "space": "DHCP", "type": "32-bit unsigned integer"}],
        "fixedaddress": [{"_ref": "fixedaddress/synthetic:192.0.2.10/default", "name": "Printer",
                          "ipv4addr": "192.0.2.10", "mac": "00:00:5e:00:53:01", "network_view": "default"}],
        "filtermac": [{"_ref": "filtermac/synthetic:Corporate", "name": "Corporate", "disable": False}],
    }


def _schema(rows):
    return {"fields": [{"name": name, "supports": "r"} for name in sorted({key for row in rows for key in row})
                       if name != "_ref"]}


def _coverage(obj, area="Test inventory", status="COMPLETE"):
    return {"Grid": "SYNTHETIC", "Area": area, "Object": obj, "Query": "raw", "Field": "",
            "Collection Status": status, "Objects Found": 1, "Notes": ""}


def _results():
    records = _family_records()
    schemas = {obj: _schema(rows) for obj, rows in records.items()}
    core = CollectionResult(grid="SYNTHETIC", grid_url=ORIGIN, wapi_version="2.13.7", collected_at=CORE_TIME,
                            records={"network": records.pop("network")}, schemas={"network": schemas.pop("network")},
                            coverage=[_coverage("network", "Networks")])
    configuration = CollectionResult(grid="SYNTHETIC", grid_url=ORIGIN, wapi_version="2.13.7", collected_at=CONFIG_TIME,
                                     records=records, schemas=schemas,
                                     coverage=[_coverage(obj) for obj in records])
    return core, configuration


def _save_result(root, result):
    store = RawStore(root, result)
    for obj, schema in result.schemas.items():
        store.save_schema(obj, schema)
    for mode, source in (("raw", result.records), ("effective", result.effective_records)):
        for obj, records in source.items():
            query = store.start_query(obj, mode, {"_paging": "1", "_return_as_object": "1"})
            store.save_page(query, 1, {"result": records, "synthetic_evidence": True})
            store.finish_query(query, "COMPLETE" if records else "EMPTY")
    store.finish_collection()


def _hashes(root):
    return {str(path.relative_to(root)): hashlib.sha256(path.read_bytes()).hexdigest()
            for path in root.rglob("*") if path.is_file()}


def _sheet_rows(sheet):
    values = list(sheet.values)
    return [dict(zip(values[0], row)) for row in values[1:]]


@responses.activate
def test_cli_all_real_collection_keeps_multiple_families_and_pages_in_one_report(tmp_path, monkeypatch):
    records = _family_records()
    second_network = {**deepcopy(records["network"][0]), "_ref": "network/synthetic:198.51.100.0/24/default",
                      "network": "198.51.100.0/24"}
    schemas = {obj: _schema(rows) for obj, rows in records.items()}
    for field in schemas["network"]["fields"]:
        if field["name"] == "options":
            field["overridden_by"] = "use_options"
    requested = []

    def serve(request):
        parsed = urlsplit(request.url)
        obj = parsed.path.removeprefix("/wapi/v2.13.7/")
        query = parse_qs(parsed.query)
        requested.append((request.method, obj, query))
        if not obj:
            payload = {"supported_objects": list(records), "supported_versions": ["2.13.7"]}
        elif "_schema" in query:
            payload = schemas[obj]
        else:
            rows = deepcopy(records[obj])
            continuation = None
            if obj == "network":
                mode = "effective" if query.get("_inheritance") == ["True"] else "raw"
                if "_page_id" in query:
                    mode = query["_page_id"][0].split(":")[0]
                    rows = [deepcopy(second_network)]
                else:
                    continuation = f"{mode}:second"
                if mode == "effective":
                    for row in rows:
                        row["options"] = [{"inherited": True, "source": "grid:dhcpproperties/synthetic:SYNTHETIC",
                                           "values": [{"num": 51, "value": "28800", "use_option": False}]}]
            payload = {"result": rows}
            if continuation:
                payload["next_page_id"] = continuation
        return 200, {"Content-Type": "application/json"}, json.dumps(payload)

    responses.add_callback(responses.GET, re.compile(re.escape(BASE) + ".*"), callback=serve)
    config = tmp_path / "grids.yaml"
    config.write_text("grids:\n  - name: SYNTHETIC\n    url: " + ORIGIN +
                      "\n    wapi_version: 2.13.7\n    username: readonly\n", encoding="utf-8")
    monkeypatch.setenv("INFOBLOX_PASSWORD", "synthetic-test-password")
    output = tmp_path / "all"
    assert cli.main(["--config", str(config), "--all", "--output-dir", str(output)]) == 0
    collected, = load_raw(output)
    assert set(collected.records) == set(records)
    assert len(collected.records["network"]) == len(collected.effective_records["network"]) == 2
    assert all(method == "GET" for method, _, _ in requested)
    assert {obj for _, obj, query in requested if obj and "_schema" not in query} == set(records)
    assert any("_page_id" in query for _, _, query in requested)
    assert all("/wapi/v2.13.7/" in call.request.url for call in responses.calls)
    validation = validate_xlsx(output / "current_state_inventory.xlsx")
    assert validation["sheet_rows"]["Networks"] == 2
    for name in ("Network_Templates", "Option_Definitions", "Fixed_Reservations", "MAC_Filters"):
        assert validation["sheet_rows"][name] == 1
    workbook = load_workbook(output / "current_state_inventory.xlsx")
    try:
        assert {row["Object"] for row in _sheet_rows(workbook["Grid_Summary"])} == set(records)
        assert any(row["raw_value"] == "43200" and row["effective_value"] == "28800"
                   for row in _sheet_rows(workbook["DHCP_Options"]))
    finally:
        workbook.close()


def test_repeated_offline_archives_merge_same_grid_without_network_or_raw_mutation(tmp_path, monkeypatch):
    core, configuration = _results()
    first, second = tmp_path / "core", tmp_path / "configuration"
    _save_result(first, core)
    _save_result(second, configuration)
    before = _hashes(first), _hashes(second)
    monkeypatch.setattr(cli, "InfobloxClient", lambda *args, **kwargs: pytest.fail("Offline replay reached live client"))
    output = tmp_path / "combined"
    assert cli.main(["--offline", str(first), "--offline", str(second), "--output-dir", str(output)]) == 0
    assert (_hashes(first), _hashes(second)) == before
    assert not (output / "raw").exists()
    validation = validate_xlsx(output / "current_state_inventory.xlsx")
    for name in ("Networks", "Network_Templates", "Option_Definitions", "Fixed_Reservations", "MAC_Filters"):
        assert validation["sheet_rows"][name] == 1
    workbook = load_workbook(output / "current_state_inventory.xlsx")
    try:
        summary = _sheet_rows(workbook["Grid_Summary"])
        assert len(summary) == 5
        assert {row["Grid"] for row in summary} == {"SYNTHETIC"}
    finally:
        workbook.close()


def test_merge_disjoint_families_is_deterministic_deep_copied_and_preserves_source_times():
    core, configuration = _results()
    originals = deepcopy((core, configuration))
    merged, = combine_collections([configuration, core])
    reversed_result, = combine_collections([core, configuration])
    assert merged == reversed_result
    assert set(merged.records) == set(_family_records())
    assert (core, configuration) == originals
    assert {row["Collected At"] for row in merged.collection_sources} == {CORE_TIME, CONFIG_TIME}
    assert {row["Collected At"] for row in merged.collection_sources if row["Object"] == "network"} == {CORE_TIME}
    assert {row["Collected At"] for row in merged.collection_sources if row["Object"] == "filtermac"} == {CONFIG_TIME}
    merged.records["network"][0]["options"][0]["value"] = "changed"
    merged.schemas["network"]["fields"][0]["name"] = "changed"
    assert (core, configuration) == originals


def test_identical_overlap_is_deduplicated_and_merge_can_be_applied_again():
    core, configuration = _results()
    duplicate = deepcopy(core)
    merged, = combine_collections([core, duplicate, configuration])
    assert len(merged.records["network"]) == 1
    assert merged.coverage.count(core.coverage[0]) == 1
    repeated, = combine_collections([merged])
    assert repeated == merged


def test_archive_identity_preserves_distinct_families_at_one_time_and_deduplicates_replay():
    core, configuration = _results()
    configuration.collected_at = core.collected_at
    merged, = combine_collections([core, deepcopy(core), configuration])
    identities = {row["Archive ID"] for row in merged.collection_sources}
    assert len(identities) == 2
    assert {row["Archive ID"] for row in merged.collection_sources if row["Object"] == "network"}.isdisjoint(
        {row["Archive ID"] for row in merged.collection_sources if row["Object"] == "filtermac"})
    assert combine_collections([merged])[0] == merged
    assert combine_collections([configuration, core])[0] == merged


def test_archive_identity_normalizes_equivalent_timezone_metadata_but_preserves_original_text():
    core, _ = _results()
    same = deepcopy(core)
    same.collected_at = "2026-09-20T14:00:00+02:00"
    merged, = combine_collections([core, same])
    assert len({row["Archive ID"] for row in merged.collection_sources}) == 1
    assert {row["Collected At"] for row in merged.collection_sources} == {CORE_TIME, same.collected_at}


def test_recombined_legacy_source_rows_are_not_assigned_an_invented_archive_identity():
    core, _ = _results()
    core.collection_sources = [{"Grid": core.grid, "Object": "network", "Collected At": CORE_TIME}]
    merged, = combine_collections([core])
    assert merged.collection_sources == core.collection_sources
    assert "Archive ID" not in merged.collection_sources[0]


@pytest.mark.parametrize("field", ["grid_url", "wapi_version", "records", "effective_records", "schemas"])
def test_conflicting_same_grid_snapshots_are_rejected_without_mutating_input(field):
    core, _ = _results()
    core.effective_records = {"network": [{"_ref": core.records["network"][0]["_ref"], "options": []}]}
    second = deepcopy(core)
    if field == "grid_url":
        second.grid_url = "https://different-origin.example"
    elif field == "wapi_version":
        second.wapi_version = "2.12"
    elif field == "records":
        second.records["network"][0]["comment"] = "Conflicting snapshot"
    elif field == "effective_records":
        second.effective_records["network"][0]["options"] = [{"num": 51, "value": "999"}]
    else:
        second.schemas["network"]["fields"].append({"name": "conflicting_field", "supports": "r"})
    before = deepcopy((core, second))
    with pytest.raises(ValueError):
        combine_collections([core, second])
    assert (core, second) == before


def test_distinct_grids_remain_separate():
    core, _ = _results()
    second = deepcopy(core)
    second.grid, second.grid_url = "OTHER", "https://other-grid.example"
    combined = combine_collections([core, second])
    assert {result.grid for result in combined} == {"SYNTHETIC", "OTHER"}
    assert all(len(result.records["network"]) == 1 for result in combined)


def test_raw_and_effective_from_different_snapshots_cannot_be_joined():
    core, _ = _results()
    later = CollectionResult(grid=core.grid, grid_url=core.grid_url, wapi_version=core.wapi_version,
                             collected_at=CONFIG_TIME, schemas=deepcopy(core.schemas),
                             effective_records={"network": [{"_ref": core.records["network"][0]["_ref"],
                                                              "options": []}]})
    before = deepcopy((core, later))
    with pytest.raises(ValueError):
        combine_collections([core, later])
    assert (core, later) == before


def test_only_superseded_increment_placeholder_is_removed():
    core, configuration = _results()
    placeholder = {"Grid": "SYNTHETIC", "Area": "Existing network templates", "Object": "", "Query": "",
                   "Field": "", "Collection Status": "PARTIAL", "Objects Found": "",
                   "Notes": "Assessment area is not fully covered by this increment"}
    unrelated = {**placeholder, "Area": "DDNS settings"}
    real_partial = {**placeholder, "Object": "networktemplate", "Notes": "Field unavailable in schema"}
    manual = {**placeholder, "Area": "Approval process", "Collection Status": "MANUAL_REVIEW_REQUIRED",
              "Notes": "Confirm with Grid owners"}
    core.coverage.extend([placeholder, unrelated, real_partial, manual])
    configuration.coverage.append(_coverage("networktemplate", "Existing network templates"))
    merged, = combine_collections([core, configuration])
    assert placeholder not in merged.coverage
    assert unrelated in merged.coverage
    assert real_partial in merged.coverage
    assert manual in merged.coverage


def test_offline_conflict_fails_without_workbook_and_preserves_archives(tmp_path, monkeypatch):
    core, _ = _results()
    conflict = deepcopy(core)
    conflict.records["network"][0]["network"] = "198.51.100.0/24"
    first, second = tmp_path / "first", tmp_path / "second"
    _save_result(first, core)
    _save_result(second, conflict)
    before = _hashes(first), _hashes(second)
    monkeypatch.setattr(cli, "InfobloxClient", lambda *args, **kwargs: pytest.fail("Offline conflict reached live client"))
    output = tmp_path / "conflict"
    assert cli.main(["--offline", str(first), "--offline", str(second), "--output-dir", str(output)]) == 1
    assert not (output / "current_state_inventory.xlsx").exists()
    assert (_hashes(first), _hashes(second)) == before


def test_report_itself_combines_results_before_normalization(tmp_path):
    core, configuration = _results()
    write_reports([core, deepcopy(core), configuration], tmp_path)
    validation = validate_xlsx(tmp_path / "current_state_inventory.xlsx")
    assert validation["sheet_rows"]["Networks"] == 1
    assert validation["sheet_rows"]["Network_Templates"] == 1
    workbook = load_workbook(tmp_path / "current_state_inventory.xlsx")
    try:
        assert len(_sheet_rows(workbook["Grid_Summary"])) == 5
    finally:
        workbook.close()


def test_collection_sources_preserve_failed_and_unexposed_object_timestamps(tmp_path):
    core, _ = _results()
    core.coverage.extend([
        {**_coverage("range", "Ranges", "NOT_EXPOSED_BY_WAPI"), "Query": "schema",
         "Objects Found": "", "Notes": "Object absent from root schema"},
        {**_coverage("filtermac", "DHCP filters and MAC filters", "ERROR"), "Query": "schema",
         "Objects Found": "", "Notes": "GET failed: HTTP 403"},
    ])
    core.errors.append({"grid": core.grid, "area": "DHCP filters and MAC filters",
                        "object_type": "filtermac", "error": "GET failed: HTTP 403"})
    before = deepcopy(core)
    merged, = combine_collections([core])
    assert core == before
    for object_type in ("range", "filtermac"):
        source, = [row for row in merged.collection_sources if row["Object"] == object_type]
        assert source["Collected At"] == CORE_TIME
        assert source["Raw Count"] is None and source["Effective Count"] is None
        assert object_type not in merged.records and object_type not in merged.effective_records
    write_reports([core], tmp_path)
    workbook = load_workbook(tmp_path / "current_state_inventory.xlsx")
    try:
        sources = _sheet_rows(workbook["Collection_Sources"])
        for object_type in ("range", "filtermac"):
            row, = [row for row in sources if row["Object"] == object_type]
            assert row["Collected At"] == CORE_TIME
            assert row["Raw Count"] is None and row["Effective Count"] is None
    finally:
        workbook.close()
