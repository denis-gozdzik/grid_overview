"""Live empty/predefined inventories and explicitly synthetic populated report cases."""
from copy import deepcopy
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook
import requests

from infoblox_inventory.models import CollectionResult
from infoblox_inventory.report import write_reports
from infoblox_inventory.topology import TOPOLOGY_SHEETS, normalize_topology


CAPTURE = Path(__file__).parent / "fixtures" / "topology_lab_20260920"


def live_result():
    provenance = json.loads((CAPTURE / "provenance.json").read_text(encoding="utf-8"))
    return CollectionResult(
        grid="LAB-GRID", wapi_version="2.13.7", collected_at="2026-09-20T06:32:29Z",
        records={obj: json.loads((CAPTURE / "responses" / f"{obj}.json").read_text(encoding="utf-8"))["result"]
                 for obj in TOPOLOGY_SHEETS},
        schemas={obj: json.loads((CAPTURE / "schemas" / f"{obj}.json").read_text(encoding="utf-8"))
                 for obj in TOPOLOGY_SHEETS},
        coverage=[{"Grid": "LAB-GRID", "Area": obj, "Object": obj, "Query": "raw",
                   "Collection Status": "COMPLETE" if count else "EMPTY", "Objects Found": count}
                  for obj, count in provenance["live_counts"].items()],
    )


def rows(sheet):
    values = list(sheet.values)
    return [dict(zip(values[0], row)) for row in values[1:]]


def matrix(path):
    workbook = load_workbook(path)
    try:
        return {sheet.title: list(sheet.values) for sheet in workbook}
    finally:
        workbook.close()


def test_actual_topology_capture_hashes_counts_and_typed_definitions():
    provenance = json.loads((CAPTURE / "provenance.json").read_text(encoding="utf-8"))
    assert provenance["kind"] == "actual_live_lab_topology_capture"
    for relative, expected in provenance["sha256"].items():
        assert hashlib.sha256((CAPTURE / relative).read_bytes()).hexdigest() == expected
    result = live_result()
    assert {obj: len(items) for obj, items in result.records.items()} == {
        "dhcpfailover": 0, "networktemplate": 0, "rangetemplate": 0,
        "fixedaddresstemplate": 0, "dhcpoptionspace": 1, "dhcpoptiondefinition": 95,
    }
    records = normalize_topology(result)
    assert len(records) == 96 and all(record.status == "COMPLETE" for record in records)
    space = next(record for record in records if record.object_type == "dhcpoptionspace")
    assert space.name == "DHCP" and space.space_type == "PREDEFINED_DHCP"
    assert space.comment is None
    definitions = [record for record in records if record.object_type == "dhcpoptiondefinition"]
    assert set(space.option_definitions) == {record.name for record in definitions}
    assert all(record.space == "DHCP" for record in definitions)
    lease = next(record for record in definitions if record.code == 51)
    assert lease.name == "dhcp-lease-time" and lease.type == "32-bit unsigned integer"


def test_live_capture_reports_preserve_empty_headers_coverage_and_repeat_offline(tmp_path, monkeypatch):
    def no_network(*args, **kwargs):
        raise AssertionError("Report generation must remain offline")
    monkeypatch.setattr(requests.Session, "request", no_network)
    result = live_result()
    before = deepcopy(result)
    write_reports([result], tmp_path / "first")
    write_reports([result], tmp_path / "second")
    assert result == before
    assert matrix(tmp_path / "first/current_state_inventory.xlsx") == matrix(tmp_path / "second/current_state_inventory.xlsx")
    workbook = load_workbook(tmp_path / "first/current_state_inventory.xlsx")
    assert set(TOPOLOGY_SHEETS.values()) | {"Template_Options"} <= set(workbook.sheetnames)
    for obj in ("dhcpfailover", "networktemplate", "rangetemplate", "fixedaddresstemplate"):
        sheet = workbook[TOPOLOGY_SHEETS[obj]]
        assert sheet.max_row == 1
        assert {"name", "object_ref", "normalization_status", "data_representation"} <= {cell.value for cell in sheet[1]}
    assert len(rows(workbook["Option_Definitions"])) == 95
    assert len(rows(workbook["Option_Spaces"])) == 1
    assert workbook["Template_Options"].max_row == 1
    coverage = rows(workbook["Coverage"])
    assert sum(row.get("Query") == "raw" and row["Collection Status"] == "EMPTY" for row in coverage) == 4
    assert sum(row.get("Query") == "normalization" and row["Collection Status"] == "COMPLETE" for row in coverage) == 2
    assert all(sheet.freeze_panes == "A2" and sheet.auto_filter.ref is None for sheet in workbook)
    assert all(table.autoFilter.ref == table.ref for sheet in workbook for table in sheet.tables.values())
    workbook.close()


def test_synthetic_multigrid_templates_are_stored_rows_not_effective_claims(tmp_path):
    results = []
    for grid, stored in (("A", "43200"), ("B", "86400")):
        result = CollectionResult(grid=grid, records={
            "networktemplate": [{"_ref": "networktemplate/synthetic:shared", "name": "=1+1",
                                 "bootfile": "=boot.efi", "use_bootfile": False,
                                 "options": [{"num": 51, "value": stored, "use_option": True,
                                              "future_option": ["opaque"]}],
                                 "use_options": False,
                                 "enable_ddns": True, "use_enable_ddns": False,
                                 "ddns_domainname": "corp.example", "use_ddns_domainname": True,
                                 "ddns_generate_hostname": False, "use_ddns_generate_hostname": True,
                                 "ddns_ttl": 3600, "use_ddns_ttl": True,
                                 "ddns_update_fixed_addresses": True,
                                 "use_ddns_update_fixed_addresses": False,
                                 "ddns_use_option81": True, "use_ddns_use_option81": True,
                                 "update_dns_on_lease_renewal": True,
                                 "use_update_dns_on_lease_renewal": True,
                                 "future_field": {"unknown": True}}],
            "rangetemplate": [{"_ref": "rangetemplate/synthetic:clients", "name": "clients",
                               "offset": 10, "number_of_addresses": 50,
                               "server_association_type": "MEMBER",
                               "enable_ddns": True, "use_enable_ddns": True,
                               "ddns_domainname": "ranges.example", "use_ddns_domainname": False,
                               "ddns_generate_hostname": True, "use_ddns_generate_hostname": True,
                               "update_dns_on_lease_renewal": False,
                               "use_update_dns_on_lease_renewal": True}],
            "member": [{"_ref": "member/synthetic", "name": "member"}],
            "dhcpfailover": [{"name": "pair", "primary": "primary", "secondary": "secondary"}],
            "fixedaddresstemplate": [{"name": "fixed-template", "number_of_addresses": "invalid",
                                      "enable_ddns": True, "use_enable_ddns": True,
                                      "ddns_domainname": "fixed.example", "use_ddns_domainname": True,
                                      "ddns_hostname": "host.example"}],
        }, schemas={"networktemplate": live_result().schemas["networktemplate"]},
            coverage=[{"Grid": grid, "Area": "Existing network templates", "Object": "networktemplate",
                       "Query": "raw", "Collection Status": "PARTIAL", "Objects Found": 1,
                       "Notes": "Synthetic second-page failure"}])
        results.append(result)
    before = deepcopy(results)
    write_reports(results, tmp_path)
    assert results == before
    workbook = load_workbook(tmp_path / "current_state_inventory.xlsx")
    templates = rows(workbook["Network_Templates"])
    assert {row["grid"] for row in templates} == {"A", "B"}
    assert all(row["name"] == "=1+1" and row["use_bootfile"] is False for row in templates)
    assert all(json.loads(row["extra_fields"])["future_field"] == {"unknown": True} for row in templates)
    assert all(row["enable_ddns"] is True and row["use_enable_ddns"] is False for row in templates)
    assert all(row["ddns_domainname"] == "corp.example" and row["use_ddns_domainname"] is True
               for row in templates)
    assert all(row["ddns_ttl"] == 3600 and row["use_ddns_ttl"] is True for row in templates)
    assert all(row["ddns_use_option81"] is True and row["use_ddns_use_option81"] is True
               for row in templates)
    range_templates = rows(workbook["Range_Templates"])
    assert all(row["enable_ddns"] is True and row["use_enable_ddns"] is True
               for row in range_templates)
    assert all(row["ddns_domainname"] == "ranges.example" and row["use_ddns_domainname"] is False
               for row in range_templates)
    fixed_templates = rows(workbook["Fixed_Address_Templates"])
    assert all(row["enable_ddns"] is True and row["ddns_hostname"] == "host.example"
               for row in fixed_templates)
    options = rows(workbook["Template_Options"])
    assert {(row["grid"], row["stored_value"]) for row in options} == {("A", "43200"), ("B", "86400")}
    assert all(row["data_representation"] == "RAW_CONFIGURATION" and row["use_options"] is False for row in options)
    assert all("effective_value" not in row for row in options)
    assert {"Data_Index", "Template_Archetypes", "Template_Grid_Map",
            "Template_Variants", "Template_Instances", "Template_Review",
            "Template_Bundle_Models", "Template_Bundles",
            "Template_Models", "Template_Grid_Matrix",
            "Template_Assignments", "Template_Semantics"} <= set(workbook.sheetnames)
    assert workbook["Template_Semantics"].sheet_state == "hidden"
    assert workbook["Overview"]["A1"].value == "Infoblox DHCP Template & Current-State Assessment"
    assert workbook["Overview"]["A8"].value == "Archetypes"
    assert workbook["Overview"]["B8"].hyperlink.location == "'Template_Archetypes'!A1"
    assert workbook["Overview"]["A11"].value == "Provisioning archetypes"
    model_rows = rows(workbook["Template_Models"])
    assert len(model_rows) == 2  # one Network shape and one Range shape; Fixed Address is out of scope
    network_models = [row for row in model_rows if row["Template Type"] == "networktemplate"]
    assert len(network_models) == 1
    # Stored lease values differ, but use_options=False makes them inactive; inactive literals
    # must not split the reusable shape.
    assert network_models[0]["Template Count"] == 2
    assert network_models[0]["Grid Count"] == 2
    assignments = rows(workbook["Template_Assignments"])
    assert {(row["Grid"], row["Template Type"]) for row in assignments} >= {
        ("A", "networktemplate"), ("B", "networktemplate"),
        ("A", "rangetemplate"), ("B", "rangetemplate"),
    }
    bundle_rows = rows(workbook["Template_Bundles"])
    assert len(bundle_rows) == 4
    assert {row["Bundle Status"] for row in bundle_rows} == {
        "NETWORK_ONLY", "ORPHAN_RANGE_TEMPLATE"
    }
    visible = {sheet.title for sheet in workbook.worksheets if sheet.sheet_state == "visible"}
    assert visible == {
        "Overview", "Data_Index", "Template_Archetypes", "Template_Grid_Map",
        "Template_Variants", "Template_Instances", "Template_Review",
    }
    assert "Template_Bundle_Models" not in visible
    assert "Template_Bundles" not in visible
    assert "Template_Models" not in visible
    assert "Template_Grid_Matrix" not in visible
    assert "Template_Assignments" not in visible
    assert "Standardization" not in visible
    assert "Decisions" not in visible
    assert "Exceptions" not in visible
    assert "Errors" not in visible
    assert "Manual_Review" not in visible
    assert workbook["Network_Templates"].sheet_state == "hidden"
    assert workbook["Profile_Dimensions"].sheet_state == "hidden"
    index_rows = rows(workbook["Data_Index"])
    index_by_sheet = {row["Sheet"]: row for row in index_rows}
    assert index_by_sheet["Template_Archetypes"]["Visibility"] == "VISIBLE"
    assert index_by_sheet["Template_Bundle_Models"]["Visibility"] == "HIDDEN"
    assert index_by_sheet["DHCP_Raw"]["Visibility"] == "HIDDEN"
    assert index_by_sheet["DHCP_Raw"]["Open"] == "Unhide for drill-down"
    assert index_by_sheet["Standardization"]["Category"] == "Decision workflow"
    assert index_by_sheet["Standardization"]["Open"] == "Unhide for decisions"
    assert workbook["Overview"]["A12"].value == "Archetype"
    assert "Model Signature" in {
        cell.value for cell in workbook["Template_Bundle_Models"][1]
    }
    bundle_model_headers = {cell.value: cell.column for cell in workbook["Template_Bundle_Models"][1]}
    for technical in ("Network Models", "Range Models", "Geometry Signature", "Canonical Shape", "Full SHA256"):
        assert workbook["Template_Bundle_Models"].column_dimensions[
            workbook["Template_Bundle_Models"].cell(1, bundle_model_headers[technical]).column_letter
        ].hidden
    bundle_headers = {cell.value: cell.column for cell in workbook["Template_Bundles"][1]}
    for technical in ("Network Model ID", "Range Model ID", "Geometry Signature", "Stored Enabled Options"):
        assert workbook["Template_Bundles"].column_dimensions[
            workbook["Template_Bundles"].cell(1, bundle_headers[technical]).column_letter
        ].hidden
    model_headers = {cell.value: cell.column for cell in workbook["Template_Models"][1]}
    assert workbook["Template_Models"].column_dimensions[
        workbook["Template_Models"].cell(1, model_headers["Canonical Shape"]).column_letter
    ].hidden
    assert workbook["Template_Models"].column_dimensions[
        workbook["Template_Models"].cell(1, model_headers["Full SHA256"]).column_letter
    ].hidden
    matrix_rows = rows(workbook["Template_Grid_Matrix"])
    assert {"A", "B"} <= {cell.value for cell in workbook["Template_Grid_Matrix"][1]}
    assert all(json.loads(row["extra_fields"])["future_option"] == ["opaque"] for row in options)
    assert workbook["DHCP_Effective"].max_row == 1
    assert workbook["DHCP_Options"].max_row == 1
    assert {row["name"] for row in rows(workbook["Members_Failover"])} == {"member"}
    assert {row["name"] for row in rows(workbook["DHCP_Failover"])} == {"pair"}
    assert all(row["normalization_status"] == "PARTIAL" for row in rows(workbook["Fixed_Address_Templates"]))
    assert sum(row["Collection Status"] == "PARTIAL" and row["Query"] == "raw"
               for row in rows(workbook["Coverage"])) == 2
    assert not any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row)
    workbook.close()
    summary = (tmp_path / "current_state_summary.md").read_text(encoding="utf-8")
    assert "RAW_CONFIGURATION" in summary and "raw-query coverage" in summary


def test_optional_null_template_fields_do_not_mark_record_partial():
    result = CollectionResult(grid="LAB", records={
        "networktemplate": [{
            "_ref": "networktemplate/nulls:one",
            "name": "null-safe",
            "netmask": 24,
            "members": [{"name": "member-a", "ipv4addr": "10.0.0.10", "ipv6addr": None}],
            "delegated_member": None,
        }],
        "rangetemplate": [{
            "_ref": "rangetemplate/nulls:one",
            "name": "range-null-safe",
            "offset": 10,
            "number_of_addresses": 20,
            "member": {"name": "member-a", "ipv4addr": "10.0.0.10", "ipv6addr": None},
            "ms_server": None,
        }],
    })

    records = normalize_topology(result)
    by_type = {record.object_type: record for record in records}
    assert by_type["networktemplate"].status == "COMPLETE"
    assert by_type["rangetemplate"].status == "COMPLETE"
    assert by_type["networktemplate"].issues == []
    assert by_type["rangetemplate"].issues == []


def test_core_only_report_does_not_gain_topology_sheets(tmp_path):
    write_reports([CollectionResult(grid="CORE", records={"networkview": [{"name": "default"}]})], tmp_path)
    workbook = load_workbook(tmp_path / "current_state_inventory.xlsx")
    assert not set(TOPOLOGY_SHEETS.values()).intersection(workbook.sheetnames)
    assert "Template_Options" not in workbook.sheetnames
    workbook.close()
