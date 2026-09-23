"""Regression tests across HTTP collection, immutable evidence and offline reports."""
from copy import deepcopy
import json
from urllib.parse import parse_qs, urlsplit

from openpyxl import load_workbook
import pytest
import responses

from infoblox_inventory import cli
from infoblox_inventory.client import InfobloxClient
from infoblox_inventory.collectors.core import OBJECTS, collect_grid
from infoblox_inventory.collectors.topology import TOPOLOGY_OBJECTS
from infoblox_inventory.config import GridConfig
from infoblox_inventory.models import CollectionResult
from infoblox_inventory.normalize import normalize_options
from infoblox_inventory.storage import RawStore, load_raw, path_component


BASE = "https://grid.example/wapi/v2.13.7"


def schema(*names, overrides=None, unreadable=()):
    overrides = overrides or {}
    return {"fields": [
        {"name": name, "supports": "wu" if name in unreadable else "rwu",
         **({"overridden_by": overrides[name]} if name in overrides else {})}
        for name in names
    ]}


def add_root(objects):
    responses.add(responses.GET, BASE + "/", json={
        "supported_objects": objects, "supported_versions": ["2.14", "2.13.7", "2.9"],
    })


def client():
    return InfobloxClient(GridConfig("LAB-GRID", "https://grid.example"), "readonly", "secret")


def normalized(result):
    return normalize_options(result.grid, result.grid_url, result.wapi_version,
                             result.records, result.effective_records, result.schemas)


@responses.activate
def test_core_selection_queries_only_the_six_existing_lab_objects():
    expected = {"networkview", "grid:dhcpproperties", "member", "member:dhcpproperties", "network", "range"}
    add_root([*sorted(expected), "fixedaddress", "dhcpfailover"])
    for object_type in sorted(expected):
        field = "options" if object_type == "grid:dhcpproperties" else "name" if object_type in {"member", "member:dhcpproperties", "networkview", "range"} else "network"
        responses.add(responses.GET, BASE + "/" + object_type, json=schema(field))
        responses.add(responses.GET, BASE + "/" + object_type, json={"result": []})
    with client() as connection:
        result = collect_grid(connection, only="core")
    assert set(result.records) == expected
    assert not result.errors
    assert len(responses.calls) == 13  # One root schema and schema/data for six objects.
    assert {urlsplit(call.request.url).path.removeprefix("/wapi/v2.13.7/") for call in responses.calls} == {"", *expected}
    assert all(call.request.method == "GET" for call in responses.calls)


@responses.activate
def test_template_collector_requests_profile_comparable_ddns_and_schema_use_flags():
    objects = ("networktemplate", "rangetemplate", "fixedaddresstemplate")
    add_root(list(objects))
    overrides = {
        "networktemplate": {
            "enable_ddns": "use_enable_ddns",
            "ddns_domainname": "use_ddns_domainname",
            "ddns_generate_hostname": "use_ddns_generate_hostname",
            "ddns_ttl": "use_ddns_ttl",
            "ddns_update_fixed_addresses": "use_ddns_update_fixed_addresses",
            "ddns_use_option81": "use_ddns_use_option81",
            "update_dns_on_lease_renewal": "use_update_dns_on_lease_renewal",
        },
        "rangetemplate": {
            "enable_ddns": "use_enable_ddns",
            "ddns_domainname": "use_ddns_domainname",
            "ddns_generate_hostname": "use_ddns_generate_hostname",
            "update_dns_on_lease_renewal": "use_update_dns_on_lease_renewal",
        },
        "fixedaddresstemplate": {
            "enable_ddns": "use_enable_ddns",
            "ddns_domainname": "use_ddns_domainname",
        },
    }
    for object_type in objects:
        requested = TOPOLOGY_OBJECTS[object_type][2]
        flags = list(overrides[object_type].values())
        responses.add(
            responses.GET, BASE + "/" + object_type,
            json=schema(*requested, *flags, overrides=overrides[object_type]),
        )
        responses.add(responses.GET, BASE + "/" + object_type, json={"result": []})

    with client() as connection:
        result = collect_grid(connection, only="templates")

    assert not result.errors
    data_queries = {}
    for call in responses.calls:
        query = parse_qs(urlsplit(call.request.url).query)
        if "_return_fields" not in query:
            continue
        object_type = urlsplit(call.request.url).path.rsplit("/", 1)[-1]
        data_queries[object_type] = set(query["_return_fields"][0].split(","))

    assert set(data_queries) == set(objects)
    for object_type in objects:
        assert set(overrides[object_type]) <= data_queries[object_type]
        assert set(overrides[object_type].values()) <= data_queries[object_type]
    assert "ddns_ttl" in data_queries["networktemplate"]
    assert "ddns_ttl" not in data_queries["rangetemplate"]
    assert "ddns_hostname" in data_queries["fixedaddresstemplate"]
    assert all(
        row["Collection Status"] == "EMPTY"
        for row in result.coverage
        if row.get("Object") in objects and row.get("Query") == "raw"
    )


@responses.activate
def test_live_cli_collects_all_dhcp_levels_and_offline_report_replays_same_evidence(tmp_path, monkeypatch):
    grid_ref = "grid:dhcpproperties/a:LAB-GRID"
    member_ref = "member:dhcpproperties/b:infobloxlab.local"
    network_ref = "network/c:10.77.10.0/24/default"
    range_ref = "range/d:10.77.10.100-10.77.10.199/default"
    lease = {"name": "dhcp-lease-time", "num": 51, "vendor_class": "DHCP", "value": "28800", "use_option": False}
    identities = {
        "grid:dhcpproperties": {"_ref": grid_ref},
        "member:dhcpproperties": {"_ref": member_ref, "name": "infobloxlab.local"},
        "network": {"_ref": network_ref, "network": "10.77.10.0/24", "network_view": "default"},
        "range": {"_ref": range_ref, "network": "10.77.10.0/24", "network_view": "default",
                  "start_addr": "10.77.10.100", "end_addr": "10.77.10.199"},
    }
    add_root(list(identities))
    pages = {}
    for object_type, identity in identities.items():
        root = object_type == "grid:dhcpproperties"
        fields = [key for key in identity if key != "_ref"] + ["options", "nextserver"]
        overrides = {} if root else {"options": "use_options", "nextserver": "custom_override_flag"}
        fields += list(overrides.values())
        responses.add(responses.GET, BASE + "/" + object_type,
                      json=schema(*fields, overrides=overrides),
                      match=[responses.matchers.query_param_matcher(
                          {"_schema": "1", "_schema_version": "2", "_schema_searchable": "1"})])
        raw = {**identity, "options": [{**lease, "value": "28800" if root else "43200", "use_option": root}],
               "nextserver": "10.77.0.10" if root else "0.0.0.0",
               "unknown_extension": {"preserve": [None, False, "zażółć"]}}
        if not root:
            raw.update(use_options=False, custom_override_flag=False)
        effective = {**identity, "options": [{"inherited": True, "source": grid_ref, "values": [lease]}],
                     "nextserver": {"inherited": True, "multisource": False, "source": member_ref,
                                    "value": "10.77.0.20"}}
        if object_type == "range":
            effective["options"] += [
                {"inherited": True, "source": network_ref, "values": [
                    {"name": "routers", "num": 3, "value": "10.77.10.1", "use_option": False}]},
            ]
        for mode, row in [("raw", raw)] + ([] if root else [("effective", effective)]):
            pages[(object_type, mode)] = [{"result": [row], "opaque_metadata": {"mode": mode}}]
            if object_type == "network":
                pages[(object_type, mode)][0]["next_page_id"] = f"network:{mode}:2"
                pages[(object_type, mode)].append({"result": [{**row, "_ref": "network/e:10.77.20.0/24/default",
                                                              "network": "10.77.20.0/24"}]})

    requested_queries = []

    def serve_query(request):
        parsed = urlsplit(request.url)
        object_type = parsed.path.rsplit("/", 1)[-1]
        query = parse_qs(parsed.query)
        if "_page_id" in query:
            _, mode, number = query["_page_id"][0].split(":")
            page = int(number) - 1
        else:
            mode = "effective" if query.get("_inheritance") == ["True"] else "raw"
            page = 0
            requested_queries.append((object_type, mode, query))
        return 200, {"Content-Type": "application/json"}, json.dumps(pages[(object_type, mode)][page])

    for object_type in identities:
        responses.add_callback(responses.GET, BASE + "/" + object_type, callback=serve_query)
    config = tmp_path / "grids.yaml"
    config.write_text("grids:\n  - name: LAB-GRID\n    url: https://grid.example\n    username: readonly\n", encoding="utf-8")
    monkeypatch.setenv("INFOBLOX_PASSWORD", "secret")
    live = tmp_path / "live"
    assert cli.main(["--config", str(config), "--collector", "options", "--output-dir", str(live)]) == 0
    collected = load_raw(live)[0]
    assert set(collected.records) == set(identities)
    assert set(collected.effective_records) == set(identities) - {"grid:dhcpproperties"}
    assert len(collected.records["network"]) == len(collected.effective_records["network"]) == 2
    member_lease = next(row for row in normalized(collected)
                        if row["object_type"] == "member:dhcpproperties" and row["option_number"] == 51)
    assert member_lease["raw_value"] == "43200"
    assert member_lease["effective_value"] == "28800"
    assert member_lease["source_ref"] == grid_ref
    router = next(row for row in normalized(collected) if row["parameter"] == "routers")
    assert router["effective_value"] == "10.77.10.1"
    assert router["configured_here"] is False
    assert router["source_level"] == "Network"
    assert all(call.request.method == "GET" for call in responses.calls)
    assert len([call for call in responses.calls if urlsplit(call.request.url).path.endswith("/")]) == 1
    assert all("/wapi/v2.13.7/" in call.request.url for call in responses.calls)
    for object_type, _mode, query in requested_queries:
        assert query["_paging"] == query["_return_as_object"] == ["1"]
        if object_type != "grid:dhcpproperties":
            assert "custom_override_flag" in query["_return_fields"][0].split(",")
    for (object_type, mode), envelopes in pages.items():
        for number, envelope in enumerate(envelopes, 1):
            path = live / "raw" / "LAB-GRID" / path_component(object_type) / mode / f"page-{number:06d}.json"
            assert json.loads(path.read_text(encoding="utf-8")) == envelope

    def forbidden(*args, **kwargs):
        pytest.fail("Offline replay attempted live configuration, authentication or networking")

    monkeypatch.setattr(cli, "InfobloxClient", forbidden)
    monkeypatch.setattr(cli, "load_config", forbidden)
    monkeypatch.setattr(cli.getpass, "getpass", forbidden)
    monkeypatch.delenv("INFOBLOX_PASSWORD")
    offline = tmp_path / "offline"
    assert cli.main(["--fixture", str(live), "--output-dir", str(offline)]) == 0
    assert (offline / "current_state_summary.md").read_bytes() == (live / "current_state_summary.md").read_bytes()
    live_book = load_workbook(live / "current_state_inventory.xlsx")
    offline_book = load_workbook(offline / "current_state_inventory.xlsx")
    assert live_book.sheetnames == offline_book.sheetnames
    for sheet in live_book:
        assert list(sheet.values) == list(offline_book[sheet.title].values)
        assert sheet.freeze_panes == "A2"
    live_book.close()
    offline_book.close()


@responses.activate
def test_write_only_network_template_origin_is_coverage_only_not_collection_partial():
    add_root(["network"])
    requested = OBJECTS["network"][2]
    responses.add(
        responses.GET, BASE + "/network",
        json=schema(*requested, "template", unreadable={"template"}),
    )
    responses.add(responses.GET, BASE + "/network", json={"result": []})

    with client() as connection:
        result = collect_grid(connection, only="network")

    raw = next(row for row in result.coverage
               if row.get("Object") == "network" and row.get("Query") == "raw")
    assert raw["Collection Status"] == "EMPTY"
    origin = next(row for row in result.coverage
                  if row.get("Object") == "network" and row.get("Field") == "template")
    assert origin["Collection Status"] == "NOT_EXPOSED_BY_WAPI"
    query = parse_qs(urlsplit(responses.calls[-1].request.url).query)
    assert "template" not in set(query["_return_fields"][0].split(","))


@responses.activate
def test_readable_network_template_origin_is_preserved_without_becoming_required():
    add_root(["network"])
    requested = OBJECTS["network"][2]
    responses.add(responses.GET, BASE + "/network", json=schema(*requested, "template"))
    responses.add(
        responses.GET, BASE + "/network",
        json={"result": [{"_ref": "network/a", "network": "10.0.0.0/24",
                           "template": "Branch_LAN"}]},
    )

    with client() as connection:
        result = collect_grid(connection, only="network")

    assert result.records["network"][0]["template"] == "Branch_LAN"
    assert not any(
        row.get("Object") == "network"
        and row.get("Field") == "template"
        and row.get("Collection Status") == "NOT_EXPOSED_BY_WAPI"
        for row in result.coverage
    )
    query = parse_qs(urlsplit(responses.calls[-1].request.url).query)
    assert "template" in set(query["_return_fields"][0].split(","))


@responses.activate
def test_missing_unreadable_fields_and_dynamic_override_flag_are_reported(tmp_path):
    add_root(["network"])
    responses.add(responses.GET, BASE + "/network", json=schema(
        "network", "nextserver", "strange_override", "options", overrides={"nextserver": "strange_override"},
        unreadable={"options"}))
    responses.add(responses.GET, BASE + "/network", json={"result": [{"_ref": "network/a", "network": "10.0.0.0/24"}]})
    responses.add(responses.GET, BASE + "/network", json={"result": []})
    with client() as connection:
        result = collect_grid(connection, str(tmp_path), "networks")
    queries = [parse_qs(urlsplit(call.request.url).query) for call in responses.calls[2:]]
    assert len(queries) == 2
    assert set(queries[0]["_return_fields"][0].split(",")) == {"network", "nextserver", "strange_override"}
    assert queries[1]["_inheritance"] == ["True"]
    unsupported = {row.get("Field") for row in result.coverage if row["Collection Status"] == "NOT_EXPOSED_BY_WAPI"}
    assert {"options", "bootfile"} <= unsupported
    assert any(row.get("Query") == "raw" and row["Collection Status"] == "PARTIAL" for row in result.coverage)
    assert not result.errors


@responses.activate
def test_object_absence_is_not_an_http_schema_failure():
    add_root([])
    with client() as connection:
        missing = collect_grid(connection, only="network")
    assert not missing.errors
    assert any(row.get("Object") == "network" and row["Collection Status"] == "NOT_EXPOSED_BY_WAPI"
               for row in missing.coverage)
    assert len(responses.calls) == 1
    responses.reset()
    add_root(["network"])
    responses.add(responses.GET, BASE + "/network", status=403)
    with client() as connection:
        denied = collect_grid(connection, only="network")
    assert denied.errors[0]["object_type"] == "network"
    assert "403" in denied.errors[0]["error"]
    assert any(row["Collection Status"] == "ERROR" for row in denied.coverage)
    assert not any(row["Collection Status"] == "NOT_EXPOSED_BY_WAPI" for row in denied.coverage)
    assert len(responses.calls) == 2


@responses.activate
def test_missing_override_metadata_does_not_invent_effective_inheritance():
    add_root(["network"])
    responses.add(responses.GET, BASE + "/network", json=schema("network", "options", "use_options"))
    responses.add(responses.GET, BASE + "/network", json={"result": [{
        "_ref": "network/a", "network": "10.0.0.0/24", "use_options": False,
        "options": [{"name": "dhcp-lease-time", "num": 51, "value": "43200", "use_option": False}],
    }]})
    with client() as connection:
        result = collect_grid(connection, only="network")
    assert not result.effective_records
    assert normalized(result)[0]["effective_value"] is None
    assert normalized(result)[0]["status"] == "PARTIAL"
    assert len(responses.calls) == 3


@responses.activate
@pytest.mark.parametrize("malformed", [False, True])
def test_failed_second_page_preserves_partial_records_and_raw_evidence(tmp_path, malformed):
    add_root(["networkview"])
    responses.add(responses.GET, BASE + "/networkview", json=schema("name", "comment", "extattrs"))
    first = {"result": [{"_ref": "networkview/a:default", "name": "default"}], "next_page_id": "later-page"}
    responses.add(responses.GET, BASE + "/networkview", json=first)
    if malformed:
        responses.add(responses.GET, BASE + "/networkview", json={"result": "unexpected server response"})
    else:
        responses.add(responses.GET, BASE + "/networkview", status=403)
    with client() as connection:
        result = collect_grid(connection, str(tmp_path), "views")
    assert result.records["networkview"] == first["result"]
    assert result.errors[0]["query"] == "raw"
    assert any(row.get("Query") == "raw" and row["Collection Status"] == "PARTIAL"
               and row["Objects Found"] == 1 for row in result.coverage)
    archive = tmp_path / "LAB-GRID"
    assert json.loads((archive / "networkview/raw/page-000001.json").read_text()) == first
    if malformed:
        assert json.loads((archive / "networkview/raw/page-000002.json").read_text()) == {
            "result": "unexpected server response"}
    manifest = json.loads((archive / "manifest.json").read_text())
    assert manifest["queries"][0]["status"] == "PARTIAL"
    replay = load_raw(tmp_path)[0]
    assert replay.records == result.records
    assert replay.errors == result.errors
    assert replay.coverage == result.coverage


def test_archives_preserve_unknown_json_and_refuse_overwrites(tmp_path):
    result = CollectionResult("../Grid:żółć", grid_url="https://grid.example", wapi_version="2.13.7")
    store = RawStore(tmp_path, result)
    page = {"result": [{"_ref": "network/a", "unknown": [None, True, {"text": "zażółć", "value": 1.25}]}],
            "unknown_envelope": {"flag": False}}
    original = deepcopy(page)
    query = store.start_query("network", "raw", {"_return_fields": "network"})
    store.save_page(query, 1, page)
    store.finish_query(query, "COMPLETE")
    store.finish_collection()
    assert page == original
    assert store.path.resolve().is_relative_to(tmp_path.resolve())
    saved = (store.path / query["pages"][0]).read_bytes()
    with pytest.raises(FileExistsError):
        store.save_page(query, 1, {"result": []})
    with pytest.raises(ValueError, match="already exists"):
        RawStore(tmp_path, result)
    assert (store.path / query["pages"][0]).read_bytes() == saved
    replay = load_raw(tmp_path)[0]
    assert replay.records["network"] == original["result"]
    assert not replay.errors


def test_offline_rejects_archive_paths_outside_grid_directory(tmp_path):
    store = RawStore(tmp_path, CollectionResult("LAB-GRID"))
    manifest_path = store.path / "manifest.json"
    manifest = json.loads(manifest_path.read_text())
    manifest["schemas"] = {"network": "../external.json"}
    (tmp_path / "external.json").write_text('{"fields": []}')
    manifest_path.write_text(json.dumps(manifest))
    with pytest.raises(ValueError, match="escapes"):
        load_raw(tmp_path)


@responses.activate
@pytest.mark.parametrize("invalid_kind", ["bare_list", "error_with_result", "invalid_token"])
def test_rejected_effective_pages_are_saved_but_never_normalized(tmp_path, invalid_kind):
    add_root(["network"])
    responses.add(responses.GET, BASE + "/network", json=schema(
        "network", "options", "use_options", overrides={"options": "use_options"}))
    raw = {"_ref": "network/a:10.0.0.0/24/default", "network": "10.0.0.0/24",
           "use_options": False, "options": [{"name": "routers", "num": 3, "value": "shadow"}]}
    responses.add(responses.GET, BASE + "/network", json={"result": [raw]})
    untrusted_rows = [{"_ref": raw["_ref"], "options": [{"inherited": False, "source": "", "values": [
        {"name": "routers", "num": 3, "value": "must-not-be-effective"}]}]}]
    invalid_page = (untrusted_rows if invalid_kind == "bare_list" else
                    {"Error": "query failed", "result": untrusted_rows} if invalid_kind == "error_with_result" else
                    {"result": untrusted_rows, "next_page_id": 42})
    responses.add(responses.GET, BASE + "/network", json=invalid_page)
    with client() as connection:
        result = collect_grid(connection, str(tmp_path), "network")
    archived = tmp_path / "LAB-GRID/network/effective/page-000001.json"
    assert json.loads(archived.read_text(encoding="utf-8")) == invalid_page
    assert result.effective_records["network"] == []
    assert result.errors
    replay = load_raw(tmp_path)[0]
    assert replay.effective_records == result.effective_records
    assert all(row["effective_value"] is None for row in normalized(result))
    assert normalized(replay) == normalized(result)
