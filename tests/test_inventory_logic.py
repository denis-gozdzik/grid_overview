from infoblox_inventory.analysis import compare_options, naming_analysis
from infoblox_inventory.models import CollectionResult
from infoblox_inventory.normalize import normalize_options
from infoblox_inventory.report import write_reports
from openpyxl import load_workbook


def test_plain_options_do_not_claim_effective_or_inherited_configuration():
    records = {"network": [{"_ref": "network/1", "network": "10.0.0.0/24", "network_view": "default",
                            "options": [{"name": "routers", "use_option": True, "value": "10.0.0.1"},
                                        {"name": "domain-name", "use_option": False, "value": ""}]}]}
    rows = normalize_options("JB", "https://grid", "2.13.7", records)
    assert {row["raw_value"] for row in rows} == {"10.0.0.1", ""}
    assert all(row["status"] == "PARTIAL" for row in rows)
    assert all(row["effective_value"] is None for row in rows)
    assert all(row["inherited"] is None for row in rows)


def test_comparison_calls_common_observed_value_not_standard():
    rows = [
        {"grid": "JB", "object_type": "network", "source_level": "Network", "parameter": "routers", "object_name": "site-a", "effective_value": "10.0.0.1", "status": "COMPLETE"},
        {"grid": "GRID2", "object_type": "network", "source_level": "Network", "parameter": "routers", "object_name": "site-a", "effective_value": "10.0.0.254", "status": "COMPLETE"},
    ]
    differences = compare_options(rows)
    assert {row["common_observed_value"] for row in differences} == {None}
    assert all("standard_value" not in row for row in differences)
    assert {row["classification"] for row in differences} == {"DIFFERENT"}


def test_comparison_matches_queried_objects_despite_different_inheritance_sources():
    common = {"object_type": "range", "object_name": "10.0.0.100-10.0.0.199", "network_view": "default",
              "parameter": "routers", "effective_value": "10.0.0.1", "status": "COMPLETE"}
    rows = [{**common, "grid": "A", "source_level": "Grid"},
            {**common, "grid": "B", "source_level": "Network"}]
    differences = compare_options(rows)
    assert len(differences) == 2
    assert {row["classification"] for row in differences} == {"SOURCE_DIFFERENCE"}
    assert {row["common_observed_value"] for row in differences} == {"10.0.0.1"}
    assert {row["occurrence_count"] for row in differences} == {2}


def test_comparison_excludes_unresolved_and_multisource_rows():
    assert compare_options([
        {"status": "PARTIAL", "effective_value": None},
        {"status": "NOT_CONFIGURED", "effective_value": None},
        {"status": "COMPLETE", "multisource": True, "effective_value": ["a", "b"]},
    ]) == []


def test_comparison_keeps_network_views_and_option_spaces_separate():
    common = {"grid": "A", "object_type": "network", "object_name": "10.0.0.0/24", "parameter": "custom",
              "effective_value": "one", "status": "COMPLETE", "option_number": 200, "network_view": "default"}
    rows = [{**common, "vendor_class": "A"}, {**common, "vendor_class": "B"},
            {**common, "vendor_class": "A", "network_view": "other"}]
    assert len(compare_options(rows)) == 3


def test_naming_patterns_are_inferred():
    patterns = naming_analysis({"network": [{"name": "NYC-core"}, {"name": "NYC-edge"}, {"name": "LON-core"}]})
    assert any(row["observed_pattern"] == "prefix:NYC" and row["confidence"] == "OBSERVED" for row in patterns)


def test_reports_separate_raw_effective_and_preserve_literals_and_metadata(tmp_path):
    ref = "network/example:10.0.0.0/24/default"
    result = CollectionResult(grid="LAB", grid_url="https://grid", wapi_version="2.13.7", collected_at="2026-09-19T10:00:00Z",
        records={"network": [{"_ref": ref, "network": "10.0.0.0/24", "comment": "=1+1", "nextserver": "shadow",
                               "extattrs": {"Owner": {"value": "network-team"}}, "use_nextserver": False,
                               "options": [{"name": "routers", "num": 3, "value": "shadow", "use_option": False}]}]},
        effective_records={"network": [{"_ref": ref,
            "nextserver": {"inherited": True, "multisource": False, "source": "grid:dhcpproperties/example:LAB", "value": "10.0.0.10"},
            "options": [{"inherited": True, "source": "grid:dhcpproperties/example:LAB",
                          "values": [{"name": "routers", "num": 3, "value": "10.0.0.1", "use_option": False}]}]}]},
        schemas={"network": {"fields": [{"name": "nextserver", "overridden_by": "use_nextserver"}]}},
        coverage=[{"Grid": "LAB", "Area": "DHCP options", "Collection Status": "COMPLETE", "Notes": ""}])
    failed = CollectionResult(grid="FAILED", collected_at="2026-09-19T10:00:00Z", wapi_version="2.13.7")
    failed.fail("Network Views", "networkview", RuntimeError("read denied"))
    write_reports([result, failed], tmp_path)
    workbook = load_workbook(tmp_path / "current_state_inventory.xlsx")
    assert {"DHCP_Raw", "DHCP_Effective", "DHCP_Options", "Coverage", "Manual_Review"} <= set(workbook.sheetnames)
    options = workbook["DHCP_Options"]
    headers = [cell.value for cell in options[1]]
    option_row = dict(zip(headers, [cell.value for cell in options[2]]))
    assert option_row["raw_value"] == "shadow"
    assert option_row["effective_value"] == "10.0.0.1"
    assert option_row["wapi_version"] == "2.13.7"
    assert option_row["grid_url"] == "https://grid"
    assert '"Owner"' in option_row["extensible_attributes"]
    assert any(cell.value == "=1+1" and cell.data_type == "s" for sheet in workbook for row in sheet for cell in row)
    assert not any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row)
    grid_headers = [cell.value for cell in workbook["Grid_Summary"][1]]
    grid_index = grid_headers.index("Grid")
    assert {row[grid_index].value for row in workbook["Grid_Summary"].iter_rows(min_row=2)} == {"LAB", "FAILED"}
    summary = (tmp_path / "current_state_summary.md").read_text(encoding="utf-8")
    assert "read denied" in summary
    assert "2026-09-19T10:00:00Z" in summary
    assert "COMPLETE: 1" in summary
    assert "remain to be implemented" in summary
    assert all(sheet.freeze_panes == "A2" and sheet.auto_filter.ref is None for sheet in workbook)
    assert all(table.autoFilter.ref == table.ref for sheet in workbook for table in sheet.tables.values())


def test_markdown_aggregates_field_coverage_into_concise_area_rows(tmp_path):
    result = CollectionResult(grid="LAB", coverage=[
        {"Grid": "LAB", "Area": "DHCP", "Collection Status": "NOT_EXPOSED_BY_WAPI", "Field": f"field{number}"}
        for number in range(100)
    ])
    write_reports([result], tmp_path)
    summary = (tmp_path / "current_state_summary.md").read_text(encoding="utf-8")
    assert "NOT_EXPOSED_BY_WAPI: 100" in summary
    assert len(summary.splitlines()) < 80
