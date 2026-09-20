"""Validate one saved LAB-GRID capture and rebuild its reports without networking.

Usage: python scripts/validate_lab_snapshot.py output/<timestamped-live-run>
The raw archive is never edited. Existing report/validation outputs are refused.
"""
from __future__ import annotations

import argparse
from contextlib import ExitStack, contextmanager
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import sys
from typing import Any, Iterator
from unittest.mock import patch

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from infoblox_inventory import cli  # noqa: E402
from infoblox_inventory.normalize import normalize_options, normalize_scalars, parse_source_ref  # noqa: E402
from infoblox_inventory.storage import load_raw, page_records  # noqa: E402


GRID = "LAB-GRID"
VERSION = "2.13.7"
MEMBER = "infobloxlab.local"
NETWORK = "10.77.10.0/24"
RANGE = "10.77.10.100-10.77.10.199"
OBJECTS = {"networkview", "grid:dhcpproperties", "member", "member:dhcpproperties", "network", "range"}
EFFECTIVE_OBJECTS = {"member:dhcpproperties", "network", "range"}
EXPECTED_OPTIONS = {
    "51": ("dhcp-lease-time", "3600", "Range", RANGE),
    "3": ("routers", "10.77.10.1", "Network", NETWORK),
    "6": ("domain-name-servers", "10.77.0.60,10.77.0.61", "Member", MEMBER),
    "15": ("domain-name", "lab.local", "Grid", GRID),
}


def contained(base: Path, relative: str) -> Path:
    destination = (base / relative).resolve()
    if destination == base.resolve() or not destination.is_relative_to(base.resolve()):
        raise ValueError("Validation path escapes its containing directory")
    return destination


def read_json(base: Path, relative: str) -> Any:
    return json.loads(contained(base, relative).read_text(encoding="utf-8"))


def hashes(raw: Path) -> dict[str, str]:
    if not raw.is_dir():
        raise ValueError("The run has no raw directory")
    result = {}
    for path in sorted(raw.rglob("*")):
        if path.is_file():
            resolved = contained(raw, path.relative_to(raw).as_posix())
            result[path.relative_to(raw).as_posix()] = hashlib.sha256(resolved.read_bytes()).hexdigest()
    if not result:
        raise ValueError("The raw directory is empty")
    return result


@contextmanager
def no_network(attempts: list[str]) -> Iterator[None]:
    def denied(label: str):
        def reject(*_args: Any, **_kwargs: Any) -> None:
            attempts.append(label)
            raise RuntimeError(f"Network access forbidden during offline validation: {label}")
        return reject

    with ExitStack() as stack:
        for target in ("requests.sessions.Session.request", "requests.sessions.Session.send",
                       "socket.socket.connect", "socket.socket.connect_ex", "socket.create_connection",
                       "socket.getaddrinfo"):
            stack.enter_context(patch(target, denied(target)))
        yield


def check(report: dict[str, Any], name: str, expected: Any, observed: Any,
          *, passed: bool | None = None, not_observed: bool = False) -> None:
    status = "NOT_OBSERVED" if not_observed else "PASS" if (observed == expected if passed is None else passed) else "FAIL"
    report["checks"].append({"check": name, "status": status, "expected": expected, "observed": observed})


def member_present(value: Any) -> bool:
    if isinstance(value, dict):
        return any(member_present(item) for item in value.values())
    if isinstance(value, list):
        return any(member_present(item) for item in value)
    return isinstance(value, str) and (value == MEMBER or parse_source_ref(value)["source_object"] == MEMBER)


def normalized(result: Any) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    args = (result.grid, result.grid_url, result.wapi_version, result.records, result.effective_records, result.schemas)
    return normalize_scalars(*args), normalize_options(*args)


def range_row(row: dict[str, Any]) -> bool:
    return row.get("object_type") == "range" and row.get("object_name") == RANGE and row.get("network_view") == "default"


def projection(row: dict[str, Any]) -> dict[str, Any]:
    return {key: row.get(key) for key in ("parameter", "option_number", "effective_value", "configured_here",
            "inherited", "source_level", "source_object", "source_ref", "status", "raw_value", "raw_use_option")}


def pagination(report: dict[str, Any], archive: Path, manifest: dict[str, Any]) -> None:
    expected = {(item, "raw") for item in OBJECTS} | {(item, "effective") for item in EFFECTIVE_OBJECTS}
    observed = {(item["object_type"], item["mode"]) for item in manifest["queries"]}
    check(report, "query_scope", sorted(expected), sorted(observed))
    for query in manifest["queries"]:
        pages = [read_json(archive, relative) for relative in query["pages"]]
        count = sum(len(page_records(page)) for page in pages)
        tokens = [page.get("next_page_id") for page in pages]
        complete = (bool(pages) and tokens[-1] is None and all(tokens[:-1])
                    and len(tokens[:-1]) == len(set(tokens[:-1]))
                    and query["status"] in {"COMPLETE", "EMPTY"}
                    and count == query["record_count"])
        detail = {"object_type": query["object_type"], "mode": query["mode"], "pages": len(pages),
                  "record_count": count, "manifest_record_count": query["record_count"],
                  "manifest_status": query["status"], "final_page_token_exhausted": bool(pages) and tokens[-1] is None,
                  "all_intermediate_pages_have_tokens": all(tokens[:-1]), "complete": complete}
        report["pagination"].append(detail)
        check(report, f"pagination:{query['object_type']}:{query['mode']}", True, complete)


def validate_normalized(report: dict[str, Any], scalars: list[dict[str, Any]],
                        options: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    scalar_rows = [row for row in scalars if range_row(row) and row.get("parameter") == "nextserver"]
    option_rows = [row for row in options if range_row(row)]
    check(report, "range_nextserver_row_count", 1, len(scalar_rows))
    check(report, "range_effective_option_numbers", sorted(EXPECTED_OPTIONS), sorted(str(row.get("option_number")) for row in option_rows))
    expected_rows = [("nextserver", scalar_rows, "10.77.0.40", "Range", RANGE)]
    for number, (name, value, level, source) in EXPECTED_OPTIONS.items():
        matches = [row for row in option_rows if str(row.get("option_number")) == number]
        check(report, f"range_option_{number}_row_count", 1, len(matches))
        expected_rows.append((name, matches, value, level, source))
    for name, matches, value, level, source in expected_rows:
        expected = {"parameter": name, "effective_value": value, "configured_here": level == "Range",
                    "inherited": level != "Range", "source_level": level, "source_object": source,
                    "status": "COMPLETE"}
        observed = {key: matches[0].get(key) for key in expected} if len(matches) == 1 else None
        # NIOS option values are textual; scalar lease numbers may also be returned as numbers.
        equivalent = dict(observed) if observed is not None else None
        if equivalent is not None:
            equivalent["effective_value"] = str(equivalent["effective_value"])
        check(report, f"range_effective:{name}", expected, observed, passed=equivalent == expected)
        report["expected_vs_observed"].append({"parameter": name, "expected_value": value,
            "expected_source_level": level, "observed": [projection(row) for row in matches]})
        if len(matches) == 1:
            row = matches[0]
            wrapper = row.get("inheritance_details")
            original = wrapper.get("source") if isinstance(wrapper, dict) else None
            original_valid = original == "" if level == "Range" else isinstance(original, str) and bool(original)
            expected_ref = row.get("object_ref") if level == "Range" else original
            preserved = original_valid and bool(expected_ref) and row.get("source_ref") == expected_ref
            check(report, f"source_reference:{name}", "Original wrapper source retained; normalized ref identifies its source",
                  {"original_source": original, "normalized_source_ref": row.get("source_ref"),
                   "object_ref": row.get("object_ref")}, passed=preserved)
    return scalar_rows, option_rows


def shadow_evidence(report: dict[str, Any], result: Any, options: list[dict[str, Any]]) -> None:
    # Inspect original records independently; a broken normalizer must not turn a
    # present regression case into a misleading NOT_OBSERVED result.
    shadows = 0
    observations = []
    for object_type, records in result.records.items():
        effective = {record.get("_ref"): record for record in result.effective_records.get(object_type, [])}
        for raw in records:
            for stored in raw.get("options", []):
                if not isinstance(stored, dict) or str(stored.get("value")) != "43200" or stored.get("use_option") is not False:
                    continue
                shadows += 1
                key = (str(stored.get("num")), str(stored.get("vendor_class") or "DHCP"))
                inherited_options = effective.get(raw.get("_ref"), {}).get("options", [])
                supplied = [item for group in inherited_options if isinstance(group, dict)
                            for item in group.get("values", []) if isinstance(item, dict)
                            and (str(item.get("num")), str(item.get("vendor_class") or "DHCP")) == key]
                if len(supplied) != 1 or str(supplied[0].get("value")) == "43200":
                    continue
                matching = [row for row in options if row.get("object_type") == object_type
                            and row.get("object_ref") == raw.get("_ref")
                            and (str(row.get("option_number")), str(row.get("vendor_class") or "DHCP")) == key]
                correct = (len(matching) == 1 and matching[0].get("effective_value") == supplied[0].get("value")
                           and matching[0].get("raw_value") == stored["value"]
                           and matching[0].get("raw_use_option") is False and matching[0].get("status") == "COMPLETE")
                observations.append({"object_type": object_type, "object_ref": raw.get("_ref"),
                    "original_raw_option": stored, "original_effective_option": supplied[0],
                    "normalized": [projection(row) for row in matching], "matches_effective_wrapper": correct})
    report["raw_shadow_evidence"] = {"raw_43200_use_option_false_rows": shadows,
                                    "different_confirmed_effective_values": observations}
    check(report, "live_43200_shadow_does_not_override_effective", "At least one preserved 43200/use_option=false shadow with differing effective wrapper",
          observations, passed=bool(observations) and all(item["matches_effective_wrapper"] for item in observations),
          not_observed=not observations)


def member_inheritance_evidence(report: dict[str, Any], result: Any,
                                scalars: list[dict[str, Any]], options: list[dict[str, Any]]) -> None:
    """A completed inheritance GET can still supply no effective wrappers."""
    object_type = "member:dhcpproperties"
    raw_by_ref = {row.get("_ref"): row for row in result.records.get(object_type, [])}
    observations = []
    for record in result.effective_records.get(object_type, []):
        original_options = record.get("options")
        scalar = record.get("nextserver")
        unwrapped = (not isinstance(scalar, dict) and isinstance(original_options, list)
                     and all(isinstance(item, dict) and "inherited" not in item and "multisource" not in item
                             for item in original_options))
        if not unwrapped:
            continue
        matches = [row for row in scalars + options if row.get("object_type") == object_type
                   and row.get("object_ref") == record.get("_ref")]
        required = {"nextserver", "dhcp-lease-time", "domain-name-servers"}
        safe = (required.issubset({row.get("parameter") for row in matches})
                and all(row.get("status") == "PARTIAL" and row.get("effective_value") is None
                        and row.get("configured_here") is None for row in matches))
        observation = {"object_ref": record.get("_ref"), "inheritance_response_has_raw_shape": True,
                       "same_decoded_record_as_raw": raw_by_ref.get(record.get("_ref")) == record,
                       "original_nextserver": scalar, "original_options": original_options,
                       "normalized_rows": [projection(row) for row in matches],
                       "unconfirmed_values_remain_unresolved": safe}
        observations.append(observation)
        check(report, "member_unwrapped_inheritance_remains_unresolved", "PARTIAL, effective_value=None, configured_here=None",
              observation, passed=safe)
    report["member_inheritance_evidence"] = observations
    if observations:
        report.setdefault("limitations", []).append(
            "Member DHCP _inheritance=True returned plain values/options without inheritance wrappers. "
            "Its GET and pagination completed, but member effective values remain PARTIAL/None. "
            "Member raw lease 43200 is not asserted to be an effective lease; the independently observed "
            "network effective lease is 28800 and the range effective lease is 3600.")


def excel_value(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        return json.dumps(value, sort_keys=True, default=str)
    return None if value == "" else value


def workbook_snapshot(path: Path) -> dict[str, list[list[Any]]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        return {sheet.title: [list(row) for row in sheet.iter_rows(values_only=True)] for sheet in workbook.worksheets}
    finally:
        workbook.close()


def verify_excel(report: dict[str, Any], cells: dict[str, list[list[Any]]],
                 scalar_rows: list[dict[str, Any]], option_rows: list[dict[str, Any]], build: str) -> None:
    for sheet, expected_rows in (("DHCP_Effective", scalar_rows), ("DHCP_Options", option_rows)):
        matrix = cells.get(sheet, [])
        actual = [dict(zip(matrix[0], values)) for values in matrix[1:]] if matrix else []
        actual = [row for row in actual if range_row(row) and (sheet != "DHCP_Effective" or row.get("parameter") == "nextserver")]
        check(report, f"excel:{build}:{sheet}:range_row_count", len(expected_rows), len(actual))
        for expected in expected_rows:
            matches = [row for row in actual if row.get("object_ref") == expected.get("object_ref")
                       and row.get("parameter") == expected.get("parameter")
                       and row.get("option_number") == expected.get("option_number")]
            mismatches = {key: {"expected": excel_value(value), "observed": matches[0].get(key)}
                          for key, value in expected.items()
                          if len(matches) == 1 and (key not in matches[0] or matches[0].get(key) != excel_value(value))}
            check(report, f"excel:{build}:{sheet}:{expected['parameter']}", "One row matching every normalized value and metadata field",
                  {"matching_rows": len(matches), "mismatches": mismatches}, passed=len(matches) == 1 and not mismatches)
    summary = cells.get("Grid_Summary", [])
    rows = [dict(zip(summary[0], values)) for values in summary[1:]] if summary else []
    check(report, f"excel:{build}:version_and_grid", True,
          bool(rows) and all(row.get("Grid") == GRID and row.get("WAPI Version") == VERSION for row in rows))
    for sheet, association_key in (("Networks", "members"), ("Ranges", "member")):
        matrix = cells.get(sheet, [])
        rows = [dict(zip(matrix[0], values)) for values in matrix[1:]] if matrix else []
        if sheet == "Networks":
            rows = [row for row in rows if row.get("network") == NETWORK and row.get("network_view") == "default"]
        else:
            rows = [row for row in rows if row.get("start_addr") == "10.77.10.100"
                    and row.get("end_addr") == "10.77.10.199" and row.get("network_view") == "default"]
        association = rows[0].get(association_key) if len(rows) == 1 else None
        if isinstance(association, str):
            association = json.loads(association)
        valid = len(rows) == 1 and member_present(association)
        if sheet == "Ranges":
            valid = valid and rows[0].get("server_association_type") == "MEMBER"
        check(report, f"excel:{build}:{sheet}:member_association", MEMBER, association, passed=valid)
    matrix = cells.get("DHCP_Options", [])
    rows = [dict(zip(matrix[0], values)) for values in matrix[1:]] if matrix else []
    for observation in report["raw_shadow_evidence"]["different_confirmed_effective_values"]:
        supplied = observation["original_effective_option"]
        matches = [row for row in rows if row.get("object_ref") == observation["object_ref"]
                   and str(row.get("option_number")) == str(supplied.get("num"))]
        valid = (len(matches) == 1 and matches[0].get("raw_value") == "43200"
                 and matches[0].get("raw_use_option") is False
                 and matches[0].get("effective_value") == supplied["value"])
        check(report, f"excel:{build}:raw_effective:{observation['object_type']}", supplied["value"],
              [{key: row.get(key) for key in ("raw_value", "raw_use_option", "effective_value")} for row in matches], passed=valid)
    for observation in report.get("member_inheritance_evidence", []):
        valid = True
        for expected in observation["normalized_rows"]:
            sheet = "DHCP_Options" if expected.get("option_number") is not None else "DHCP_Effective"
            matrix = cells.get(sheet, [])
            member_rows = [dict(zip(matrix[0], values)) for values in matrix[1:]] if matrix else []
            matches = [row for row in member_rows if row.get("object_ref") == observation["object_ref"]
                       and row.get("parameter") == expected["parameter"]]
            valid = valid and len(matches) == 1 and all(
                row.get("status") == "PARTIAL" and row.get("effective_value") is None
                and row.get("configured_here") is None and row.get("raw_value") == excel_value(expected["raw_value"])
                for row in matches)
        check(report, f"excel:{build}:member_unwrapped_values_remain_unresolved", True, valid)


def rebuild(report: dict[str, Any], run: Path, scalars: list[dict[str, Any]], options: list[dict[str, Any]],
            scalar_rows: list[dict[str, Any]], option_rows: list[dict[str, Any]]) -> None:
    reports = contained(run, "reports")
    preserved = contained(run, "reports-first-build")
    args = ["--offline", str(contained(run, "raw")), "--grid", GRID, "--output-dir", str(reports)]
    report["offline_command_arguments"] = args
    code = cli.main(args)
    check(report, "offline_first_build_exit_code", 0, code)
    first = workbook_snapshot(reports / "current_state_inventory.xlsx")
    verify_excel(report, first, scalar_rows, option_rows, "first")
    markdown = {name: (reports / name).read_text(encoding="utf-8")
                for name in ("current_state_summary.md", "manual_review.md")}
    # Re-resolve immediately before moving and refuse a pre-existing destination.
    if reports.resolve().parent != run or preserved.resolve().parent != run or preserved.exists():
        raise ValueError("Unsafe or existing report move destination")
    reports.rename(preserved)
    check(report, "first_reports_moved_aside", True, preserved.is_dir() and not reports.exists())
    code = cli.main(args)
    check(report, "offline_second_build_exit_code", 0, code)
    second = workbook_snapshot(reports / "current_state_inventory.xlsx")
    verify_excel(report, second, scalar_rows, option_rows, "rebuilt")
    check(report, "offline_workbook_cell_matrices_identical", True, first == second)
    check(report, "offline_markdown_identical", True,
          all((reports / name).read_text(encoding="utf-8") == text for name, text in markdown.items()))
    replay = load_raw(run / "raw", GRID)
    reloaded_scalars, reloaded_options = normalized(replay[0])
    check(report, "offline_normalized_values_identical", True,
          scalars == reloaded_scalars and options == reloaded_options)


def validate(report: dict[str, Any], run: Path) -> None:
    archive = contained(run, "raw/LAB-GRID")
    manifest = read_json(archive, "manifest.json")
    results = load_raw(run / "raw", GRID)
    check(report, "one_grid_archive_loaded", 1, len(results))
    if len(results) != 1:
        raise ValueError("Expected exactly one LAB-GRID result")
    result = results[0]
    check(report, "grid_name", GRID, manifest["grid"])
    check(report, "grid_url", "https://192.168.88.243", manifest["grid_url"])
    check(report, "wapi_version", VERSION, manifest["wapi_version"])
    root = read_json(archive, manifest["schemas"]["root"])
    requested = root.get("requested_version")
    check(report, "root_schema_requested_version", VERSION, requested, not_observed=requested is None)
    check(report, "root_schema_supports_2_13_7", True, VERSION in root.get("supported_versions", []))
    check(report, "object_schema_scope", sorted(OBJECTS), sorted(set(manifest["schemas"]) - {"root"}))
    check(report, "collection_finished", True, manifest["collection_finished"])
    check(report, "collection_errors", [], result.errors)
    report["collected_at"] = manifest["collected_at"]
    report["object_counts"] = {item: {"raw": len(result.records.get(item, [])),
                                      "effective": len(result.effective_records.get(item, []))} for item in sorted(OBJECTS)}
    pagination(report, archive, manifest)
    network = [row for row in result.records.get("network", []) if row.get("network") == NETWORK and row.get("network_view") == "default"]
    ranges = [row for row in result.records.get("range", []) if row.get("start_addr") == "10.77.10.100"
              and row.get("end_addr") == "10.77.10.199" and row.get("network_view") == "default"]
    check(report, "test_network_count", 1, len(network))
    check(report, "test_range_count", 1, len(ranges))
    association = network[0].get("members") if len(network) == 1 else None
    check(report, "network_member_association", MEMBER, association, passed=member_present(association))
    association = ranges[0].get("member") if len(ranges) == 1 else None
    check(report, "range_member_association", MEMBER, association, passed=member_present(association))
    check(report, "range_server_association_type", "MEMBER", ranges[0].get("server_association_type") if len(ranges) == 1 else None)
    scalars, options = normalized(result)
    scalar_rows, option_rows = validate_normalized(report, scalars, options)
    shadow_evidence(report, result, options)
    member_inheritance_evidence(report, result, scalars, options)
    rebuild(report, run, scalars, options, scalar_rows, option_rows)


def markdown_report(report: dict[str, Any]) -> str:
    def escape(value: Any) -> str:
        return str(value).replace("|", "\\|").replace("\r", " ").replace("\n", " ")

    lines = ["# LAB-GRID offline validation", "", f"Overall result: **{report['status']}**", "",
             f"Run: `{report['run_directory']}`", "", "| Object | Raw records | Effective records |",
             "| --- | ---: | ---: |"]
    for item, counts in report["object_counts"].items():
        lines.append(f"| {item} | {counts['raw']} | {counts['effective']} |")
    lines.extend(["", "| Parameter | Expected | Observed | Expected source | Observed source |",
                  "| --- | --- | --- | --- | --- |"])
    for row in report["expected_vs_observed"]:
        observed = row["observed"]
        value = ", ".join(str(item["effective_value"]) for item in observed) or "MISSING"
        source = ", ".join(str(item["source_level"]) for item in observed) or "MISSING"
        lines.append(f"| {row['parameter']} | {escape(row['expected_value'])} | {escape(value)} | {row['expected_source_level']} | {escape(source)} |")
    lines.extend(["", "| Check | Result |", "| --- | --- |"])
    lines.extend(f"| {escape(item['check'])} | {item['status']} |" for item in report["checks"])
    if report.get("limitations"):
        lines.extend(["", "## Observed limitations", ""])
        lines.extend(f"- {escape(item)}" for item in report["limitations"])
    lines.extend(["", "NOT_OBSERVED means that this live archive does not prove the case; it is not a successful validation of that case.",
                  "Raw JSON hashes, exact observations, pagination evidence, and workbook discrepancies are in validation_results.json.",
                  "The two report builds ran with HTTP requests and socket connections blocked. Their cell values and Markdown were compared.",
                  "This utility validates a saved archive; collection request logs establish its live origin and GET-only transport."])
    return "\n".join(lines) + "\n"


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("run_directory", type=Path, help="Timestamped collection directory containing raw/LAB-GRID")
    args = parser.parse_args()
    run = args.run_directory.resolve()
    if not run.is_dir():
        parser.error("The run directory must already exist")
    for name in ("reports", "reports-first-build", "validation_results.json", "validation_report.md"):
        if contained(run, name).exists():
            parser.error(f"Refusing to overwrite existing {name}; preserve it outside these validation output names")
    report: dict[str, Any] = {"run_directory": str(run), "validated_at": datetime.now(timezone.utc).isoformat(),
        "checks": [], "object_counts": {}, "pagination": [], "expected_vs_observed": [], "network_attempts": []}
    before: dict[str, str] | None = None
    with no_network(report["network_attempts"]):
        try:
            before = hashes(contained(run, "raw"))
            report["raw_sha256_before"] = before
            validate(report, run)
        except Exception as exc:
            check(report, "validation_execution", "All validation stages completed", f"{type(exc).__name__}: {exc}", passed=False)
        finally:
            if before is not None:
                try:
                    after = hashes(contained(run, "raw"))
                    report["raw_sha256_after"] = after
                    check(report, "raw_files_unchanged", True, before == after)
                except Exception as exc:
                    check(report, "raw_files_unchanged", True, f"{type(exc).__name__}: {exc}", passed=False)
    check(report, "offline_network_attempts", [], report["network_attempts"])
    report["status"] = "FAIL" if any(item["status"] == "FAIL" for item in report["checks"]) else "PASS"
    report["not_observed"] = [item["check"] for item in report["checks"] if item["status"] == "NOT_OBSERVED"]
    with contained(run, "validation_results.json").open("x", encoding="utf-8") as stream:
        json.dump(report, stream, indent=2, ensure_ascii=False, allow_nan=False)
        stream.write("\n")
    with contained(run, "validation_report.md").open("x", encoding="utf-8") as stream:
        stream.write(markdown_report(report))
    print(f"Offline validation: {report['status']}; {len(report['not_observed'])} NOT_OBSERVED checks. Results: {run / 'validation_results.json'}")
    return 1 if report["status"] == "FAIL" else 0


if __name__ == "__main__":
    raise SystemExit(main())
