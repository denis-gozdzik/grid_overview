from __future__ import annotations

from collections import Counter, defaultdict
import json
from pathlib import Path
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.table import Table, TableStyleInfo

from .analysis import naming_analysis
from .models import CollectionResult
from .dataset import combine_collections
from .normalize import normalize_options, normalize_scalars
from .profile_readiness import (
    PROFILE_READINESS_HEADERS, PROFILE_SUMMARY_HEADERS,
    build_profile_readiness, profile_readiness_summary,
)
from .profile_population import PROFILE_POPULATION_HEADERS, profile_population_rows
from .profile_discovery import (
    PROFILE_ASSIGNMENT_HEADERS, PROFILE_FINGERPRINT_HEADERS, PROFILE_KPI_HEADERS,
    PROFILE_USEFULNESS_HEADERS, build_profile_discovery,
)
from .profile_context import (
    PROFILE_CONTEXT_HEADERS, PROFILE_RELATIONSHIP_ASSIGNMENT_HEADERS,
    PROFILE_RELATIONSHIP_HEADERS, build_profile_context, build_profile_relationships,
)
from .profile_dimensions import (
    PROFILE_COMPOSITION_HEADERS, PROFILE_DIMENSION_ASSIGNMENT_HEADERS,
    PROFILE_DIMENSION_HEADERS, PROFILE_DIMENSION_KPI_HEADERS,
    build_profile_dimensions,
)
from .topology import (TOPOLOGY_SHEETS, normalize_topology, topology_coverage,
                       topology_excel_rows, topology_headers, topology_option_rows)
from .template_semantics import (
    TEMPLATE_ASSIGNMENT_HEADERS, TEMPLATE_MODEL_HEADERS, TEMPLATE_SEMANTIC_HEADERS,
    build_template_semantic_model,
)
from .template_bundles import (
    TEMPLATE_BUNDLE_HEADERS, TEMPLATE_BUNDLE_MODEL_HEADERS,
    build_template_bundle_models, build_template_bundles,
)
from .reservations import (RESERVATION_SHEETS, normalize_reservations, reservation_coverage,
                           reservation_excel_rows, reservation_headers, reservation_option_rows,
                           reservation_relationship_rows)
from .standardization import (
    DECISION_HEADERS, DIFFERENCE_HEADERS, EXCEPTION_HEADERS, STANDARDIZATION_HEADERS,
    build_standardization, decision_rows, difference_rows, exception_rows,
    grid_comparison_rows, load_decisions, workbook_standardization_rows, write_decision_template,
)


def _json_value(value: Any) -> Any:
    return json.dumps(value, sort_keys=True, default=str) if isinstance(value, (dict, list)) else value


def _markdown_value(value: Any) -> str:
    return str(_json_value(value) if value is not None else "").replace("|", "\\|").replace("\n", " ").replace("\r", " ")


def _excel_text(value: Any) -> Any:
    value = _json_value(value)
    return (re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", lambda match: f"\\u{ord(match[0]):04x}", value)
            if isinstance(value, str) else value)


ASSESSMENT_COLUMNS = [
    ("Grid", "grid"), ("Object Type", "object_type"), ("Object", "object_name"),
    ("Network View", "network_view"), ("Parent Network", "parent_network"),
    ("Parameter", "parameter"), ("Option Number", "option_number"),
    ("Effective Value", "effective_value"), ("Configured Here", "configured_here"),
    ("Inherited", "inherited"), ("Source Level", "source_level"),
    ("Source Object", "source_object"), ("Status", "status"),
]


def assessment_table(rows, *, ddns=False):
    columns = [("DDNS Parameter" if ddns and label == "Parameter" else label, key)
               for label, key in ASSESSMENT_COLUMNS if not (ddns and key == "option_number")]
    primary = {key for _, key in columns}
    technical = sorted({key for row in rows for key in row} - primary)
    return ([{**{label: row.get(key) for label, key in columns},
              **{key: row.get(key) for key in technical}} for row in rows],
            [label for label, _ in columns] + technical)


def _effective_query_summary(result, object_type):
    present = object_type in result.effective_records
    statuses = sorted({row.get("Collection Status", "PARTIAL") for row in result.coverage
                       if row.get("Object") == object_type and row.get("Query") == "effective"})
    return {"Effective Query Rows": len(result.effective_records[object_type]) if present else None,
            "Effective Query State": (", ".join(statuses) or "QUERY_STATUS_UNAVAILABLE") if present
                                     else "NOT_REQUESTED",
            "Effective Count Meaning": "Response object count, not configured effective fields. "
                + ("See query coverage and normalized field statuses." if present else
                   "No separate effective response was captured; this does not imply absent configuration.")}


SHEET_TITLE_OVERRIDES = {
    "Profile_Relationship_Assignments": "Profile_Rel_Assignments",
}


def _excel_sheet_title(workbook: Workbook, logical_name: str) -> str:
    """Return a deterministic Excel-safe worksheet title for a logical report name."""
    configured = SHEET_TITLE_OVERRIDES.get(logical_name, logical_name)
    base = re.sub(r"[\[\]:*?/\\]", "_", str(configured)).strip().strip("'") or "Sheet"
    base = base[:31]
    existing = {sheet.title.casefold() for sheet in workbook.worksheets}
    title = base
    suffix = 2
    while title.casefold() in existing:
        ending = f"_{suffix}"
        title = f"{base[:31 - len(ending)]}{ending}"
        suffix += 1
    return title


def _write_inventory_sheet(workbook: Workbook, index: int, name: str,
                           rows: list[dict[str, Any]], headers: list[str], *,
                           preserve_order: bool = False):
    sheet = workbook.create_sheet(_excel_sheet_title(workbook, name))
    # Excel table headers must be nonempty strings and unique ignoring case.
    # Keep source keys separate so display-name repairs never move or lose data.
    labels: list[str] = []
    used: set[str] = set()
    for key in headers:
        base = str(_excel_text(key)).strip()[:255] or "Column"
        label = base
        suffix = 2
        while label.casefold() in used:
            ending = f" ({suffix})"
            label = base[:255 - len(ending)] + ending
            suffix += 1
        used.add(label.casefold())
        labels.append(label)
    sheet.append(labels)
    ordered = rows if preserve_order else sorted(rows, key=lambda row: json.dumps(row, sort_keys=True, default=str))
    for row in ordered:
        sheet.append([_excel_text(row.get(header, "")) for header in headers])
    # WAPI names/comments/options are literal data, including leading '='.
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    sheet.freeze_panes = "A2"
    for column in sheet.columns:
        width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 10), 45)
        sheet.column_dimensions[column[0].column_letter].width = width
    if rows:
        table = Table(displayName=f"InventoryTable{index}", ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        # The table owns its AutoFilter. Adding a worksheet AutoFilter over
        # this range causes Microsoft Excel to repair/reject the table parts.
        sheet.add_table(table)
    # An empty inventory keeps its headers but has neither a table nor filter.
    return sheet


def _section_title(sheet, row: int, start: int, end: int, title: str) -> None:
    sheet.merge_cells(start_row=row, start_column=start, end_row=row, end_column=end)
    cell = sheet.cell(row, start, title)
    cell.fill = PatternFill('solid', fgColor='17365D')
    cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    cell.alignment = Alignment(vertical='center')
    sheet.row_dimensions[row].height = 24


def _write_template_overview_dashboard(
    workbook: Workbook,
    results: list[CollectionResult],
    coverage: list[dict[str, Any]],
    template_models: list[dict[str, Any]],
    template_bundles: list[dict[str, Any]],
    template_bundle_models: list[dict[str, Any]],
) -> dict[int, str]:
    """Template-first landing page for discovery across multiple Grids."""
    sheet = workbook.create_sheet('Overview')
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 90
    sheet.freeze_panes = 'A2'
    widths = {'A': 24, 'B': 18, 'C': 15, 'D': 13, 'E': 15, 'F': 18, 'G': 24, 'H': 16, 'I': 34}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    sheet.merge_cells('A1:I1')
    sheet['A1'] = 'Infoblox DHCP Template & Current-State Assessment'
    sheet['A1'].fill = PatternFill('solid', fgColor='0F243E')
    sheet['A1'].font = Font(name='Calibri', size=18, bold=True, color='FFFFFF')
    sheet['A1'].alignment = Alignment(vertical='center')
    sheet.row_dimensions[1].height = 34
    sheet.merge_cells('A2:I2')
    sheet['A2'] = (
        'Start with reusable provisioning models. Grid-local IPs, domains and server names remain '
        'parameters; policy and construction differences define separate models.'
    )
    sheet['A2'].font = Font(name='Calibri', size=10, italic=True, color='666666')
    sheet['A2'].alignment = Alignment(wrap_text=True, vertical='top')
    sheet.row_dimensions[2].height = 30

    def object_count(*types: str) -> int:
        return sum(len(result.records.get(kind, [])) for result in results for kind in types)

    thin = Side(style='thin', color='D9E2F3')
    bundle_status = Counter(str(row.get('Bundle Status') or '') for row in template_bundles)
    bundle_family = Counter(str(row.get('Provisioning Family') or '') for row in template_bundles)
    model_attention = Counter(str(row.get('Attention') or '') for row in template_bundle_models)
    bundle_attention = Counter(str(row.get('Attention') or '') for row in template_bundles)

    _section_title(sheet, 4, 1, 4, 'Environment')
    _section_title(sheet, 4, 5, 9, 'Template discovery')
    environment = [
        ('Grids', len(results), 'Networks', object_count('network')),
        ('Templates', object_count('networktemplate', 'rangetemplate', 'fixedaddresstemplate'),
         'Ranges', object_count('range')),
        ('Network Templates', object_count('networktemplate'),
         'Range Templates', object_count('rangetemplate')),
        ('Template semantic models', len(template_models),
         'Provisioning bundles', len(template_bundles)),
        ('Bundle models', len(template_bundle_models),
         'Network Views', object_count('networkview')),
    ]
    discovery = [
        ('Linked bundles', bundle_status['LINKED'], 'Orphan ranges', bundle_status['ORPHAN_RANGE_TEMPLATE']),
        ('Models needing review', model_attention['REVIEW'], 'Bundles needing review', bundle_attention['REVIEW']),
        ('MS_SERVER bundles', bundle_family['MS_SERVER'], 'INFOBLOX bundles', bundle_family['INFOBLOX']),
        ('Cross-Grid models', sum((row.get('Grid Count') or 0) > 1 for row in template_bundle_models),
         'Collected Grids', len(results)),
        ('Missing linked ranges', bundle_status['MISSING_RANGE_TEMPLATE'],
         'Network-only bundles', bundle_status['NETWORK_ONLY']),
    ]
    for offset, row_values in enumerate(environment, start=5):
        for col, value in zip((1, 2, 3, 4), row_values):
            cell = sheet.cell(offset, col, value)
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            cell.font = Font(name='Calibri', size=11, bold=col in {2, 4})
            if col in {2, 4}:
                cell.fill = PatternFill('solid', fgColor='EAF2F8')
    for offset, row_values in enumerate(discovery, start=5):
        for col, value in zip((5, 6, 7, 8), row_values):
            cell = sheet.cell(offset, col, value)
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            cell.font = Font(name='Calibri', size=11, bold=col in {6, 8})
            if col in {6, 8}:
                cell.fill = PatternFill('solid', fgColor='EAF2F8')

    sheet['B9'].hyperlink = Hyperlink(ref='B9', location="'Template_Bundle_Models'!A1")
    sheet['B9'].font = Font(name='Calibri', size=11, bold=True, color='0563C1', underline='single')

    _section_title(sheet, 11, 1, 9, 'Provisioning model catalog')
    headers = [
        'Bundle Model', 'Family', 'Bundles', 'Grids', 'Grid coverage %',
        'Prefixes', 'Construction pattern', 'Attention', 'Templates',
    ]
    for column, label in enumerate(headers, start=1):
        cell = sheet.cell(12, column, label)
        cell.fill = PatternFill('solid', fgColor='D9EAF7')
        cell.font = Font(name='Calibri', size=10, bold=True, color='17365D')
        cell.alignment = Alignment(wrap_text=True, vertical='center')

    ordered_models = sorted(
        template_bundle_models,
        key=lambda row: (
            -int(row.get('Grid Count') or 0),
            -int(row.get('Bundle Count') or 0),
            str(row.get('Bundle Model ID') or ''),
        ),
    )
    for row_number, item in enumerate(ordered_models[:12], start=13):
        templates = item.get('Network Templates') or item.get('Range Templates') or ''
        values = [
            item.get('Bundle Model ID'), item.get('Provisioning Family'),
            item.get('Bundle Count'), item.get('Grid Count'), item.get('Grid Coverage %'),
            item.get('Network Prefixes'), item.get('Pattern Summary'), item.get('Attention'),
            templates,
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_number, column, _excel_text(value))
            if isinstance(cell.value, str):
                cell.data_type = 's'
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            cell.border = Border(bottom=Side(style='hair', color='E7E6E6'))
        sheet.cell(row_number, 1).hyperlink = Hyperlink(
            ref=sheet.cell(row_number, 1).coordinate, location="'Template_Bundle_Models'!A1"
        )
        sheet.cell(row_number, 1).font = Font(name='Calibri', size=10, color='0563C1', underline='single')
        sheet.cell(row_number, 5).number_format = '0.0"%"'
        attention_cell = sheet.cell(row_number, 8)
        if attention_cell.value == 'REVIEW':
            attention_cell.fill = PatternFill('solid', fgColor='FFF2CC')
            attention_cell.font = Font(name='Calibri', size=10, bold=True, color='9C6500')
        elif attention_cell.value == 'OK':
            attention_cell.fill = PatternFill('solid', fgColor='E2F0D9')

    catalog_end = 12 + min(len(ordered_models), 12)
    if ordered_models:
        table = Table(displayName='OverviewBundleModels', ref=f'A12:I{catalog_end}')
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        sheet.add_table(table)

    nav_row = catalog_end + 2
    _section_title(sheet, nav_row, 1, 9, 'Navigate')
    navigation = [
        ('Bundle model catalog', "'Template_Bundle_Models'!A1"),
        ('Bundle assignments', "'Template_Bundles'!A1"),
        ('Grid-local parameters', "'Template_Grid_Matrix'!A1"),
        ('All sheets / evidence index', "'Data_Index'!A1"),
    ]
    for offset, (label, location) in enumerate(navigation, start=nav_row + 1):
        cell = sheet.cell(offset, 1, label)
        cell.hyperlink = Hyperlink(ref=cell.coordinate, location=location)
        cell.font = Font(name='Calibri', size=10, color='0563C1', underline='single')
        cell.alignment = Alignment(vertical='center')
    coverage_counts = Counter(str(row.get('Collection Status', 'PARTIAL')) for row in coverage)
    collection_row = nav_row + 1
    sheet.cell(collection_row, 5, 'Collection errors')
    sheet.cell(collection_row, 6, sum(len(result.errors) for result in results))
    sheet.cell(collection_row + 1, 5, 'Complete observations')
    sheet.cell(collection_row + 1, 6, coverage_counts['COMPLETE'])
    sheet.cell(collection_row + 2, 5, 'Partial observations')
    sheet.cell(collection_row + 2, 6, coverage_counts['PARTIAL'])
    sheet.cell(collection_row + 3, 5, 'Manual review observations')
    sheet.cell(collection_row + 3, 6, coverage_counts['MANUAL_REVIEW_REQUIRED'])
    for row_number in range(collection_row, collection_row + 4):
        sheet.cell(row_number, 5).border = Border(bottom=thin)
        sheet.cell(row_number, 6).border = Border(bottom=thin)
        sheet.cell(row_number, 6).font = Font(name='Calibri', size=10, bold=True)

    note_row = nav_row + 6
    sheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=9)
    sheet.cell(
        note_row, 1,
        'Model prevalence is descriptive evidence, not an approved standard. Review flags identify '
        'construction inconsistencies or unresolved relationships; disabled option containers remain '
        'visible as option state, not as automatic defects.'
    )
    sheet.cell(note_row, 1).fill = PatternFill('solid', fgColor='EAF2F8')
    sheet.cell(note_row, 1).font = Font(name='Calibri', size=10, italic=True, color='17365D')
    sheet.cell(note_row, 1).alignment = Alignment(wrap_text=True, vertical='center')
    sheet.row_dimensions[note_row].height = 34

    sheet.print_area = f'A1:I{note_row}'
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    return {}


def _write_overview_dashboard(workbook: Workbook, results: list[CollectionResult],
                              coverage: list[dict[str, Any]], standardization: list[dict[str, Any]],
                              exceptions: list[dict[str, Any]],
                              template_models: list[dict[str, Any]] | None = None,
                              template_bundles: list[dict[str, Any]] | None = None,
                              template_bundle_models: list[dict[str, Any]] | None = None) -> dict[int, str]:
    """Create a workshop-oriented landing page; technical evidence stays on later sheets."""
    if template_bundle_models:
        return _write_template_overview_dashboard(
            workbook, results, coverage, template_models or [],
            template_bundles or [], template_bundle_models,
        )
    sheet = workbook.create_sheet('Overview')
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 90
    sheet.freeze_panes = 'A2'
    widths = {'A': 28, 'B': 28, 'C': 14, 'D': 12, 'E': 22, 'F': 18, 'G': 18, 'H': 14, 'I': 18}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    sheet.merge_cells('A1:I1')
    sheet['A1'] = 'Infoblox Current State & Standardization Assessment'
    sheet['A1'].fill = PatternFill('solid', fgColor='0F243E')
    sheet['A1'].font = Font(name='Calibri', size=18, bold=True, color='FFFFFF')
    sheet['A1'].alignment = Alignment(vertical='center')
    sheet.row_dimensions[1].height = 34
    sheet.merge_cells('A2:I2')
    sheet['A2'] = ('Observed configuration is evidence. Common values are candidates for review; '
                   'only an explicit APPROVED decision defines a target standard.')
    sheet['A2'].font = Font(name='Calibri', size=10, italic=True, color='666666')
    sheet['A2'].alignment = Alignment(wrap_text=True, vertical='top')
    sheet.row_dimensions[2].height = 30

    def object_count(*types: str) -> int:
        return sum(len(result.records.get(kind, [])) for result in results for kind in types)

    _section_title(sheet, 4, 1, 4, 'Environment')
    _section_title(sheet, 4, 5, 9, 'Standardization')
    environment = [
        ('Grids', len(results), 'Networks', object_count('network')),
        ('Members', object_count('member'), 'Ranges', object_count('range')),
        ('Network Views', object_count('networkview'), 'Reservations', object_count('fixedaddress')),
        ('Templates', object_count('networktemplate', 'rangetemplate', 'fixedaddresstemplate'),
         'Filters', object_count('filtermac', 'filteroption', 'filterrelayagent', 'filterfingerprint', 'filternac')),
        ('Template models', len(template_models or []), 'Bundle models',
         len(template_bundle_models or [])),
    ]
    classifications = Counter(str(row.get('Consistency Classification', '')) for row in standardization)
    decisions = Counter(str(row.get('Decision Status', 'PENDING')) for row in standardization)
    blocked = sum(row.get('Standardization Candidate') == 'BLOCKED_BY_DATA' for row in standardization)
    local_parameters = sum((row.get('Local Override Count') or 0) > 0 for row in standardization)
    standard = [
        ('Standardization questions', len(standardization), 'With confirmed evidence', sum((row.get('Confirmed Objects') or 0) > 0 for row in standardization)),
        ('Multiple observed values', classifications['MULTIPLE_VALUES'], 'Parameters with local overrides', local_parameters),
        ('Pending decisions', decisions['PENDING'], 'Approved targets', decisions['APPROVED']),
        ('Blocked by data', blocked, 'Exception / review rows', len(exceptions)),
    ]
    thin = Side(style='thin', color='D9E2F3')
    for offset, (left_label, left_value, right_label, right_value) in enumerate(environment, start=5):
        for col, value in ((1, left_label), (2, left_value), (3, right_label), (4, right_value)):
            cell = sheet.cell(offset, col, value)
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            cell.font = Font(name='Calibri', size=11, bold=col in {2, 4})
            if col in {2, 4}:
                cell.fill = PatternFill('solid', fgColor='EAF2F8')
    if template_models:
        model_cell = sheet.cell(9, 2)
        model_cell.hyperlink = Hyperlink(ref=model_cell.coordinate, location="'Template_Models'!A1")
        model_cell.font = Font(name='Calibri', size=11, bold=True, color='0563C1', underline='single')
    if template_bundles:
        bundle_cell = sheet.cell(9, 4)
        bundle_cell.hyperlink = Hyperlink(ref=bundle_cell.coordinate, location="'Template_Bundle_Models'!A1")
        bundle_cell.font = Font(name='Calibri', size=11, bold=True, color='0563C1', underline='single')
    for offset, (left_label, left_value, right_label, right_value) in enumerate(standard, start=5):
        for col, value in ((5, left_label), (6, left_value), (7, right_label), (8, right_value)):
            cell = sheet.cell(offset, col, value)
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            cell.font = Font(name='Calibri', size=11, bold=col in {6, 8})
            if col in {6, 8}:
                cell.fill = PatternFill('solid', fgColor='EAF2F8')

    _section_title(sheet, 10, 1, 4, 'Assessment confidence / coverage')
    _section_title(sheet, 10, 5, 9, 'Collection')
    coverage_counts = Counter(str(row.get('Collection Status', 'PARTIAL')) for row in coverage)
    coverage_rows = [
        ('COMPLETE', coverage_counts['COMPLETE'], 'PARTIAL', coverage_counts['PARTIAL']),
        ('EMPTY', coverage_counts['EMPTY'], 'NOT CONFIGURED', coverage_counts['NOT_CONFIGURED']),
        ('NOT EXPOSED', coverage_counts['NOT_EXPOSED_BY_WAPI'], 'MANUAL REVIEW', coverage_counts['MANUAL_REVIEW_REQUIRED']),
        ('ERROR', coverage_counts['ERROR'], '', ''),
    ]
    sources = [row for result in results for row in result.collection_sources]
    timestamps = sorted({str(row.get('Collected At')) for row in sources if row.get('Collected At')})
    archive_ids = {str(row.get('Archive ID')) for row in sources if row.get('Archive ID')}
    versions = sorted({str(result.wapi_version) for result in results if result.wapi_version})
    collection_rows = [
        ('Earliest collection', timestamps[0] if timestamps else 'Unavailable', 'Latest collection', timestamps[-1] if timestamps else 'Unavailable'),
        ('Archives combined', len(archive_ids) if archive_ids else ('PARTIAL' if sources else 0), 'WAPI version(s)', ', '.join(versions) or 'Unavailable'),
        ('Collection errors', sum(len(result.errors) for result in results), 'Manual review items', coverage_counts['MANUAL_REVIEW_REQUIRED']),
        ('Reference definitions', object_count('dhcpoptiondefinition', 'dhcpoptionspace', 'extensibleattributedef'), '', ''),
    ]
    for offset, row_values in enumerate(coverage_rows, start=11):
        for col, value in zip((1, 2, 3, 4), row_values):
            cell = sheet.cell(offset, col, value)
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            if col in {2, 4}:
                cell.font = Font(name='Calibri', size=11, bold=True)
        status = str(sheet.cell(offset, 1).value or '')
        if status == 'ERROR' and sheet.cell(offset, 2).value:
            sheet.cell(offset, 2).fill = PatternFill('solid', fgColor='F4CCCC')
        elif status in {'PARTIAL', 'NOT EXPOSED'}:
            sheet.cell(offset, 2).fill = PatternFill('solid', fgColor='FFF2CC')
    for offset, row_values in enumerate(collection_rows, start=11):
        for col, value in zip((5, 6, 7, 8), row_values):
            cell = sheet.cell(offset, col, value)
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            if col in {6, 8}:
                cell.font = Font(name='Calibri', size=10, bold=True)

    _section_title(sheet, 16, 1, 9, 'Top standardization hotspots')
    hotspot_headers = ['Parameter', 'Observed values', 'Confirmed', 'Grids', 'Consistency', 'Local overrides', 'Override %', 'Coverage', 'Decision']
    for col, header in enumerate(hotspot_headers, start=1):
        cell = sheet.cell(17, col, header)
        cell.fill = PatternFill('solid', fgColor='D9EAF7')
        cell.font = Font(name='Calibri', size=10, bold=True, color='17365D')
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    priority = {'MULTIPLE_VALUES': 0, 'ONLY_IN_SOME_GRIDS': 1, 'DIFFERENT_SOURCE': 2,
                'LOCAL_OVERRIDES': 3, 'INSUFFICIENT_DATA': 4, 'NOT_CONFIGURED': 5,
                'CONSISTENT': 6, 'NO_OBJECTS_IN_SCOPE': 7}
    ordered = sorted(standardization, key=lambda row: (
        priority.get(str(row.get('Consistency Classification')), 9),
        -(row.get('Local Override Count') or 0), str(row.get('Parameter'))))
    meaningful = [row for row in ordered if (row.get('Confirmed Objects') or 0) > 0
                  and row.get('Consistency Classification') != 'CONSISTENT']
    blocked_rows = [row for row in ordered if row.get('Standardization Candidate') == 'BLOCKED_BY_DATA']
    hotspots = (meaningful[:8] + [row for row in blocked_rows if row not in meaningful][:2])
    if not hotspots:
        hotspots = [row for row in ordered if (row.get('Confirmed Objects') or 0) > 0][:8]
    link_targets: dict[int, str] = {}
    for idx, item in enumerate(hotspots, start=18):
        parameter_label = f"{item.get('Parameter')} [{item.get('Scope')}]"
        confirmed = f"{item.get('Confirmed Objects') or 0}/{item.get('Population Objects') or 0}"
        query_coverage = item.get('Query Coverage %')
        resolved_coverage = item.get('Resolved Evidence %')
        coverage_parts = [str(item.get('Coverage') or '')]
        if query_coverage is not None:
            coverage_parts.append(f"collection {query_coverage:.1f}%")
        if resolved_coverage is not None:
            coverage_parts.append(f"evidence {resolved_coverage:.1f}%")
        coverage_display = " | ".join(part for part in coverage_parts if part)
        values = [parameter_label, item.get('Observed Values') or item.get('Evidence Status'),
                  confirmed, item.get('Grids Assessed'), item.get('Consistency Classification'),
                  item.get('Local Override Count'), item.get('Local Override %'), coverage_display,
                  item.get('Decision Status')]
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(idx, col, value)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.border = Border(bottom=Side(style='hair', color='E7E6E6'))
        parameter_id = str(item.get('Parameter ID') or '')
        if parameter_id:
            link_targets[idx] = parameter_id
        sheet.cell(idx, 1).font = Font(name='Calibri', size=10, color='0563C1', underline='single')
        if item.get('Local Override %') is not None:
            sheet.cell(idx, 7).number_format = '0.0"%"'
        consistency = str(item.get('Consistency Classification'))
        fill = 'E2F0D9' if consistency == 'CONSISTENT' else 'FFF2CC' if consistency != 'INSUFFICIENT_DATA' else 'E7E6E6'
        sheet.cell(idx, 5).fill = PatternFill('solid', fgColor=fill)
    if hotspots:
        table = Table(displayName='OverviewHotspots', ref=f'A17:I{17 + len(hotspots)}')
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        sheet.add_table(table)

    note_row = 19 + len(hotspots)
    sheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=10)
    sheet.cell(note_row, 1, 'Workflow: review Standardization → record human decisions in YAML/Decisions → use Exceptions as the remediation/exception backlog.')
    sheet.cell(note_row, 1).fill = PatternFill('solid', fgColor='EAF2F8')
    sheet.cell(note_row, 1).font = Font(name='Calibri', size=10, italic=True, color='17365D')
    sheet.cell(note_row, 1).alignment = Alignment(wrap_text=True, vertical='center')
    sheet.row_dimensions[note_row].height = 30
    sheet.print_area = f'A1:J{note_row}'
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    return link_targets


def _wire_overview_standardization_links(workbook: Workbook, link_targets: dict[int, str]) -> None:
    """Resolve Overview drill-down links against the final Standardization row order."""
    if not link_targets:
        return
    overview = workbook['Overview']
    standardization = workbook['Standardization']
    header_map = {cell.value: cell.column for cell in standardization[1]}
    parameter_column = header_map.get('Parameter ID')
    if parameter_column is None:
        raise ValueError("Standardization sheet is missing the 'Parameter ID' column")
    target_rows: dict[str, int] = {}
    for row_number in range(2, standardization.max_row + 1):
        value = standardization.cell(row_number, parameter_column).value
        if value in (None, ''):
            continue
        parameter_id = str(value)
        if parameter_id in target_rows:
            raise ValueError(f"Duplicate Parameter ID in Standardization sheet: {parameter_id}")
        target_rows[parameter_id] = row_number
    missing = sorted({parameter_id for parameter_id in link_targets.values()
                      if parameter_id not in target_rows})
    if missing:
        raise ValueError("Overview hotspot Parameter ID missing from Standardization sheet: " + ", ".join(missing))
    for overview_row, parameter_id in link_targets.items():
        cell = overview.cell(overview_row, 1)
        cell.hyperlink = Hyperlink(ref=cell.coordinate,
                                   location=f"'Standardization'!A{target_rows[parameter_id]}")


def _write_overview_profile_summary(workbook: Workbook, summaries: list[dict[str, Any]],
                                    kpis: list[dict[str, Any]] | None = None,
                                    relationships: list[dict[str, Any]] | None = None) -> None:
    """Append input readiness without moving existing evidence or hotspot links."""
    sheet = workbook['Overview']
    matrix = workbook['Profile_Readiness']
    profile_column = next(cell.column for cell in matrix[1] if cell.value == 'Profile')
    targets: dict[str, int] = {}
    for row_number in range(2, matrix.max_row + 1):
        profile = matrix.cell(row_number, profile_column).value
        if profile:
            targets.setdefault(str(profile), row_number)
    start = sheet.max_row + 2
    _section_title(sheet, start, 1, 10, 'Profile discovery readiness')
    for column, label in enumerate([*PROFILE_SUMMARY_HEADERS, 'Input details'], start=1):
        cell = sheet.cell(start + 1, column, label)
        cell.fill = PatternFill('solid', fgColor='D9EAF7')
        cell.font = Font(name='Calibri', size=10, bold=True, color='17365D')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
    sheet.row_dimensions[start + 1].height = 30
    for row_number, summary in enumerate(summaries, start=start + 2):
        for column, header in enumerate(PROFILE_SUMMARY_HEADERS, start=1):
            cell = sheet.cell(row_number, column, _excel_text(summary.get(header)))
            if isinstance(cell.value, str):
                cell.data_type = 's'
            cell.alignment = Alignment(wrap_text=True, vertical='center')
            cell.border = Border(bottom=Side(style='thin', color='D9E2F3'))
            if header == 'Ready Input %':
                cell.number_format = '0.0"%"'
        target = targets[str(summary['Profile'])]
        link_cell = sheet.cell(row_number, 10, 'View inputs')
        link_cell.hyperlink = Hyperlink(ref=link_cell.coordinate, location=f"'Profile_Readiness'!A{target}")
        link_cell.font = Font(name='Calibri', size=10, color='0563C1', underline='single')
        link_cell.alignment = Alignment(wrap_text=True, vertical='center')
        link_cell.border = Border(bottom=Side(style='thin', color='D9E2F3'))
    note_row = start + len(summaries) + 2
    sheet.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=9)
    cell = sheet.cell(note_row, 1,
                      'Ready Input % = READY inputs / applicable candidate inputs; DEFERRED semantic inputs are excluded; '
                      'this is not object coverage. Network readiness uses DHCP-relevant candidate Networks; '
                      'Range readiness uses parameter-applicable association populations. Standardization hotspots above keep '
                      'their full scoped populations, so those denominators are not directly comparable. '
                      'READY means usable evidence for future profile discovery, not an approved standard or configuration compliance.')
    cell.font = Font(name='Calibri', size=10, italic=True, color='666666')
    cell.alignment = Alignment(wrap_text=True, vertical='center')
    sheet.row_dimensions[note_row].height = 30

    end_row = note_row
    if kpis:
        kpi_start = note_row + 2
        _section_title(sheet, kpi_start, 1, 9, 'Observed profile discovery')
        labels = [
            'Profile', 'Applicable', 'Profiled', 'Profiled %', 'Distinct profiles',
            'Top-1 share of profiled %', 'Top-3 share of profiled %', 'Singleton profiles', 'Profile details',
        ]
        for column, label in enumerate(labels, start=1):
            cell = sheet.cell(kpi_start + 1, column, label)
            cell.fill = PatternFill('solid', fgColor='D9EAF7')
            cell.font = Font(name='Calibri', size=10, bold=True, color='17365D')
            cell.alignment = Alignment(wrap_text=True, vertical='center')

        kpi_sheet = workbook['Profile_KPIs']
        kpi_profile_column = next(cell.column for cell in kpi_sheet[1] if cell.value == 'Profile')
        kpi_targets = {
            str(kpi_sheet.cell(row, kpi_profile_column).value): row
            for row in range(2, kpi_sheet.max_row + 1)
            if kpi_sheet.cell(row, kpi_profile_column).value
        }
        for row_number, item in enumerate(kpis, start=kpi_start + 2):
            values = [
                item.get('Profile'), item.get('Applicable Objects'), item.get('Profiled Objects'),
                item.get('Profiled %'), item.get('Distinct Profiles'), item.get('Top-1 Share of Profiled %'),
                item.get('Top-3 Share of Profiled %'), item.get('Singleton Profiles'),
            ]
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row_number, column, _excel_text(value))
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.alignment = Alignment(wrap_text=True, vertical='center')
                cell.border = Border(bottom=Side(style='thin', color='D9E2F3'))
                if column in {4, 6, 7}:
                    cell.number_format = '0.0"%"'
            target = kpi_targets.get(str(item.get('Profile')))
            link_cell = sheet.cell(row_number, 9, 'View profile KPIs')
            if target is not None:
                link_cell.hyperlink = Hyperlink(ref=link_cell.coordinate, location=f"'Profile_KPIs'!A{target}")
                link_cell.font = Font(name='Calibri', size=10, color='0563C1', underline='single')
            link_cell.alignment = Alignment(wrap_text=True, vertical='center')
            link_cell.border = Border(bottom=Side(style='thin', color='D9E2F3'))

        discovery_note = kpi_start + len(kpis) + 2
        sheet.merge_cells(start_row=discovery_note, start_column=1, end_row=discovery_note, end_column=9)
        cell = sheet.cell(
            discovery_note, 1,
            'Profile prevalence is descriptive current-state evidence. The most common fingerprint is not an approved standard; '
            'objects with unresolved core inputs remain outside normal profiles.'
        )
        cell.font = Font(name='Calibri', size=10, italic=True, color='666666')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        sheet.row_dimensions[discovery_note].height = 30
        end_row = discovery_note

    if relationships:
        relationship_start = end_row + 2
        _section_title(sheet, relationship_start, 1, 9, 'Observed Network↔Range profile relationships')
        paired_rows = [row for row in relationships if row.get('Relationship Status') == 'PROFILE_PAIR']
        associated = sum(int(row.get('Range Count') or 0) for row in relationships)
        paired = sum(int(row.get('Range Count') or 0) for row in paired_rows)
        top_pair = sorted(
            paired_rows,
            key=lambda row: (-int(row.get('Range Count') or 0),
                             str(row.get('Network Fingerprint ID') or ''),
                             str(row.get('Range Fingerprint ID') or '')),
        )[0] if paired_rows else {}
        values = [
            ['Associated ranges', associated, 'Paired profiled ranges', paired,
             'Paired %', round(paired * 100.0 / associated, 1) if associated else None,
             'Distinct profiled pairs', len(paired_rows), 'Relationship details'],
            ['Top pair count', top_pair.get('Range Count'),
             'Top pair % associated', top_pair.get('Share of Associated Ranges %'),
             'Top pair % paired', top_pair.get('Share of Paired Profiled Ranges %'),
             'Association family', top_pair.get('Association Family'), 'View relationships'],
        ]
        for offset, row_values in enumerate(values, start=1):
            row_number = relationship_start + offset
            for column, value in enumerate(row_values, start=1):
                cell = sheet.cell(row_number, column, _excel_text(value))
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.alignment = Alignment(wrap_text=True, vertical='center')
                cell.border = Border(bottom=Side(style='thin', color='D9E2F3'))
            for column in (6,):
                sheet.cell(row_number, column).number_format = '0.0"%"'
        sheet.cell(relationship_start + 2, 4).number_format = '0.0"%"'
        link_cell = sheet.cell(relationship_start + 1, 9)
        link_cell.value = 'View relationships'
        link_cell.hyperlink = Hyperlink(ref=link_cell.coordinate, location="'Profile_Relationships'!A2")
        link_cell.font = Font(name='Calibri', size=10, color='0563C1', underline='single')
        link_cell2 = sheet.cell(relationship_start + 2, 9)
        link_cell2.value = 'View relationships'
        link_cell2.hyperlink = Hyperlink(ref=link_cell2.coordinate, location="'Profile_Relationships'!A2")
        link_cell2.font = Font(name='Calibri', size=10, color='0563C1', underline='single')

        relationship_note = relationship_start + 4
        sheet.merge_cells(start_row=relationship_note, start_column=1, end_row=relationship_note, end_column=9)
        cell = sheet.cell(
            relationship_note, 1,
            'A recurring Network↔Range pair is descriptive topology/context evidence only. '
            'It does not make either fingerprint an approved standard.'
        )
        cell.font = Font(name='Calibri', size=10, italic=True, color='666666')
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        sheet.row_dimensions[relationship_note].height = 30
        end_row = relationship_note

    sheet.print_area = f'A1:I{end_row}'


def _style_decision_support(sheet, name: str) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 85
    sheet.freeze_panes = 'A2'
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='17365D')
        cell.font = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    input_headers = {
        'Decisions': {'Proposed / Discussed Target', 'Approved Target', 'Status', 'Exceptions Allowed',
                      'Exception Rule', 'Owner', 'Decision Date', 'Notes'},
        'Standardization': {'Decision Status', 'Proposed / Discussed Target', 'Approved Target', 'Exceptions Allowed',
                            'Exception Rule', 'Decision Owner', 'Decision Date', 'Decision Notes'},
    }.get(name, set())
    header_map = {cell.value: cell.column for cell in sheet[1]}
    for header in input_headers:
        column = header_map.get(header)
        if column:
            for cell in [sheet.cell(row, column) for row in range(2, sheet.max_row + 1)]:
                cell.fill = PatternFill('solid', fgColor='FFF2CC')
    if name == 'Exceptions':
        state_col = header_map.get('Assessment State')
        if state_col:
            fills = {'DEVIATION': 'F4CCCC', 'PENDING_DECISION': 'FFF2CC', 'APPROVED_EXCEPTION': 'E2F0D9',
                     'INSUFFICIENT_DATA': 'E7E6E6'}
            for row in range(2, sheet.max_row + 1):
                value = str(sheet.cell(row, state_col).value or '')
                if value in fills:
                    sheet.cell(row, state_col).fill = PatternFill('solid', fgColor=fills[value])
    if name == 'Profile_Readiness':
        fills = {'READY': 'E2F0D9', 'CONDITIONAL': 'FFF2CC', 'NOT_READY': 'F4CCCC',
                 'DEFERRED': 'D9EAF7', 'NOT_APPLICABLE': 'E7E6E6'}
        for row in range(2, sheet.max_row + 1):
            cell = sheet.cell(row, header_map['Readiness'])
            cell.fill = PatternFill('solid', fgColor=fills[str(cell.value)])
            for header, column in header_map.items():
                if header.endswith('%'):
                    sheet.cell(row, column).number_format = '0.0"%"'
    if name == 'Profile_Assignments':
        fills = {
            'PROFILED': 'E2F0D9',
            'UNRESOLVED_PROFILE_INPUTS': 'FFF2CC',
            'UNRESOLVED_ASSOCIATION': 'FFF2CC',
            'NOT_APPLICABLE_TO_DHCP_PROFILE': 'E7E6E6',
        }
        status_column = header_map.get('Population Status')
        if status_column:
            for row in range(2, sheet.max_row + 1):
                cell = sheet.cell(row, status_column)
                if str(cell.value) in fills:
                    cell.fill = PatternFill('solid', fgColor=fills[str(cell.value)])
    if name == 'Profile_Dimension_Assignments':
        fills = {
            'PROFILED': 'E2F0D9',
            'UNRESOLVED_DIMENSION_INPUTS': 'FFF2CC',
            'UNRESOLVED_ASSOCIATION': 'FFF2CC',
            'NOT_APPLICABLE_TO_DHCP_PROFILE': 'E7E6E6',
        }
        status_column = header_map.get('Population Status')
        if status_column:
            for row in range(2, sheet.max_row + 1):
                cell = sheet.cell(row, status_column)
                if str(cell.value) in fills:
                    cell.fill = PatternFill('solid', fgColor=fills[str(cell.value)])
    if name == 'Profile_Compositions':
        fills = {
            'COMPLETE': 'E2F0D9',
            'PARTIAL': 'FFF2CC',
            'UNRESOLVED_ASSOCIATION': 'FFF2CC',
            'NOT_APPLICABLE_TO_DHCP_PROFILE': 'E7E6E6',
        }
        status_column = header_map.get('Composition Status')
        if status_column:
            for row in range(2, sheet.max_row + 1):
                cell = sheet.cell(row, status_column)
                if str(cell.value) in fills:
                    cell.fill = PatternFill('solid', fgColor=fills[str(cell.value)])
    if name == 'Profile_Usefulness':
        role_column = header_map.get('Fingerprint Role')
        role_fills = {'CORE': 'E2F0D9', 'DERIVED_CORE': 'D9EAF7', 'OVERLAY': 'FFF2CC', 'DEFERRED': 'E7E6E6'}
        if role_column:
            for row in range(2, sheet.max_row + 1):
                cell = sheet.cell(row, role_column)
                if str(cell.value) in role_fills:
                    cell.fill = PatternFill('solid', fgColor=role_fills[str(cell.value)])
    if name in {'Profile_Relationships', 'Profile_Relationship_Assignments'}:
        status_column = header_map.get('Relationship Status')
        relationship_fills = {
            'PROFILE_PAIR': 'E2F0D9',
            'PARENT_PROFILE_UNRESOLVED': 'FFF2CC',
            'PARENT_NOT_PROFILE_CANDIDATE': 'FFF2CC',
            'PARENT_NETWORK_NOT_FOUND': 'F4CCCC',
            'RANGE_PROFILE_UNRESOLVED': 'FFF2CC',
            'RANGE_ASSOCIATION_UNRESOLVED': 'FFF2CC',
            'RANGE_NOT_APPLICABLE': 'E7E6E6',
        }
        if status_column:
            for row in range(2, sheet.max_row + 1):
                cell = sheet.cell(row, status_column)
                if str(cell.value) in relationship_fills:
                    cell.fill = PatternFill('solid', fgColor=relationship_fills[str(cell.value)])
    if name in {
        'Profile_Readiness', 'Profile_Usefulness', 'Profile_Fingerprints',
        'Profile_Assignments', 'Profile_KPIs', 'Profile_Dimensions',
        'Profile_Dimension_Assignments', 'Profile_Dimension_KPIs',
        'Profile_Compositions', 'Profile_Context', 'Profile_Relationships',
        'Profile_Relationship_Assignments', 'Standardization', 'Decisions',
        'Exceptions', 'Grid_Comparison',
    }:
        for column in sheet.columns:
            for cell in column:
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for header, column in header_map.items():
            if isinstance(header, str) and header.endswith('%'):
                for row in range(2, sheet.max_row + 1):
                    sheet.cell(row, column).number_format = '0.0"%"'
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True

def write_reports(results: list[CollectionResult], output_dir: str | Path, *, decisions_path: str | Path | None = None) -> None:
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    all_options: list[dict[str, Any]] = []
    all_scalars: list[dict[str, Any]] = []
    naming: list[dict[str, Any]] = []
    topology_records = []
    reservation_records = []
    stored_only = set(TOPOLOGY_SHEETS) | (set(RESERVATION_SHEETS) - {"fixedaddress"})
    results = combine_collections(results)
    for result in results:
        # Stored reusable configuration has its own rows, without effective-value claims.
        arguments = (result.grid, result.grid_url, result.wapi_version,
                     {key: rows for key, rows in result.records.items() if key not in stored_only},
                     {key: rows for key, rows in result.effective_records.items() if key not in stored_only},
                     result.schemas)
        all_options.extend(normalize_options(*arguments))
        all_scalars.extend(normalize_scalars(*arguments))
        naming.extend({**row, "grid": result.grid} for row in naming_analysis(result.records))
        topology_records.extend(normalize_topology(result))
        reservation_records.extend(normalize_reservations(result))
    coverage = ([row for result in results for row in result.coverage]
                + topology_coverage(topology_records) + reservation_coverage(reservation_records))
    for result in results:
        filter_count = sum(len(rows) for key, rows in result.records.items()
                           if key in RESERVATION_SHEETS and (key.startswith("filter") or key == "macfilteraddress"))
        if filter_count:
            coverage.append({"Grid": result.grid, "Area": "DHCP filters and MAC filters", "Object": "",
                             "Query": "policy_interpretation", "Collection Status": "MANUAL_REVIEW_REQUIRED",
                             "Objects Found": filter_count,
                             "Notes": "Configured names, expressions and references do not establish deployed permit/deny policy; "
                                      "consumer assignments outside this increment are not assessed."})
    decisions = load_decisions(decisions_path)
    standardization = build_standardization(results, coverage, all_scalars, all_options, decisions) if results else []
    profile_readiness = build_profile_readiness(
        standardization, results=results, scalars=all_scalars, options=all_options
    )
    profile_summaries = profile_readiness_summary(profile_readiness)
    profile_populations = profile_population_rows(results)
    profile_usefulness, profile_fingerprints, profile_assignments, profile_kpis = (
        build_profile_discovery(results, all_scalars, all_options, profile_readiness)
        if results else ([], [], [], [])
    )
    profile_dimensions, profile_dimension_assignments, profile_dimension_kpis, profile_compositions = (
        build_profile_dimensions(results, all_scalars, all_options)
        if results else ([], [], [], [])
    )
    profile_context = (
        build_profile_context(results, all_scalars, all_options, profile_assignments)
        if results else []
    )
    profile_relationships, profile_relationship_assignments = (
        build_profile_relationships(results, profile_assignments)
        if results else ([], [])
    )
    template_models, template_assignments, template_grid_matrix, template_semantics, template_grid_headers = (
        build_template_semantic_model(topology_records, [result.grid for result in results])
        if any(record.object_type in {"networktemplate", "rangetemplate"} for record in topology_records)
        else ([], [], [], [], ["Model ID", "Template Type", "Dimension", "Parameter", "Parameterization Role"])
    )
    template_bundles = (
        build_template_bundles(topology_records, template_assignments)
        if template_assignments else []
    )
    template_bundle_models = (
        build_template_bundle_models(template_bundles, [result.grid for result in results])
        if template_bundles else []
    )
    standardization_excel = workbook_standardization_rows(standardization)
    decisions_excel = decision_rows(standardization)
    exceptions_excel = exception_rows(standardization)
    grids = [result.grid for result in results]
    grid_comparison, grid_comparison_headers = (grid_comparison_rows(standardization, grids) if results else ([], ["Category", "Parameter ID", "Parameter", "Scope", "Object Type", "Distinct Values", "Common Observed Value", "Consistency", "Local Overrides", "Coverage", "Notes"]))
    differences = difference_rows(standardization)
    effective_rows, effective_headers = assessment_table(all_scalars)
    manual_rows = [row for row in coverage if row.get("Collection Status") == "MANUAL_REVIEW_REQUIRED"]
    # Decision-support sheets intentionally precede technical evidence sheets.
    sheets = {
        **({
            "Template_Bundle_Models": template_bundle_models,
            "Template_Bundles": template_bundles,
            "Template_Models": template_models,
            "Template_Grid_Matrix": template_grid_matrix,
            "Template_Assignments": template_assignments,
            "Template_Semantics": template_semantics,
        } if template_models or template_assignments else {}),
        "Profile_Readiness": profile_readiness,
        "Profile_Populations": profile_populations,
        "Profile_Usefulness": profile_usefulness,
        "Profile_Fingerprints": profile_fingerprints,
        "Profile_Assignments": profile_assignments,
        "Profile_KPIs": profile_kpis,
        "Profile_Dimensions": profile_dimensions,
        "Profile_Dimension_Assignments": profile_dimension_assignments,
        "Profile_Dimension_KPIs": profile_dimension_kpis,
        "Profile_Compositions": profile_compositions,
        "Profile_Context": profile_context,
        "Profile_Relationships": profile_relationships,
        "Profile_Relationship_Assignments": profile_relationship_assignments,
        "Standardization": standardization_excel,
        "Decisions": decisions_excel,
        "Exceptions": exceptions_excel,
        "Grid_Comparison": grid_comparison,
        "Coverage": coverage,
        "Manual_Review": manual_rows,
        "Grid_Summary": [{"Grid": result.grid, "URL": result.grid_url, "WAPI Version": result.wapi_version,
                          "Collected At": result.collected_at, "Object": key,
                          "Raw Count": len(result.records.get(key, [])),
                          **_effective_query_summary(result, key), "Errors": len(result.errors)}
                         for result in results
                         for key in (sorted(set(result.records) | set(result.effective_records)) or [""])],
        "Errors": [row for result in results for row in result.errors],
        "DHCP_Effective": effective_rows,
        "DHCP_Options": all_options,
        "DHCP_Raw": [{"grid": result.grid, "object_type": object_type, **row}
                     for result in results for object_type, records in sorted(result.records.items())
                     for row in records if "options" in row or object_type.endswith(":dhcpproperties")
                     or any(key.startswith("use_") for key in row)],
        "Differences": differences,
        "Naming_Analysis": naming,
        "Collection_Sources": [row for result in results for row in result.collection_sources],
    }
    for result in results:
        for object_type, rows in sorted(result.records.items()):
            if object_type in TOPOLOGY_SHEETS:
                sheets.setdefault(TOPOLOGY_SHEETS[object_type], [])
                continue
            if object_type in RESERVATION_SHEETS:
                sheets.setdefault(RESERVATION_SHEETS[object_type], [])
                continue
            sheets.setdefault(_sheet_name(object_type), []).extend(
                {"grid": result.grid, **{key: _json_value(value) for key, value in row.items()}}
                for row in rows)
    topology_rows = topology_excel_rows(topology_records)
    sheet_headers = {
        'Template_Bundle_Models': TEMPLATE_BUNDLE_MODEL_HEADERS,
        'Template_Bundles': TEMPLATE_BUNDLE_HEADERS,
        'Template_Models': TEMPLATE_MODEL_HEADERS,
        'Template_Grid_Matrix': template_grid_headers,
        'Template_Assignments': TEMPLATE_ASSIGNMENT_HEADERS,
        'Template_Semantics': TEMPLATE_SEMANTIC_HEADERS,
        'Profile_Readiness': PROFILE_READINESS_HEADERS,
        'Profile_Populations': PROFILE_POPULATION_HEADERS,
        'Profile_Usefulness': PROFILE_USEFULNESS_HEADERS,
        'Profile_Fingerprints': PROFILE_FINGERPRINT_HEADERS,
        'Profile_Assignments': PROFILE_ASSIGNMENT_HEADERS,
        'Profile_KPIs': PROFILE_KPI_HEADERS,
        'Profile_Dimensions': PROFILE_DIMENSION_HEADERS,
        'Profile_Dimension_Assignments': PROFILE_DIMENSION_ASSIGNMENT_HEADERS,
        'Profile_Dimension_KPIs': PROFILE_DIMENSION_KPI_HEADERS,
        'Profile_Compositions': PROFILE_COMPOSITION_HEADERS,
        'Profile_Context': PROFILE_CONTEXT_HEADERS,
        'Profile_Relationships': PROFILE_RELATIONSHIP_HEADERS,
        'Profile_Relationship_Assignments': PROFILE_RELATIONSHIP_ASSIGNMENT_HEADERS,
        'Standardization': STANDARDIZATION_HEADERS,
        'Decisions': DECISION_HEADERS,
        'Exceptions': EXCEPTION_HEADERS,
        'Grid_Comparison': grid_comparison_headers,
        'Differences': DIFFERENCE_HEADERS,
        'DHCP_Effective': effective_headers,
    }
    for object_type, name in TOPOLOGY_SHEETS.items():
        if name in sheets:
            sheets[name].extend(topology_rows[object_type])
            sheet_headers[name] = topology_headers(object_type)
    reservation_rows = reservation_excel_rows(reservation_records)
    for object_type, name in RESERVATION_SHEETS.items():
        if name in sheets:
            sheets[name].extend(reservation_rows[object_type])
            sheet_headers[name] = reservation_headers(object_type)
    if any(key in RESERVATION_SHEETS for result in results for key in result.records):
        sheets["Reservation_Options"] = reservation_option_rows(reservation_records)
        sheet_headers["Reservation_Options"] = [
            "grid", "object_type", "object_ref", "name", "data_representation", "normalization_status",
            "issues", "option_field", "option_index", "option_name", "code", "vendor", "user_class",
            "option_type", "stored_value", "use_option", "use_options", "extra_fields",
        ]
        sheets["Reservation_Links"] = reservation_relationship_rows(reservation_records)
        sheet_headers["Reservation_Links"] = [
            "grid", "object_type", "object_ref", "name", "data_representation", "normalization_status",
            "issues", "relationship_field", "relationship_index", "target", "configured_rule_type",
            "use_logic_filter_rules", "relationship_data", "interpretation",
        ]
    if any(object_type.endswith("template") for result in results for object_type in result.records
           if object_type in TOPOLOGY_SHEETS):
        sheets["Template_Options"] = topology_option_rows(topology_records)
        sheet_headers["Template_Options"] = [
            "grid", "object_type", "object_ref", "name", "data_representation", "normalization_status",
            "issues", "option_name", "code", "vendor", "stored_value", "use_option", "use_options", "extra_fields",
        ]
    workbook = Workbook()
    workbook.remove(workbook.active)
    overview_links: dict[int, str] = {}
    if results:
        overview_links = _write_overview_dashboard(
            workbook, results, coverage, standardization, exceptions_excel,
            template_models, template_bundles, template_bundle_models
        )
    table_index = 2 if results else 1
    for index, (name, rows) in enumerate(sheets.items(), start=table_index):
        headers = sheet_headers.get(name) or sorted({key for row in rows for key in row}) or ["Status"]
        inventory_sheet = _write_inventory_sheet(
            workbook, index, name, rows, headers,
            preserve_order=name in {
                'Template_Bundle_Models', 'Template_Bundles', 'Template_Models',
                'Template_Grid_Matrix', 'Template_Assignments',
                'Template_Semantics', 'Profile_Readiness', 'Profile_Populations', 'Profile_Usefulness',
                'Profile_Fingerprints', 'Profile_Assignments', 'Profile_KPIs',
                'Profile_Dimensions', 'Profile_Dimension_Assignments',
                'Profile_Dimension_KPIs', 'Profile_Compositions',
                'Profile_Context', 'Profile_Relationships',
                'Profile_Relationship_Assignments',
            },
        )
        if name in {'Template_Models', 'Template_Bundle_Models'}:
            template_header_map = {cell.value: cell.column for cell in inventory_sheet[1]}
            technical_headers = ['Canonical Shape', 'Full SHA256']
            if name == 'Template_Bundle_Models':
                technical_headers.extend(['Network Models', 'Range Models', 'Geometry Signature'])
            for technical_header in technical_headers:
                technical_column = template_header_map.get(technical_header)
                if technical_column:
                    inventory_sheet.column_dimensions[
                        inventory_sheet.cell(1, technical_column).column_letter
                    ].hidden = True
        if name in {'Template_Bundles', 'Template_Bundle_Models'}:
            bundle_headers = {cell.value: cell.column for cell in inventory_sheet[1]}
            attention_column = bundle_headers.get('Attention')
            review_column = bundle_headers.get('Review Flags')
            if attention_column:
                for row_number in range(2, inventory_sheet.max_row + 1):
                    cell = inventory_sheet.cell(row_number, attention_column)
                    if cell.value == 'REVIEW':
                        cell.fill = PatternFill('solid', fgColor='FFF2CC')
                        cell.font = Font(name='Calibri', size=10, bold=True, color='9C6500')
                    elif cell.value == 'OK':
                        cell.fill = PatternFill('solid', fgColor='E2F0D9')
            if review_column:
                for row_number in range(2, inventory_sheet.max_row + 1):
                    cell = inventory_sheet.cell(row_number, review_column)
                    if cell.value:
                        cell.fill = PatternFill('solid', fgColor='FFF2CC')
                        cell.font = Font(name='Calibri', size=10, bold=True, color='9C6500')
        if name == 'Template_Bundles':
            bundle_headers = {cell.value: cell.column for cell in inventory_sheet[1]}
            for technical_header in (
                'Network Model ID', 'Network Member Family', 'Range Model ID',
                'Range Association', 'Range Addresses', 'Geometry Signature',
                'Fixed Address Templates', 'Stored Enabled Options',
            ):
                technical_column = bundle_headers.get(technical_header)
                if technical_column:
                    inventory_sheet.column_dimensions[
                        inventory_sheet.cell(1, technical_column).column_letter
                    ].hidden = True
        if name == 'Profile_Populations':
            population_sheet = inventory_sheet
            population_headers = {cell.value: cell.column for cell in population_sheet[1]}
            share_column = population_headers.get('Share %')
            if share_column:
                for row_number in range(2, population_sheet.max_row + 1):
                    population_sheet.cell(row_number, share_column).number_format = '0.0"%"'
        if name in {
            'Template_Bundle_Models', 'Template_Bundles', 'Template_Models',
            'Template_Grid_Matrix', 'Template_Assignments',
            'Profile_Readiness', 'Profile_Usefulness', 'Profile_Fingerprints',
            'Profile_Assignments', 'Profile_KPIs', 'Profile_Dimensions',
            'Profile_Dimension_Assignments', 'Profile_Dimension_KPIs',
            'Profile_Compositions', 'Profile_Context', 'Profile_Relationships',
            'Profile_Relationship_Assignments', 'Standardization', 'Decisions',
            'Exceptions', 'Grid_Comparison',
        }:
            _style_decision_support(inventory_sheet, name)
    # Keep the workbook navigable at nine-Grid scale: decision/catalog views stay
    # visible, while detailed evidence remains in the same workbook as drill-down.
    if template_bundle_models:
        user_facing_sheets = {
            'Overview', 'Data_Index', 'Template_Bundle_Models', 'Template_Bundles',
            'Template_Models', 'Template_Grid_Matrix', 'Grid_Summary',
        }
        if sheets.get('Manual_Review'):
            user_facing_sheets.add('Manual_Review')
        if any(result.errors for result in results):
            user_facing_sheets.add('Errors')
    else:
        user_facing_sheets = {
            'Overview', 'Data_Index', 'Standardization', 'Decisions', 'Exceptions',
            'Grid_Summary', 'Manual_Review', 'Errors',
        }
    for worksheet in workbook.worksheets:
        if worksheet.title not in user_facing_sheets:
            worksheet.sheet_state = 'hidden'

    # One navigable catalog replaces dozens of visible technical tabs.
    if results:
        data_index = workbook.create_sheet('Data_Index', 1)
        data_index.sheet_view.showGridLines = False
        index_headers = ['Category', 'Sheet', 'Rows', 'Visibility', 'Purpose', 'Open']
        for column, header in enumerate(index_headers, start=1):
            cell = data_index.cell(1, column, header)
            cell.fill = PatternFill('solid', fgColor='0F243E')
            cell.font = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        purposes = {
            'Template_Bundle_Models': 'Cross-Grid catalog of normalized provisioning bundle shapes.',
            'Template_Bundles': 'NetworkTemplate→RangeTemplate assignments, geometry and review flags.',
            'Template_Models': 'Reusable Network/Range template semantic shapes.',
            'Template_Grid_Matrix': 'Grid-local values behind reusable template models.',
            'Template_Assignments': 'Template→semantic model assignments.',
            'Standardization': 'Parameter evidence and standardization questions.',
            'Decisions': 'Human approval/target decisions.',
            'Exceptions': 'Deviation and exception backlog.',
            'Grid_Summary': 'Per-Grid collection/object counts.',
            'Coverage': 'Detailed collection/evidence completeness.',
            'DHCP_Effective': 'Normalized effective DHCP scalar evidence.',
            'DHCP_Options': 'Normalized effective DHCP option evidence.',
            'DHCP_Raw': 'Raw DHCP configuration evidence.',
            'Profile_Compositions': 'Observed Network/Range dimension composition.',
            'Network_Templates': 'Typed raw NetworkTemplate evidence.',
            'Range_Templates': 'Typed raw RangeTemplate evidence.',
            'Template_Options': 'Stored template option evidence and use flags.',
        }
        index_row = 2
        for worksheet in workbook.worksheets:
            if worksheet.title in {'Overview', 'Data_Index'}:
                continue
            visible = worksheet.sheet_state == 'visible'
            if visible:
                category = 'Primary analysis'
            elif worksheet.title in {'Standardization', 'Decisions', 'Exceptions'}:
                category = 'Decision workflow'
            else:
                category = 'Technical evidence'
            values = [
                category,
                worksheet.title,
                max(worksheet.max_row - 1, 0),
                'VISIBLE' if visible else 'HIDDEN',
                purposes.get(worksheet.title, 'Supporting inventory/evidence view.'),
                'Open' if visible else (
                    'Unhide for decisions' if category == 'Decision workflow'
                    else 'Unhide for drill-down'
                ),
            ]
            for column, value in enumerate(values, start=1):
                cell = data_index.cell(index_row, column, value)
                cell.alignment = Alignment(wrap_text=True, vertical='top')
            if visible:
                link_cell = data_index.cell(index_row, 6)
                link_cell.hyperlink = Hyperlink(
                    ref=link_cell.coordinate, location=f"'{worksheet.title}'!A1"
                )
                link_cell.font = Font(name='Calibri', size=10, color='0563C1', underline='single')
            index_row += 1
        data_index.freeze_panes = 'A2'
        for column, width in {'A': 18, 'B': 30, 'C': 10, 'D': 12, 'E': 55, 'F': 20}.items():
            data_index.column_dimensions[column].width = width
        if index_row > 2:
            index_table = Table(displayName='DataIndexTable', ref=f'A1:F{index_row - 1}')
            index_table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
            data_index.add_table(index_table)
    if results:
        _wire_overview_standardization_links(workbook, overview_links)
        if not template_bundle_models:
            _write_overview_profile_summary(
                workbook, profile_summaries, profile_kpis, profile_relationships
            )
    workbook.save(output / "current_state_inventory.xlsx")
    if results:
        write_decision_template(standardization, output / "standardization_decisions.template.yaml")
    summary = ["# Infoblox current-state inventory", "",
               "This is an evidence-gathering report; common values are not approved standards.", "",
               "Raw/local values and WAPI effective inheritance evidence are reported separately. "
               "PARTIAL normalized rows have no confirmed effective value.", "", "## Grids collected", "",
               "| Grid | WAPI version | Collection timestamp (UTC) |", "| --- | --- | --- |"]
    for result in results:
        summary.append(f"| {_markdown_value(result.grid)} | {_markdown_value(result.wapi_version)} | "
                       f"{_markdown_value(result.collected_at)} |")
    summary.extend(["", "Start with the workbook Overview for inventory counts, coverage and observed effective DHCP values. "
                    "Collection_Sources retains each object's collection timestamp. Combining archives is not an atomic "
                    "appliance snapshot; conflicting object snapshots are rejected. RAW and effective evidence remain separate."])
    summary.extend(["", "## Collection errors", ""])
    errors = [row for result in results for row in result.errors]
    summary.extend(f"- {_markdown_value(row.get('grid'))}: {_markdown_value(row.get('object_type'))}: "
                   f"{_markdown_value(row.get('error'))}" for row in errors)
    if not errors:
        summary.append("No collection errors recorded.")
    summary.extend(["", "## Coverage", "",
                    "Status counts summarize collection, query, and field observations. "
                    "The workbook Coverage sheet contains object/query/field details and notes.", "",
                    "| Grid | Area | Status counts |", "| --- | --- | --- |"])
    coverage_counts: dict[tuple[str, str], Counter[str]] = defaultdict(Counter)
    for row in coverage:
        coverage_counts[(str(row.get("Grid", "")), str(row.get("Area", "")))][str(row.get("Collection Status", ""))] += 1
    for (grid, area), counts in sorted(coverage_counts.items()):
        summary.append(f"| {_markdown_value(grid)} | {_markdown_value(area)} | "
                       + ", ".join(f"{_markdown_value(status)}: {count}" for status, count in sorted(counts.items())) + " |")
    if any(object_type in TOPOLOGY_SHEETS for result in results for object_type in result.records):
        summary.extend(["", "## Reusable DHCP configuration", "",
                        "Topology sheets show typed RAW_CONFIGURATION values and observed use flags, "
                        "not effective values for objects created from templates. Template_Options exposes "
                        "stored option values separately. Empty inventories retain column headers; "
                        "consult Coverage to distinguish EMPTY, unavailable, and failed queries. "
                        "Normalization coverage describes parsed records only; raw-query coverage "
                        "remains authoritative for collection completeness. Unknown fields remain in "
                        "extra_fields and complete responses remain in the raw archive."])
    if template_bundles:
        bundle_status = Counter(str(row.get("Bundle Status") or "") for row in template_bundles)
        review_bundles = sum(str(row.get("Attention") or "") == "REVIEW" for row in template_bundles)
        summary.extend([
            "", "## Template provisioning overview", "",
            f"Provisioning bundles: {len(template_bundles)}; linked: {bundle_status['LINKED']}; "
            f"network-only: {bundle_status['NETWORK_ONLY']}; missing linked RangeTemplate: "
            f"{bundle_status['MISSING_RANGE_TEMPLATE']}; orphan RangeTemplate: "
            f"{bundle_status['ORPHAN_RANGE_TEMPLATE']}; bundles needing review: {review_bundles}.",
            "",
            "Start with Template_Bundle_Models for normalized provisioning patterns, then "
            "Template_Bundles for concrete NetworkTemplate→RangeTemplate assignments and "
            "Template_Grid_Matrix for local parameter values. "
            "Detailed template evidence remains in hidden technical sheets.",
        ])

    if any(key in RESERVATION_SHEETS for result in results for key in result.records):
        summary.extend(["", "## DHCP reservations and filters", "",
                        "Reservation/filter sheets contain typed RAW_CONFIGURATION evidence. "
                        "Reservation_Options and Reservation_Links retain nested values, use flags and source indexes. "
                        "Fixed-address effective evidence continues through the existing DHCP sheets. "
                        "Host DHCP configuration status concerns explicit IPv4 configure_for_dhcp flags only; "
                        "it does not assess IPv6, service activation or lease availability. Roaming hosts are not "
                        "classified as fixed-IP reservations. Template origin is unavailable when WAPI exposes "
                        "only a write-only template field. Filter names, including Corporate or Allowed, imply "
                        "no policy meaning. Consumers outside this collection are not assessed. "
                        "DDNS/EA fields are retained solely as reservation metadata. For superhostchild, "
                        "an empty complete parent inventory means no child request; Coverage records that dependency."])
    if profile_kpis:
        summary.extend([
            "", "## Observed profile discovery", "",
            "Fingerprint prevalence is descriptive current-state evidence. The most common fingerprint is not an approved standard. "
            "Objects with unresolved required core inputs remain outside normal profiles.", "",
            "| Profile | Applicable | Profiled | Profiled % | Unresolved | Distinct profiles | Top-1 share of profiled % | Top-3 share of profiled % | Singletons |",
            "| --- | ---: | ---: | ---: | ---: | ---: | ---: | ---: | ---: |",
        ])
        for item in profile_kpis:
            summary.append(
                f"| {_markdown_value(item.get('Profile'))} | {_markdown_value(item.get('Applicable Objects'))} | "
                f"{_markdown_value(item.get('Profiled Objects'))} | {_markdown_value(item.get('Profiled %'))} | "
                f"{_markdown_value(item.get('Unresolved Profile Objects'))} | {_markdown_value(item.get('Distinct Profiles'))} | "
                f"{_markdown_value(item.get('Top-1 Share of Profiled %'))} | "
                f"{_markdown_value(item.get('Top-3 Share of Profiled %'))} | "
                f"{_markdown_value(item.get('Singleton Profiles'))} |"
            )

    if any((item.get("Applicable Objects") or 0) > 0 for item in profile_dimension_kpis):
        summary.extend([
            "", "## Observed profile dimensions", "",
            "Dimension fingerprints separate DHCP Core, DNS/DDNS, PXE, Service and Range Association. "
            "They are descriptive current-state evidence, not approved standards. "
            "NOT_APPLICABLE is an explicit applicability state; deferred composite inputs are excluded from identity hashes.", "",
            "| Profile | Dimension | Applicable | Profiled | Profiled % | Unresolved | Distinct fingerprints | Top-1 share of profiled % | Singletons |",
            "| --- | --- | ---: | ---: | ---: | ---: | ---: | ---: | ---: |",
        ])
        for item in profile_dimension_kpis:
            summary.append(
                f"| {_markdown_value(item.get('Profile'))} | {_markdown_value(item.get('Dimension'))} | "
                f"{_markdown_value(item.get('Applicable Objects'))} | {_markdown_value(item.get('Profiled Objects'))} | "
                f"{_markdown_value(item.get('Profiled %'))} | {_markdown_value(item.get('Unresolved Dimension Objects'))} | "
                f"{_markdown_value(item.get('Distinct Fingerprints'))} | "
                f"{_markdown_value(item.get('Top-1 Share of Profiled %'))} | "
                f"{_markdown_value(item.get('Singleton Fingerprints'))} |"
            )

    if profile_relationships:
        associated = sum(int(row.get("Range Count") or 0) for row in profile_relationships)
        paired_rows = [row for row in profile_relationships if row.get("Relationship Status") == "PROFILE_PAIR"]
        paired = sum(int(row.get("Range Count") or 0) for row in paired_rows)
        summary.extend([
            "", "## Observed Network↔Range profile relationships", "",
            "Relationships describe current topology/context only; recurring pairs are not approved standards.", "",
            f"Associated ranges: {associated}; ranges with both parent Network and Range fingerprint resolved: {paired}.", "",
            "| Relationship | Network fingerprint | Range fingerprint | Association | Range count | Share associated % | Share paired-profiled % |",
            "| --- | --- | --- | --- | ---: | ---: | ---: |",
        ])
        for row in profile_relationships[:20]:
            summary.append(
                f"| {_markdown_value(row.get('Relationship Status'))} | "
                f"{_markdown_value(row.get('Network Fingerprint ID'))} | "
                f"{_markdown_value(row.get('Range Fingerprint ID'))} | "
                f"{_markdown_value(row.get('Association Family'))} | "
                f"{_markdown_value(row.get('Range Count'))} | "
                f"{_markdown_value(row.get('Share of Associated Ranges %'))} | "
                f"{_markdown_value(row.get('Share of Paired Profiled Ranges %'))} |"
            )

    changed = [row for row in differences if row.get("Classification") not in {"CONSISTENT", "NOT_CONFIGURED", "NO_OBJECTS_IN_SCOPE"}]
    summary.extend(["", "## Standardization observations", "",
                    "Observed differences and common values are descriptive. They are not approved standards. "
                    "Use Standardization and Decisions to record human review, then Exceptions for actionable deviations. "
                    "Additional DDNS/EA/DNS collection depth and broader policy interpretation remain to be implemented.", ""])
    summary.extend(
        f"- {_markdown_value(row.get('Parameter'))} [{_markdown_value(row.get('Scope'))}]: "
        f"{_markdown_value(row.get('Observed Values') or row.get('Classification'))}; "
        f"classification={_markdown_value(row.get('Classification'))}, local overrides={_markdown_value(row.get('Local Overrides'))}."
        for row in changed[:20]
    )
    if not changed:
        summary.append("No standardization differences identified in the currently confirmed parameter set.")
    candidates = [row for row in all_options + all_scalars
                  if row["configured_here"] is True and row["object_type"] != "grid:dhcpproperties"]
    summary.extend(["", "## Local exception candidates", "",
                    f"{len(candidates)} confirmed local parameter/option overrides; local configuration alone "
                    "does not establish a policy exception."])
    summary.extend(f"- {_markdown_value(row['grid'])}: {_markdown_value(row['object_name'])}, "
                   f"{_markdown_value(row['parameter'])}: {_markdown_value(row['effective_value'])}."
                   for row in candidates[:20])
    inherited = [row for row in all_options + all_scalars if row["inherited"] is True]
    multiple = [row for row in all_options + all_scalars if row["multisource"] is True]
    summary.extend(["", "## Inheritance observations", "",
                    f"{len(inherited)} rows indicate inheritance; {len(multiple)} rows indicate multiple sources. "
                    "Source references and complete inheritance groups are retained in the workbook.", "",
                    "## Manual review", "",
                    "Approval process and documentation requirements require confirmation with Grid owners. "
                    "WAPI observations do not establish organizational policy.", "", "## Files generated", "",
                    "- current_state_inventory.xlsx", "- standardization_decisions.template.yaml", "- current_state_summary.md", "- manual_review.md"])
    (output / "current_state_summary.md").write_text("\n".join(summary) + "\n", encoding="utf-8")
    manual = ["# Manual review required", "", "Approval workflow and formal documentation policy are not reliably represented by DHCP WAPI objects.", ""]
    manual.extend(f"- {result.grid}: confirm approval process and documentation policy with Grid owners." for result in results)
    (output / "manual_review.md").write_text("\n".join(manual) + "\n", encoding="utf-8")


def _sheet_name(object_type: str) -> str:
    return {"networkview": "Network_Views", "network": "Networks", "range": "Ranges",
            "grid:dhcpproperties": "Grid_DHCP", "member:dhcpproperties": "Member_DHCP",
            "member": "Members_Failover", **TOPOLOGY_SHEETS, **RESERVATION_SHEETS,
            "extensibleattributedef": "Extensible_Attributes"}.get(
                object_type, re.sub(r"[\\/*?:\[\]]", "_", object_type)[:31])
